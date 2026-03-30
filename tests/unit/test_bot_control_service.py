import asyncio
import json
from dataclasses import dataclass
from datetime import UTC, datetime
from io import BytesIO
from uuid import UUID

import pytest
from openpyxl import load_workbook

from extg_bot_ui.application.bot_control import (
    BotOperatorContext,
    MigrationBotControlService,
    MigrationRunOptions,
)
from extg_bot_ui.application.identity_matrix_workbook import IdentityMatrixWorkbookRow
from extg_bot_ui.bootstrap.config import BotUiSettings
from extg_migration_runtime.application.identity import IdentityLookupResult
from extg_migration_runtime.infrastructure.archive.telegram_export_archive_parser import (
    TelegramExportArchiveParser,
)
from extg_migration_runtime.infrastructure.archive.stage_store import (
    InMemoryTelegramExportArchiveStageStore,
)
from extg_migration_runtime.application.use_cases.reconcile_migration import (
    ReconcileMigrationResult,
)
from extg_migration_runtime.application.target_chat_provisioning import (
    ParticipantAccessApplyResult,
    ParticipantAccessFailure,
)
from extg_migration_runtime.infrastructure.express.fake_gateway import FakeExpressGateway
from extg_migration_runtime.infrastructure.persistence.in_memory import (
    InMemoryChatMappingRepository,
    InMemoryMigrationJobRepository,
    InMemoryOutboxRepository,
    InMemoryTelegramExportSnapshotRepository,
)
from extg_shared.config.common import (
    BackfillSettings,
    BotSettings,
    ExpressSettings,
    PostgresSettings,
    RetrySettings,
    TelegramSettings,
)
from extg_shared.contracts.errors import ConfigurationError, FatalItemError
from extg_shared.contracts.models import (
    ChatMappingRecord,
    ChatMigrationConfigRecord,
    HistoryCursor,
    IdentityMappingRecord,
    IntegrationOutboxStatus,
    MessageImportStatus,
    MigrationCheckpoint,
    MigrationJobRecord,
    MigrationJobStatus,
    SourceChannelAccessProfile,
    SourceDialog,
    SourceParticipant,
    SourceTopic,
)
from extg_shared.contracts.ports import MessageStatusSummary
from extg_shared.utils import get_current_express_cts_host
from extg_telethon_service.application.telegram_session_context import TelegramSessionContext


@dataclass
class StubUseCase:
    result: object | None = None

    def __post_init__(self) -> None:
        self.commands = []
        self.called = asyncio.Event()

    async def execute(self, command):
        self.commands.append(command)
        self.called.set()
        return self.result


class StubCheckpointRepository:
    def __init__(self, checkpoints=None) -> None:
        self._checkpoints = checkpoints or {}

    async def get(self, migration_id: str, source_chat_id: str, mode: str):
        return self._checkpoints.get((migration_id, source_chat_id, mode))


class StubMessageMappingRepository:
    def __init__(self, summaries=None) -> None:
        self._summaries = summaries or []

    async def summarize_by_chat(self, migration_id: str):
        return list(self._summaries)


class StubMigrationStateRepository:
    async def get(self, migration_id: str):
        return None


class StubChatMappingRepository(InMemoryChatMappingRepository):
    pass


class StubTelegramGateway:
    def __init__(self, dialogs=None) -> None:
        self._dialogs = dialogs or [
            SourceDialog(
                dialog_id="chat-1",
                chat_type="supergroup",
                title="Imported Chat",
                message_count=10,
                media_count=2,
                approximate_bytes=1000,
            ),
        ]
        self._participants = {}
        self._participant_errors = {}
        self._topics = {}
        self._channel_access_profiles = {}

    async def list_available_dialogs(
        self,
        *,
        limit: int = 100,
        query: str | None = None,
        source_backend: str = "telethon_user_session",
    ):
        return self._dialogs[:limit]

    async def list_participants(
        self,
        dialog_id: str,
        *,
        source_backend: str = "telethon_user_session",
    ):
        error = self._participant_errors.get(dialog_id)
        if error is not None:
            raise error
        return list(self._participants.get(dialog_id, []))

    async def list_topics(
        self,
        dialog_id: str,
        *,
        source_backend: str = "telethon_user_session",
    ):
        return list(self._topics.get(dialog_id, []))

    async def get_channel_access_profile(
        self,
        dialog_id: str,
        *,
        source_backend: str = "telethon_user_session",
    ):
        return self._channel_access_profiles.get(
            dialog_id,
            SourceChannelAccessProfile(
                dialog_id=dialog_id,
                is_admin=True,
                can_list_participants=True,
            ),
        )


class StubChatMigrationConfigRepository:
    def __init__(self) -> None:
        self._items = {}

    async def get(self, migration_id: str, source_chat_id: str):
        return self._items.get((migration_id, source_chat_id))

    async def list_by_migration(self, migration_id: str):
        return [
            record
            for (current_migration_id, _), record in self._items.items()
            if current_migration_id == migration_id
        ]

    async def save(self, record: ChatMigrationConfigRecord):
        key = (record.migration_id, record.source_chat_id)
        existing = self._items.get(key)
        if existing is not None and existing.updated_by_huid != record.updated_by_huid:
            return existing
        self._items[key] = record
        return record

    async def delete(self, migration_id: str, source_chat_id: str):
        return self._items.pop((migration_id, source_chat_id), None) is not None


class StubAuditRepository:
    def __init__(self) -> None:
        self.events = []

    async def add(self, event):
        self.events.append(event)


class StubLogger:
    def exception(self, *_args, **_kwargs) -> None:
        return None


class StubIdentityMappingRepository:
    def __init__(self) -> None:
        self._items_by_user_id = {}
        self._items_by_username = {}

    async def get_by_user_id(self, telegram_user_id: str, *, express_host: str | None = None):
        return self._items_by_user_id.get(telegram_user_id.strip())

    async def get_by_username(self, telegram_username: str, *, express_host: str | None = None):
        key = telegram_username.strip().lstrip("@").lower()
        return self._items_by_username.get(key)

    async def save(self, record: IdentityMappingRecord):
        normalized_user_id = record.telegram_user_id.strip() if record.telegram_user_id else None
        normalized_username = (
            record.telegram_username.strip().lstrip("@").lower()
            if record.telegram_username
            else None
        )
        normalized = IdentityMappingRecord(
            telegram_user_id=normalized_user_id,
            telegram_username=normalized_username,
            telegram_display_name=record.telegram_display_name,
            corporate_email=record.corporate_email.strip().lower()
            if record.corporate_email is not None
            else None,
            target_huid=record.target_huid,
            created_at=record.created_at,
            updated_at=record.updated_at,
            last_resolved_at=record.last_resolved_at,
        )
        if normalized_user_id is not None:
            self._items_by_user_id[normalized_user_id] = normalized
        if normalized_username is not None:
            self._items_by_username[normalized_username] = normalized
        return normalized


class StubIdentityDirectory:
    def __init__(
        self,
        *,
        identity_mapping_repository: StubIdentityMappingRepository,
        user_huid_by_email: dict[str, str] | None = None,
        user_huid_by_ad_login: dict[str, str] | None = None,
        user_huid_by_other_id: dict[str, str] | None = None,
    ) -> None:
        self._identity_mapping_repository = identity_mapping_repository
        self._user_huid_by_email = {
            email.strip().lower(): huid
            for email, huid in (user_huid_by_email or {}).items()
        }
        self._user_huid_by_ad_login = {
            login.strip().lower(): huid
            for login, huid in (user_huid_by_ad_login or {}).items()
        }
        self._user_huid_by_other_id = {
            other_id.strip(): huid
            for other_id, huid in (user_huid_by_other_id or {}).items()
        }

    def normalize_username(self, telegram_username: str | None) -> str | None:
        if telegram_username is None:
            return None
        normalized = telegram_username.strip().lstrip("@").lower()
        return normalized or None

    async def resolve_username(self, telegram_username: str | None) -> IdentityLookupResult:
        return await self.resolve_telegram_identity(
            telegram_user_id=None,
            telegram_username=telegram_username,
            telegram_display_name=None,
        )

    async def resolve_telegram_identity(
        self,
        *,
        telegram_user_id: str | None,
        telegram_username: str | None,
        telegram_display_name: str | None = None,
    ) -> IdentityLookupResult:
        normalized_user_id = telegram_user_id.strip() if telegram_user_id else None
        normalized = self.normalize_username(telegram_username)
        mapping = None
        if normalized_user_id is not None:
            mapping = await self._identity_mapping_repository.get_by_user_id(normalized_user_id)
        if mapping is None and normalized is not None:
            mapping = await self._identity_mapping_repository.get_by_username(normalized)
        if normalized_user_id is None and normalized is None:
            return IdentityLookupResult(
                telegram_user_id=None,
                telegram_username=None,
                telegram_display_name=telegram_display_name,
                corporate_email=None,
                target_huid=None,
                resolution_source="display_only",
                reason="missing_identity",
            )
        if mapping is None:
            return IdentityLookupResult(
                telegram_user_id=normalized_user_id,
                telegram_username=normalized,
                telegram_display_name=telegram_display_name,
                corporate_email=None,
                target_huid=None,
                resolution_source="display_only",
                reason="identity_not_mapped",
            )
        target_huid = mapping.target_huid or self._user_huid_by_email.get(
            mapping.corporate_email,
        )
        if target_huid and mapping.target_huid != target_huid:
            await self._identity_mapping_repository.save(
                IdentityMappingRecord(
                    telegram_user_id=normalized_user_id or mapping.telegram_user_id,
                    telegram_username=mapping.telegram_username,
                    telegram_display_name=telegram_display_name or mapping.telegram_display_name,
                    corporate_email=mapping.corporate_email,
                    target_huid=target_huid,
                    created_at=mapping.created_at,
                    updated_at=datetime(2026, 3, 18, 10, 0, tzinfo=UTC),
                    last_resolved_at=datetime(2026, 3, 18, 10, 0, tzinfo=UTC),
                ),
            )
        return IdentityLookupResult(
            telegram_user_id=normalized_user_id or mapping.telegram_user_id,
            telegram_username=normalized,
            telegram_display_name=telegram_display_name or mapping.telegram_display_name,
            corporate_email=mapping.corporate_email,
            target_huid=target_huid,
            resolution_source=(
                "identity_map_email" if target_huid and mapping.target_huid is None else "identity_map_cache"
            ),
            reason=None if target_huid else "email_not_found_in_express",
        )

    async def upsert_mapping(
        self,
        *,
        telegram_user_id: str | None,
        telegram_username: str | None,
        telegram_display_name: str | None,
        corporate_email: str,
    ) -> IdentityLookupResult:
        normalized_user_id = telegram_user_id.strip() if telegram_user_id else None
        normalized_username = self.normalize_username(telegram_username)
        saved = await self._identity_mapping_repository.save(
            IdentityMappingRecord(
                telegram_user_id=normalized_user_id,
                telegram_username=normalized_username,
                telegram_display_name=telegram_display_name,
                corporate_email=corporate_email.strip().lower(),
                target_huid=None,
                created_at=datetime(2026, 3, 18, 10, 0, tzinfo=UTC),
                updated_at=datetime(2026, 3, 18, 10, 0, tzinfo=UTC),
                last_resolved_at=None,
            ),
        )
        return await self.resolve_telegram_identity(
            telegram_user_id=saved.telegram_user_id,
            telegram_username=saved.telegram_username,
            telegram_display_name=saved.telegram_display_name,
        )

    async def upsert_direct_mapping(
        self,
        *,
        telegram_user_id: str | None,
        telegram_username: str | None,
        telegram_display_name: str | None,
        target_huid: str,
        corporate_email: str | None = None,
    ) -> IdentityLookupResult:
        normalized_user_id = telegram_user_id.strip() if telegram_user_id else None
        normalized_username = self.normalize_username(telegram_username)
        if not await self._identity_mapping_repository.save(
            IdentityMappingRecord(
                telegram_user_id=normalized_user_id,
                telegram_username=normalized_username,
                telegram_display_name=telegram_display_name,
                corporate_email=corporate_email.strip().lower()
                if corporate_email is not None
                else None,
                target_huid=target_huid.strip(),
                created_at=datetime(2026, 3, 18, 10, 0, tzinfo=UTC),
                updated_at=datetime(2026, 3, 18, 10, 0, tzinfo=UTC),
                last_resolved_at=datetime(2026, 3, 18, 10, 0, tzinfo=UTC),
            ),
        ):
            raise AssertionError("save should return a record")
        return await self.resolve_telegram_identity(
            telegram_user_id=normalized_user_id,
            telegram_username=normalized_username,
            telegram_display_name=telegram_display_name,
        )

    async def resolve_corporate_selector(
        self,
        *,
        corporate_email: str | None = None,
        target_huid: str | None = None,
        ad_login: str | None = None,
        other_id: str | None = None,
        default_ad_domain: str | None = None,
    ) -> IdentityLookupResult:
        selector_count = sum(
            1
            for value in (corporate_email, target_huid, ad_login, other_id)
            if value is not None and str(value).strip()
        )
        if selector_count != 1:
            raise ConfigurationError("exactly one selector is required")
        if corporate_email:
            normalized_email = corporate_email.strip().lower()
            return IdentityLookupResult(
                telegram_user_id=None,
                telegram_username=None,
                corporate_email=normalized_email,
                target_huid=self._user_huid_by_email.get(normalized_email),
                resolution_source="manual_selector_email",
                reason=None
                if normalized_email in self._user_huid_by_email
                else "email_not_found_in_express",
            )
        if target_huid:
            normalized_huid = target_huid.strip()
            return IdentityLookupResult(
                telegram_user_id=None,
                telegram_username=None,
                corporate_email=None,
                target_huid=normalized_huid,
                resolution_source="manual_selector_huid",
            )
        if ad_login:
            normalized_login = ad_login.strip().lower()
            return IdentityLookupResult(
                telegram_user_id=None,
                telegram_username=None,
                corporate_email=None,
                target_huid=self._user_huid_by_ad_login.get(normalized_login),
                resolution_source="manual_selector_ad_login",
                reason=None
                if normalized_login in self._user_huid_by_ad_login
                else "ad_login_not_found_in_express",
            )
        normalized_other_id = (other_id or "").strip()
        return IdentityLookupResult(
            telegram_user_id=None,
            telegram_username=None,
            corporate_email=None,
            target_huid=self._user_huid_by_other_id.get(normalized_other_id),
            resolution_source="manual_selector_other_id",
            reason=None
            if normalized_other_id in self._user_huid_by_other_id
            else "other_id_not_found_in_express",
        )


class StubTargetChatProvisioningService:
    def __init__(
        self,
        *,
        validation_errors=None,
        telegram_gateway: StubTelegramGateway | None = None,
        identity_directory: StubIdentityDirectory | None = None,
    ) -> None:
        self._validation_errors = validation_errors or {}
        self._telegram_gateway = telegram_gateway
        self._identity_directory = identity_directory
        self.validated_dialogs = []
        self.member_add_calls = []
        self.member_add_result = ParticipantAccessApplyResult()

    async def validate_dialog_preconditions(self, *, migration_id: str, dialog) -> None:
        self.validated_dialogs.append((migration_id, dialog.source_chat_id))
        error = self._validation_errors.get(dialog.source_chat_id)
        if error is not None:
            raise error
        if dialog.source_chat_type == "channel":
            profile = await self.get_channel_access_profile(
                source_chat_id=dialog.source_chat_id,
                source_backend=dialog.source_backend,
            )
            if not profile.is_admin:
                raise FatalItemError(
                    "channel migration requires Telegram admin rights for the connected user session; "
                    f"source_chat_id={dialog.source_chat_id}. Reconnect the Telegram account that is an admin "
                    "of this channel or delegate the migration to a channel administrator.",
                )

    async def list_participant_matrix(
        self,
        *,
        source_chat_id: str,
        participant_source_chat_id: str | None = None,
        source_backend: str = "telethon_user_session",
    ):
        participants = await self._telegram_gateway.list_participants(
            participant_source_chat_id or source_chat_id,
            source_backend=source_backend,
        )
        rows = []
        for participant in participants:
            lookup = await self._identity_directory.resolve_telegram_identity(
                telegram_user_id=participant.external_id,
                telegram_username=participant.username,
                telegram_display_name=participant.display_name,
            )
            rows.append(
                type(
                    "ParticipantMatrixRow",
                    (),
                    {
                        "telegram_user_id": lookup.telegram_user_id or participant.external_id,
                        "telegram_username": lookup.telegram_username or participant.username,
                        "telegram_display_name": participant.display_name,
                        "is_self": participant.is_self,
                        "corporate_email": lookup.corporate_email,
                        "target_huid": lookup.target_huid,
                        "resolution_source": lookup.resolution_source,
                        "reason": lookup.reason,
                    },
                )(),
            )
        return rows

    async def get_channel_access_profile(
        self,
        *,
        source_chat_id: str,
        source_backend: str = "telethon_user_session",
    ):
        return await self._telegram_gateway.get_channel_access_profile(
            source_chat_id,
            source_backend=source_backend,
        )

    async def add_resolved_members_to_existing_chat(
        self,
        *,
        migration_id: str,
        source_chat_id: str,
        source_chat_type: str,
        source_backend: str,
        target_chat_id: str,
        target_chat_title: str,
        participant_targets: list,
        initiator_huid: str | None,
        access_strategy: str,
        anchor_cts_host: str | None,
    ):
        self.member_add_calls.append(
            {
                "migration_id": migration_id,
                "source_chat_id": source_chat_id,
                "source_chat_type": source_chat_type,
                "source_backend": source_backend,
                "target_chat_id": target_chat_id,
                "target_chat_title": target_chat_title,
                "participant_targets": participant_targets,
                "initiator_huid": initiator_huid,
                "access_strategy": access_strategy,
                "anchor_cts_host": anchor_cts_host,
            },
        )
        return self.member_add_result


class StubOperatorTelegramSessionRepository:
    async def get_by_operator(self, operator_huid: str):
        return None

    async def save(self, record):
        return record

    async def delete(self, operator_huid: str):
        return False


class StubTelegramSessionService:
    def __init__(
        self,
        *,
        resolved_session_string: str | None = None,
        error: Exception | None = None,
    ) -> None:
        self._resolved_session_string = resolved_session_string
        self._error = error
        self.calls: list[tuple[str, bool]] = []

    async def resolve_session_string(
        self,
        *,
        operator_huid: str,
        mark_used: bool = True,
    ) -> str:
        self.calls.append((operator_huid, mark_used))
        if self._error is not None:
            raise self._error
        if self._resolved_session_string is None:
            raise ConfigurationError("telegram session is not connected for this user")
        return self._resolved_session_string


def build_settings(*, operator_huids=("operator-1",)) -> BotUiSettings:
    return BotUiSettings.model_construct(
        environment="test",
        timezone_name="Europe/Moscow",
        retry=RetrySettings(),
        backfill=BackfillSettings(batch_size=20),
        telegram=TelegramSettings(),
        express=ExpressSettings(),
        postgres=PostgresSettings(),
        bot=BotSettings(
            migration_id="migration-bot-dynamic",
            operator_huids=list(operator_huids),
        ),
    )


def build_service(
    *,
    summaries=None,
    checkpoints=None,
    dialogs=None,
    target_validation_errors=None,
    user_huid_by_email=None,
    telegram_gateway=None,
    telegram_session_service=None,
    telegram_session_context=None,
    express_gateway=None,
    backfill_use_case=None,
    operator_huids=("operator-1",),
):
    settings = build_settings(operator_huids=operator_huids)
    chat_mapping_repository = StubChatMappingRepository()
    express_gateway = express_gateway or FakeExpressGateway(
        user_huid_by_email=user_huid_by_email,
    )
    inventory_use_case = StubUseCase()
    backfill_use_case = backfill_use_case or StubUseCase()
    delta_sync_use_case = StubUseCase()
    replay_failed_use_case = StubUseCase(result=None)
    reconcile_use_case = StubUseCase(result=None)
    pause_delta_use_case = StubUseCase(result=None)
    resume_delta_use_case = StubUseCase(result=None)
    identity_mapping_repository = StubIdentityMappingRepository()
    telegram_gateway = telegram_gateway or StubTelegramGateway(dialogs=dialogs)
    identity_directory = StubIdentityDirectory(
        identity_mapping_repository=identity_mapping_repository,
        user_huid_by_email=user_huid_by_email,
    )
    outbox_repository = InMemoryOutboxRepository()
    target_chat_provisioning_service = StubTargetChatProvisioningService(
        validation_errors=target_validation_errors,
        telegram_gateway=telegram_gateway,
        identity_directory=identity_directory,
    )
    telegram_export_snapshot_repository = InMemoryTelegramExportSnapshotRepository()
    telegram_export_archive_stage_store = InMemoryTelegramExportArchiveStageStore()
    telegram_session_context = telegram_session_context or TelegramSessionContext()
    telegram_session_service = telegram_session_service or StubTelegramSessionService()
    service = MigrationBotControlService(
        settings=settings,
        telegram_gateway=telegram_gateway,
        express_gateway=express_gateway,
        inventory_use_case=inventory_use_case,
        backfill_chat_use_case=backfill_use_case,
        delta_sync_use_case=delta_sync_use_case,
        replay_failed_use_case=replay_failed_use_case,
        reconcile_migration_use_case=reconcile_use_case,
        pause_delta_use_case=pause_delta_use_case,
        resume_delta_use_case=resume_delta_use_case,
        chat_mapping_repository=chat_mapping_repository,
        chat_migration_config_repository=StubChatMigrationConfigRepository(),
        checkpoint_repository=StubCheckpointRepository(checkpoints=checkpoints),
        message_mapping_repository=StubMessageMappingRepository(summaries=summaries),
        migration_state_repository=StubMigrationStateRepository(),
        identity_mapping_repository=identity_mapping_repository,
        telegram_export_snapshot_repository=telegram_export_snapshot_repository,
        telegram_export_archive_stage_store=telegram_export_archive_stage_store,
        operator_telegram_session_repository=StubOperatorTelegramSessionRepository(),
        telegram_session_context=telegram_session_context,
        telegram_session_service=telegram_session_service,
        telegram_export_archive_parser=TelegramExportArchiveParser(),
        identity_directory=identity_directory,
        target_chat_provisioning_service=target_chat_provisioning_service,
        audit_repository=StubAuditRepository(),
        migration_job_repository=InMemoryMigrationJobRepository(
            outbox_repository=outbox_repository,
        ),
        logger=StubLogger(),
    )
    return service, backfill_use_case


async def drain_queued_jobs(service: MigrationBotControlService) -> None:
    worker_id = "test-worker"
    while True:
        job = await service._migration_job_repository.acquire_next(
            worker_id=worker_id,
            lease_duration_seconds=30.0,
        )
        if job is None:
            return
        await service.execute_migration_job(job)
        await service._migration_job_repository.mark_completed(
            job_key=job.job_key,
            worker_id=worker_id,
        )


@pytest.mark.asyncio
async def test_start_migrate_chat_blocks_when_progress_policy_ask():
    checkpoint = MigrationCheckpoint(
        migration_id="migration-bot-dynamic",
        source_chat_id="chat-1",
        mode="backfill",
        cursor=HistoryCursor(last_source_message_id="100"),
        updated_at=datetime(2026, 3, 18, 10, 0, tzinfo=UTC),
    )
    summaries = [
        MessageStatusSummary(
            migration_id="migration-bot-dynamic",
            source_chat_id="chat-1",
            import_status=MessageImportStatus.IMPORTED,
            count=5,
        ),
    ]
    service, backfill_use_case = build_service(
        summaries=summaries,
        checkpoints={("migration-bot-dynamic", "chat-1", "backfill"): checkpoint},
        telegram_session_service=StubTelegramSessionService(
            resolved_session_string="session-string-1",
        ),
    )

    result = await service.start_migrate_chat(
        operator=BotOperatorContext(huid="operator-1", chat_id="operator-chat"),
        source_chat_id="chat-1",
        options=MigrationRunOptions(progress_policy="ask"),
    )

    assert result.status == "blocked"
    assert result.progress_hints[0].mapped_total == 5
    assert result.progress_hints[0].last_source_message_id == "100"
    assert not backfill_use_case.commands


@pytest.mark.asyncio
async def test_start_migrate_chat_blocks_when_checkpoint_exists_without_message_rows():
    checkpoint = MigrationCheckpoint(
        migration_id="migration-bot-dynamic",
        source_chat_id="chat-1",
        mode="backfill",
        cursor=HistoryCursor(last_source_message_id="100"),
        updated_at=datetime(2026, 3, 18, 10, 0, tzinfo=UTC),
    )
    service, backfill_use_case = build_service(
        summaries=[],
        checkpoints={("migration-bot-dynamic", "chat-1", "backfill"): checkpoint},
        telegram_session_service=StubTelegramSessionService(
            resolved_session_string="session-string-1",
        ),
    )

    result = await service.start_migrate_chat(
        operator=BotOperatorContext(huid="operator-1", chat_id="operator-chat"),
        source_chat_id="chat-1",
        options=MigrationRunOptions(progress_policy="resume"),
    )

    assert result.status == "blocked"
    assert "inconsistent migration progress detected" in (result.reason or "")
    assert not backfill_use_case.commands


@pytest.mark.asyncio
async def test_list_available_chats_uses_operator_telegram_session_context(tmp_path):
    session_context = TelegramSessionContext()
    observed_session_paths: list[str | None] = []

    class ContextAwareTelegramGateway(StubTelegramGateway):
        async def list_available_dialogs(
            self,
            *,
            limit: int = 100,
            query: str | None = None,
            source_backend: str = "telethon_user_session",
        ):
            observed_session_paths.append(session_context.current_session_path())
            return await super().list_available_dialogs(
                limit=limit,
                query=query,
                source_backend=source_backend,
            )

    service, _ = build_service(
        telegram_gateway=ContextAwareTelegramGateway(),
        telegram_session_service=StubTelegramSessionService(
            resolved_session_string="session-string-1",
        ),
        telegram_session_context=session_context,
    )

    await service.list_available_chats(
        operator=BotOperatorContext(huid="operator-1", chat_id="operator-chat"),
        limit=10,
    )

    assert observed_session_paths == ["session-string-1"]


@pytest.mark.asyncio
async def test_list_available_chats_marks_only_operator_owned_configs_and_progress():
    summaries = [
        MessageStatusSummary(
            migration_id="migration-bot-dynamic",
            source_chat_id="chat-2",
            import_status=MessageImportStatus.IMPORTED,
            count=5,
        ),
    ]
    service, _ = build_service(
        dialogs=[
            SourceDialog(
                dialog_id="chat-1",
                chat_type="supergroup",
                title="Chat One",
                message_count=10,
                media_count=2,
                approximate_bytes=1000,
            ),
            SourceDialog(
                dialog_id="chat-2",
                chat_type="supergroup",
                title="Chat Two",
                message_count=20,
                media_count=3,
                approximate_bytes=2000,
            ),
        ],
        summaries=summaries,
        telegram_session_service=StubTelegramSessionService(
            resolved_session_string="session-string-1",
        ),
        operator_huids=("operator-1", "operator-2"),
    )
    await service._chat_migration_config_repository.save(
        ChatMigrationConfigRecord(
            migration_id="migration-bot-dynamic",
            source_chat_id="chat-2",
            source_chat_type="supergroup",
            source_chat_title="Imported Chat",
            source_backend="telethon_user_session",
            target_strategy="create",
            target_title="Imported Chat",
            target_chat_id=None,
            include_from=None,
            include_to=None,
            migrate_media=True,
            reply_mode="inline_quote",
            identity_policy="matrix_uploaded",
            access_strategy="direct_add",
            updated_by_huid="operator-2",
            created_at=datetime(2026, 3, 27, 12, 0, tzinfo=UTC),
            updated_at=datetime(2026, 3, 27, 12, 0, tzinfo=UTC),
        ),
    )

    chats = await service.list_available_chats(
        operator=BotOperatorContext(huid="operator-1", chat_id="operator-chat"),
        limit=10,
    )

    by_id = {chat.source_chat_id: chat for chat in chats}
    assert by_id["chat-1"].configured is False
    assert by_id["chat-1"].has_progress is False
    assert by_id["chat-2"].configured is False
    assert by_id["chat-2"].has_progress is False


@pytest.mark.asyncio
async def test_start_migrate_chat_accepts_and_runs_background_job(tmp_path):
    service, backfill_use_case = build_service(
        telegram_session_service=StubTelegramSessionService(
            resolved_session_string="session-string-1",
        ),
    )

    result = await service.start_migrate_chat(
        operator=BotOperatorContext(huid="operator-1", chat_id="operator-chat"),
        source_chat_id="chat-1",
        options=MigrationRunOptions(
            reply_mode="source_id",
            batch_size=7,
            progress_policy="resume",
        ),
    )

    assert result.status == "accepted"
    await drain_queued_jobs(service)
    await asyncio.wait_for(backfill_use_case.called.wait(), timeout=1)
    assert backfill_use_case.commands
    command = backfill_use_case.commands[0]
    assert command.source_chat_id == "chat-1"
    assert command.batch_size == 7
    assert command.manifest.dialogs[0].reply_mode == "source_id"


@pytest.mark.asyncio
async def test_start_migrate_chat_persists_anchor_fields_and_worker_uses_them():
    class RoutingAwareBackfillUseCase(StubUseCase):
        def __post_init__(self) -> None:
            super().__post_init__()
            self.observed_hosts: list[str | None] = []

        async def execute(self, command):
            self.commands.append(command)
            self.observed_hosts.append(get_current_express_cts_host())
            self.called.set()
            return None

    backfill_use_case = RoutingAwareBackfillUseCase()
    service, _ = build_service(
        telegram_session_service=StubTelegramSessionService(
            resolved_session_string="session-string-1",
        ),
        backfill_use_case=backfill_use_case,
    )

    result = await service.start_migrate_chat(
        operator=BotOperatorContext(
            huid="operator-1",
            chat_id="operator-chat",
            current_cts_host="cts2.example",
            current_bot_id="bot-cts2",
        ),
        source_chat_id="chat-1",
        options=MigrationRunOptions(progress_policy="resume"),
    )

    assert result.status == "accepted"
    config = await service._chat_migration_config_repository.get(
        "migration-bot-dynamic",
        "chat-1",
    )
    assert config is not None
    assert config.anchor_cts_host == "cts2.example"
    assert config.anchor_bot_id == "bot-cts2"

    queued_job = await service._migration_job_repository.get(result.job_key)
    assert queued_job is not None
    assert queued_job.anchor_cts_host == "cts2.example"
    assert queued_job.anchor_bot_id == "bot-cts2"

    claimed_job = await service._migration_job_repository.acquire_next(
        worker_id="test-worker",
        lease_duration_seconds=30.0,
    )
    assert claimed_job is not None
    await service.execute_migration_job(claimed_job)
    await asyncio.wait_for(backfill_use_case.called.wait(), timeout=1)
    assert backfill_use_case.commands
    assert backfill_use_case.observed_hosts == ["cts2.example"]


@pytest.mark.asyncio
async def test_start_migrate_chat_prefers_primary_anchor_when_gateway_is_routing_aware():
    class RoutingAwareBackfillUseCase(StubUseCase):
        def __post_init__(self) -> None:
            super().__post_init__()
            self.observed_hosts: list[str | None] = []

        async def execute(self, command):
            self.commands.append(command)
            self.observed_hosts.append(get_current_express_cts_host())
            self.called.set()
            return None

    class RoutingAwareExpressGateway(FakeExpressGateway):
        primary_cts_host = "cts1.example"

        def bot_id_for_cts_host(self, cts_host: str | None) -> str | None:
            if cts_host == "cts1.example":
                return "bot-cts1"
            if cts_host == "cts2.example":
                return "bot-cts2"
            return None

    backfill_use_case = RoutingAwareBackfillUseCase()
    service, _ = build_service(
        telegram_session_service=StubTelegramSessionService(
            resolved_session_string="session-string-1",
        ),
        backfill_use_case=backfill_use_case,
        express_gateway=RoutingAwareExpressGateway(),
    )

    result = await service.start_migrate_chat(
        operator=BotOperatorContext(
            huid="operator-1",
            chat_id="operator-chat",
            current_cts_host="cts2.example",
            current_bot_id="bot-cts2",
        ),
        source_chat_id="chat-1",
        options=MigrationRunOptions(progress_policy="resume"),
    )

    assert result.status == "accepted"
    config = await service._chat_migration_config_repository.get(
        "migration-bot-dynamic",
        "chat-1",
    )
    assert config is not None
    assert config.anchor_cts_host == "cts1.example"
    assert config.anchor_bot_id == "bot-cts1"

    queued_job = await service._migration_job_repository.get(result.job_key)
    assert queued_job is not None
    assert queued_job.anchor_cts_host == "cts1.example"
    assert queued_job.anchor_bot_id == "bot-cts1"

    claimed_job = await service._migration_job_repository.acquire_next(
        worker_id="test-worker",
        lease_duration_seconds=30.0,
    )
    assert claimed_job is not None
    await service.execute_migration_job(claimed_job)
    await asyncio.wait_for(backfill_use_case.called.wait(), timeout=1)
    assert backfill_use_case.commands
    assert backfill_use_case.observed_hosts == ["cts1.example"]


@pytest.mark.asyncio
async def test_start_migrate_chat_accepts_uuid_operator_bot_id() -> None:
    service, _ = build_service(
        telegram_session_service=StubTelegramSessionService(
            resolved_session_string="session-string-1",
        ),
    )

    result = await service.start_migrate_chat(
        operator=BotOperatorContext(
            huid="operator-1",
            chat_id="operator-chat",
            current_cts_host="cts2.example",
            current_bot_id=UUID("043a8472-0ec8-5f35-a5a4-3f3ef3ae4aa9"),  # type: ignore[arg-type]
        ),
        source_chat_id="chat-1",
        options=MigrationRunOptions(progress_policy="resume"),
    )

    assert result.status == "accepted"
    config = await service._chat_migration_config_repository.get(
        "migration-bot-dynamic",
        "chat-1",
    )
    assert config is not None
    assert config.anchor_bot_id == "043a8472-0ec8-5f35-a5a4-3f3ef3ae4aa9"
    assert get_current_express_cts_host() is None


@pytest.mark.asyncio
async def test_start_migrate_chat_appends_shadow_outbox_event():
    service, _ = build_service(
        telegram_session_service=StubTelegramSessionService(
            resolved_session_string="session-string-1",
        ),
    )

    result = await service.start_migrate_chat(
        operator=BotOperatorContext(huid="operator-1", chat_id="operator-chat"),
        source_chat_id="chat-1",
        options=MigrationRunOptions(progress_policy="resume"),
    )

    assert result.status == "accepted"
    repository = service._migration_job_repository
    outbox_repository = repository._outbox_repository
    leased = await outbox_repository.lease_batch(
        publisher_id="publisher-1",
        limit=10,
        lease_duration_seconds=30.0,
    )

    assert len(leased) == 1
    event = leased[0]
    assert event.status is IntegrationOutboxStatus.LEASED
    assert event.aggregate_type == "migration_job"
    assert event.aggregate_id == result.job_key
    assert event.event_type == "migration.job.requested"
    assert event.payload_json["job_key"] == result.job_key


@pytest.mark.asyncio
async def test_start_migrate_all_bootstraps_missing_chats_when_some_configs_exist():
    service, _ = build_service(
        dialogs=[
            SourceDialog(
                dialog_id="chat-1",
                chat_type="supergroup",
                title="Configured Chat",
                message_count=10,
                media_count=2,
                approximate_bytes=1000,
            ),
            SourceDialog(
                dialog_id="chat-2",
                chat_type="supergroup",
                title="New Chat",
                message_count=5,
                media_count=1,
                approximate_bytes=500,
            ),
        ],
        telegram_session_service=StubTelegramSessionService(
            resolved_session_string="session-string-1",
        ),
    )
    await service._chat_migration_config_repository.save(
        ChatMigrationConfigRecord(
            migration_id="migration-bot-dynamic",
            source_chat_id="chat-1",
            source_chat_type="supergroup",
            source_chat_title="Configured Chat",
            target_strategy="create",
            target_title="Configured Chat",
            target_chat_id=None,
            source_backend="telethon_user_session",
            include_from=None,
            include_to=None,
            migrate_media=True,
            reply_mode="inline_quote",
            identity_policy="matrix_optional",
            access_strategy="direct_add",
            updated_by_huid="operator-1",
            created_at=datetime(2026, 3, 23, 18, 0, tzinfo=UTC),
            updated_at=datetime(2026, 3, 23, 18, 0, tzinfo=UTC),
            topic_strategy="single_chat",
            skip_in_all=False,
            telegram_chat_id=None,
            source_topic_id=None,
            source_thread_id=None,
            source_thread_title=None,
        ),
    )

    result = await service.start_migrate_all(
        operator=BotOperatorContext(huid="operator-1", chat_id="operator-chat"),
        options=MigrationRunOptions(progress_policy="resume"),
    )

    assert result.status == "accepted"
    assert set(result.source_chat_ids) == {"chat-1", "chat-2"}
    assert set(result.accepted_source_chat_ids) == {"chat-1", "chat-2"}
    assert len(result.job_keys) == 2
    configs = await service._chat_migration_config_repository.list_by_migration("migration-bot-dynamic")
    config_by_chat_id = {config.source_chat_id: config for config in configs}
    assert set(config_by_chat_id) == {"chat-1", "chat-2"}
    assert config_by_chat_id["chat-2"].access_strategy == "invite_link"
    leased = await service._migration_job_repository._outbox_repository.lease_batch(
        publisher_id="publisher-1",
        limit=10,
        lease_duration_seconds=30.0,
    )
    assert len(leased) == 2
    assert {event.payload_json["job_key"] for event in leased} == set(result.job_keys)


@pytest.mark.asyncio
async def test_execute_migration_job_dispatches_delta_sync_operation():
    service, _ = build_service(
        telegram_session_service=StubTelegramSessionService(
            resolved_session_string="session-string-1",
        ),
    )
    await service._chat_migration_config_repository.save(
        ChatMigrationConfigRecord(
            migration_id="migration-bot-dynamic",
            source_chat_id="chat-1",
            source_chat_type="supergroup",
            source_chat_title="Configured Chat",
            target_strategy="create",
            target_title="Configured Chat",
            target_chat_id=None,
            source_backend="telethon_user_session",
            include_from=None,
            include_to=None,
            migrate_media=True,
            reply_mode="inline_quote",
            identity_policy="matrix_optional",
            access_strategy="invite_link",
            updated_by_huid="operator-1",
            created_at=datetime(2026, 3, 24, 12, 0, tzinfo=UTC),
            updated_at=datetime(2026, 3, 24, 12, 0, tzinfo=UTC),
            topic_strategy="single_chat",
            skip_in_all=False,
            telegram_chat_id=None,
            source_topic_id=None,
            source_thread_id=None,
            source_thread_title=None,
        ),
    )
    job = MigrationJobRecord(
        job_key="job-delta-chat",
        migration_id="migration-bot-dynamic",
        operation="delta_sync",
        operator_huid="operator-1",
        source_chat_ids=("chat-1",),
        batch_size=33,
        status=MigrationJobStatus.RUNNING,
        requested_at=datetime(2026, 3, 24, 12, 0, tzinfo=UTC),
        started_at=datetime(2026, 3, 24, 12, 0, tzinfo=UTC),
    )

    await service.execute_migration_job(job)

    assert service._delta_sync_use_case.commands
    command = service._delta_sync_use_case.commands[0]
    assert command.manifest.migration_id == "migration-bot-dynamic"
    assert command.source_chat_ids == ["chat-1"]
    assert command.batch_size == 33


@pytest.mark.asyncio
async def test_start_migrate_chat_allows_parallel_job_for_different_chat_scope():
    service, backfill_use_case = build_service(
        telegram_session_service=StubTelegramSessionService(
            resolved_session_string="session-string-1",
        ),
    )
    await service._migration_job_repository.enqueue(
        MigrationJobRecord(
            job_key="job-other-chat",
            migration_id="migration-bot-dynamic",
            operation="migrate_chat",
            operator_huid="operator-1",
            source_chat_ids=("chat-2",),
            batch_size=100,
            status=MigrationJobStatus.RUNNING,
            requested_at=datetime(2026, 3, 23, 18, 0, tzinfo=UTC),
            started_at=datetime(2026, 3, 23, 18, 0, tzinfo=UTC),
            lease_expires_at=datetime(2099, 3, 23, 18, 1, tzinfo=UTC),
        ),
    )

    result = await service.start_migrate_chat(
        operator=BotOperatorContext(huid="operator-1", chat_id="operator-chat"),
        source_chat_id="chat-1",
        options=MigrationRunOptions(progress_policy="resume"),
    )

    assert result.status == "accepted"
    await drain_queued_jobs(service)
    await asyncio.wait_for(backfill_use_case.called.wait(), timeout=1)


@pytest.mark.asyncio
async def test_start_migrate_chat_ignores_stale_running_job_lease():
    service, backfill_use_case = build_service(
        telegram_session_service=StubTelegramSessionService(
            resolved_session_string="session-string-1",
        ),
    )
    await service._migration_job_repository.enqueue(
        MigrationJobRecord(
            job_key="job-stale-chat",
            migration_id="migration-bot-dynamic",
            operation="migrate_chat",
            operator_huid="operator-1",
            source_chat_ids=("chat-1",),
            batch_size=100,
            status=MigrationJobStatus.RUNNING,
            requested_at=datetime(2000, 3, 23, 18, 0, tzinfo=UTC),
            started_at=datetime(2000, 3, 23, 18, 0, tzinfo=UTC),
            lease_expires_at=datetime(2000, 3, 23, 18, 1, tzinfo=UTC),
        ),
    )

    result = await service.start_migrate_chat(
        operator=BotOperatorContext(huid="operator-1", chat_id="operator-chat"),
        source_chat_id="chat-1",
        options=MigrationRunOptions(progress_policy="resume"),
    )

    assert result.status == "accepted"
    await drain_queued_jobs(service)
    await asyncio.wait_for(backfill_use_case.called.wait(), timeout=1)


@pytest.mark.asyncio
async def test_start_migrate_chat_blocks_when_runtime_backpressure_is_active():
    service, backfill_use_case = build_service(
        telegram_session_service=StubTelegramSessionService(
            resolved_session_string="session-string-1",
        ),
    )
    service._settings.worker.backpressure_max_queued_jobs = 1
    await service._migration_job_repository.enqueue(
        MigrationJobRecord(
            job_key="job-queued-pressure",
            migration_id="migration-bot-dynamic",
            operation="migrate_chat",
            operator_huid="operator-1",
            source_chat_ids=("chat-2",),
            batch_size=100,
            status=MigrationJobStatus.QUEUED,
            requested_at=datetime(2026, 3, 23, 18, 0, tzinfo=UTC),
        ),
    )

    result = await service.start_migrate_chat(
        operator=BotOperatorContext(huid="operator-1", chat_id="operator-chat"),
        source_chat_id="chat-1",
        options=MigrationRunOptions(progress_policy="resume"),
    )

    assert result.status == "blocked"
    assert "runtime backpressure is active" in (result.reason or "")
    assert result.blocked_source_chat_ids == ("chat-1",)
    assert not backfill_use_case.commands


@pytest.mark.asyncio
async def test_list_available_chats_marks_configured_and_progress():
    checkpoint = MigrationCheckpoint(
        migration_id="migration-bot-dynamic",
        source_chat_id="chat-1",
        mode="backfill",
        cursor=HistoryCursor(last_source_message_id="55"),
        updated_at=datetime(2026, 3, 18, 10, 0, tzinfo=UTC),
    )
    summaries = [
        MessageStatusSummary(
            migration_id="migration-bot-dynamic",
            source_chat_id="chat-1",
            import_status=MessageImportStatus.IMPORTED,
            count=3,
        ),
    ]
    service, _ = build_service(
        summaries=summaries,
        checkpoints={("migration-bot-dynamic", "chat-1", "backfill"): checkpoint},
        telegram_session_service=StubTelegramSessionService(
            resolved_session_string="session-string-1",
        ),
    )

    config_result = await service.configure_chat(
        operator=BotOperatorContext(huid="operator-1", chat_id="operator-chat"),
        source_chat_id="chat-1",
        options=MigrationRunOptions(reply_mode="source_id"),
    )
    chats = await service.list_available_chats(
        operator=BotOperatorContext(huid="operator-1", chat_id="operator-chat"),
        limit=20,
        query=None,
    )

    assert config_result.reply_mode == "source_id"
    assert chats[0].configured is True
    assert chats[0].has_progress is True
    assert chats[0].imported_count == 3
    assert chats[0].last_source_message_id == "55"


@pytest.mark.asyncio
async def test_list_chat_users_returns_text_matrix_with_resolved_identities(tmp_path):
    telegram_gateway = StubTelegramGateway(
        dialogs=[
            SourceDialog(
                dialog_id="chat-1",
                chat_type="supergroup",
                title="Imported Chat",
                message_count=10,
                media_count=2,
                approximate_bytes=1000,
            ),
        ],
    )
    telegram_gateway._participants["chat-1"] = [
        SourceParticipant(
            external_id="100",
            username="peer.user",
            display_name="Peer User",
            is_self=False,
        ),
    ]
    service, _ = build_service(
        telegram_gateway=telegram_gateway,
        user_huid_by_email={"peer.user@example.com": "huid-peer-user"},
        telegram_session_service=StubTelegramSessionService(
            resolved_session_string="session-string-1",
        ),
    )
    await service.map_identity(
        operator=BotOperatorContext(huid="operator-1", chat_id="operator-chat"),
        telegram_user_id="100",
        telegram_username="@peer.user",
        telegram_display_name="Peer User",
        corporate_email="peer.user@example.com",
    )

    result = await service.list_chat_users(
        operator=BotOperatorContext(huid="operator-1", chat_id="operator-chat"),
        source_chat_id="chat-1",
    )

    assert result.source_chat_id == "chat-1"
    assert result.entries[0].telegram_user_id == "100"
    assert result.entries[0].telegram_username == "peer.user"
    assert result.entries[0].corporate_email == "peer.user@example.com"
    assert result.entries[0].target_huid == "huid-peer-user"


@pytest.mark.asyncio
async def test_start_migrate_chat_configures_new_private_chat_from_telegram_discovery():
    service, backfill_use_case = build_service(
        dialogs=[
            SourceDialog(
                dialog_id="chat-1",
                chat_type="supergroup",
                title="Imported Chat",
                message_count=10,
                media_count=2,
                approximate_bytes=1000,
            ),
            SourceDialog(
                dialog_id="1057191621",
                chat_type="private",
                title="Direct Chat",
                message_count=5,
                media_count=0,
                approximate_bytes=250,
            ),
        ],
        telegram_session_service=StubTelegramSessionService(
            resolved_session_string="session-string-1",
        ),
    )

    result = await service.start_migrate_chat(
        operator=BotOperatorContext(huid="operator-1", chat_id="operator-chat"),
        source_chat_id="1057191621",
        options=MigrationRunOptions(progress_policy="resume"),
    )

    assert result.status == "accepted"
    await drain_queued_jobs(service)
    await asyncio.wait_for(backfill_use_case.called.wait(), timeout=1)
    command = backfill_use_case.commands[0]
    assert command.manifest.migration_id == "migration-bot-dynamic"
    assert command.manifest.dialogs[0].source_chat_id == "1057191621"


@pytest.mark.asyncio
async def test_show_chat_returns_db_config():
    service, _ = build_service(
        dialogs=[
            SourceDialog(
                dialog_id="chat-1",
                chat_type="supergroup",
                title="Imported Chat",
                message_count=10,
                media_count=2,
                approximate_bytes=1000,
            ),
            SourceDialog(
                dialog_id="1057191621",
                chat_type="private",
                title="Direct Chat",
                message_count=5,
                media_count=0,
                approximate_bytes=250,
            ),
        ],
        telegram_session_service=StubTelegramSessionService(
            resolved_session_string="session-string-1",
        ),
    )

    await service.configure_chat(
        operator=BotOperatorContext(huid="operator-1", chat_id="operator-chat"),
        source_chat_id="1057191621",
        options=MigrationRunOptions(reply_mode="source_id"),
    )
    result = await service.show_chat(
        operator=BotOperatorContext(huid="operator-1", chat_id="operator-chat"),
        source_chat_id="1057191621",
    )

    assert result.configured_via == "db"
    assert result.source_chat_id == "1057191621"
    assert result.source_backend == "telethon_user_session"
    assert result.reply_mode == "source_id"


@pytest.mark.asyncio
async def test_show_or_configure_chat_bootstraps_missing_config():
    service, _ = build_service(
        dialogs=[
            SourceDialog(
                dialog_id="1057191621",
                chat_type="private",
                title="Direct Chat",
                message_count=5,
                media_count=0,
                approximate_bytes=250,
            ),
        ],
        telegram_session_service=StubTelegramSessionService(
            resolved_session_string="session-string-1",
        ),
    )

    result = await service.show_or_configure_chat(
        operator=BotOperatorContext(huid="operator-1", chat_id="operator-chat"),
        source_chat_id="1057191621",
    )

    assert result.source_chat_id == "1057191621"
    assert result.source_backend == "telethon_user_session"
    assert result.target_strategy == "create"
    assert result.target_title == "Direct Chat"


@pytest.mark.asyncio
async def test_configure_chat_rejects_removed_bot_api_live_backend():
    service, _ = build_service(
        dialogs=[
            SourceDialog(
                dialog_id="chat-1",
                chat_type="supergroup",
                title="Imported Chat",
                message_count=10,
                media_count=2,
                approximate_bytes=1000,
            ),
        ],
    )

    with pytest.raises(ConfigurationError, match="Unsupported source backend"):
        await service.configure_chat(
            operator=BotOperatorContext(huid="operator-1", chat_id="operator-chat"),
            source_chat_id="chat-1",
            options=MigrationRunOptions(source_backend="telegram_bot_api_live"),
        )


@pytest.mark.asyncio
async def test_start_archive_import_stages_snapshot_and_enqueues_job():
    service, _ = build_service()

    archive_payload = {
        "id": 5186712067,
        "name": "Archive Chat",
        "type": "private_group",
        "messages": [
            {
                "id": 1,
                "type": "message",
                "date": "2026-03-17T08:00:00",
                "date_unixtime": "1773734400",
                "from": "Alice",
                "from_id": "user1",
                "text": "hello from archive",
            },
        ],
    }

    result = await service.start_archive_import(
        operator=BotOperatorContext(huid="operator-1", chat_id="operator-chat"),
        archive_content=json.dumps(archive_payload).encode("utf-8"),
        archive_filename="result.json",
    )

    assert result.status == "accepted"
    assert result.operation == "import_archive_chat"
    assert result.source_chat_ids == ("archive:5186712067",)

    config = await service._chat_migration_config_repository.get(
        service._settings.bot.migration_id,
        "archive:5186712067",
    )
    assert config is not None
    assert config.source_backend == "telegram_export_archive"
    assert config.skip_in_all is True
    assert config.migrate_media is False
    assert config.identity_policy == "display_only"

    snapshot = await service._telegram_export_snapshot_repository.get("archive:5186712067")
    assert snapshot is not None
    assert snapshot.source_chat_title == "Archive Chat"
    assert snapshot.message_count == 1
    assert snapshot.archive_locator.startswith("in-memory-archive-")

    assert any(
        event.event_type == "telegram_export_archive_staged"
        and event.payload_json.get("source_chat_id") == "archive:5186712067"
        for event in service._audit_repository.events
    )


@pytest.mark.asyncio
async def test_start_migrate_chat_allows_partial_private_chat_without_identity_mapping():
    service, backfill_use_case = build_service(
        dialogs=[
            SourceDialog(
                dialog_id="1057191621",
                chat_type="private",
                title="Direct Chat",
                message_count=5,
                media_count=0,
                approximate_bytes=250,
            ),
        ],
        telegram_session_service=StubTelegramSessionService(
            resolved_session_string="session-string-1",
        ),
    )

    result = await service.start_migrate_chat(
        operator=BotOperatorContext(huid="operator-1", chat_id="operator-chat"),
        source_chat_id="1057191621",
        options=MigrationRunOptions(progress_policy="resume"),
    )

    assert result.status == "accepted"
    await drain_queued_jobs(service)
    await asyncio.wait_for(backfill_use_case.called.wait(), timeout=1)


@pytest.mark.asyncio
async def test_cancel_active_jobs_requests_running_and_cancels_queued_for_operator():
    service, _ = build_service(
        telegram_session_service=StubTelegramSessionService(
            resolved_session_string="session-string-1",
        ),
    )
    await service._migration_job_repository.enqueue(
        MigrationJobRecord(
            job_key="job-queued-self",
            migration_id="migration-bot-dynamic",
            operation="migrate_chat",
            operator_huid="operator-1",
            source_chat_ids=("chat-1",),
            batch_size=100,
            status=MigrationJobStatus.QUEUED,
            requested_at=datetime(2026, 3, 23, 18, 0, tzinfo=UTC),
        ),
    )
    await service._migration_job_repository.enqueue(
        MigrationJobRecord(
            job_key="job-running-self",
            migration_id="migration-bot-dynamic",
            operation="migrate_chat",
            operator_huid="operator-1",
            source_chat_ids=("chat-2",),
            batch_size=100,
            status=MigrationJobStatus.RUNNING,
            requested_at=datetime(2026, 3, 23, 18, 1, tzinfo=UTC),
            started_at=datetime(2026, 3, 23, 18, 1, tzinfo=UTC),
            lease_expires_at=datetime(2099, 3, 23, 18, 2, tzinfo=UTC),
            worker_id="worker-1",
        ),
    )
    await service._migration_job_repository.enqueue(
        MigrationJobRecord(
            job_key="job-running-other",
            migration_id="migration-bot-dynamic",
            operation="migrate_chat",
            operator_huid="operator-2",
            source_chat_ids=("chat-3",),
            batch_size=100,
            status=MigrationJobStatus.RUNNING,
            requested_at=datetime(2026, 3, 23, 18, 2, tzinfo=UTC),
            started_at=datetime(2026, 3, 23, 18, 2, tzinfo=UTC),
            lease_expires_at=datetime(2099, 3, 23, 18, 3, tzinfo=UTC),
            worker_id="worker-2",
        ),
    )

    result = await service.cancel_active_jobs(
        operator=BotOperatorContext(huid="operator-1", chat_id="operator-chat"),
    )

    assert {job.job_key for job in result.cancelled_queued_jobs} == {"job-queued-self"}
    assert {job.job_key for job in result.cancellation_requested_jobs} == {"job-running-self"}
    queued_record = await service._migration_job_repository.get("job-queued-self")
    running_record = await service._migration_job_repository.get("job-running-self")
    other_record = await service._migration_job_repository.get("job-running-other")
    assert queued_record is not None
    assert queued_record.status is MigrationJobStatus.FAILED
    assert queued_record.last_error_code == "cancelled_by_operator"
    assert running_record is not None
    assert running_record.status is MigrationJobStatus.RUNNING
    assert running_record.last_error_code == "cancel_requested_by_operator"
    assert other_record is not None
    assert other_record.last_error_code is None


@pytest.mark.asyncio
async def test_cancel_active_jobs_filters_by_source_chat_id():
    service, _ = build_service(
        telegram_session_service=StubTelegramSessionService(
            resolved_session_string="session-string-1",
        ),
    )
    await service._migration_job_repository.enqueue(
        MigrationJobRecord(
            job_key="job-queued-chat-1",
            migration_id="migration-bot-dynamic",
            operation="migrate_chat",
            operator_huid="operator-1",
            source_chat_ids=("chat-1",),
            batch_size=100,
            status=MigrationJobStatus.QUEUED,
            requested_at=datetime(2026, 3, 23, 18, 0, tzinfo=UTC),
        ),
    )
    await service._migration_job_repository.enqueue(
        MigrationJobRecord(
            job_key="job-running-chat-2",
            migration_id="migration-bot-dynamic",
            operation="migrate_chat",
            operator_huid="operator-1",
            source_chat_ids=("chat-2",),
            batch_size=100,
            status=MigrationJobStatus.RUNNING,
            requested_at=datetime(2026, 3, 23, 18, 1, tzinfo=UTC),
            started_at=datetime(2026, 3, 23, 18, 1, tzinfo=UTC),
            lease_expires_at=datetime(2099, 3, 23, 18, 2, tzinfo=UTC),
            worker_id="worker-1",
        ),
    )
    await service._migration_job_repository.enqueue(
        MigrationJobRecord(
            job_key="job-running-chat-1",
            migration_id="migration-bot-dynamic",
            operation="migrate_chat",
            operator_huid="operator-1",
            source_chat_ids=("chat-1",),
            batch_size=100,
            status=MigrationJobStatus.RUNNING,
            requested_at=datetime(2026, 3, 23, 18, 2, tzinfo=UTC),
            started_at=datetime(2026, 3, 23, 18, 2, tzinfo=UTC),
            lease_expires_at=datetime(2099, 3, 23, 18, 3, tzinfo=UTC),
            worker_id="worker-2",
        ),
    )

    result = await service.cancel_active_jobs(
        operator=BotOperatorContext(huid="operator-1", chat_id="operator-chat"),
        source_chat_id="chat-1",
    )

    assert result.source_chat_id == "chat-1"
    assert {job.job_key for job in result.cancelled_queued_jobs} == {"job-queued-chat-1"}
    assert {job.job_key for job in result.cancellation_requested_jobs} == {"job-running-chat-1"}
    queued_record = await service._migration_job_repository.get("job-queued-chat-1")
    running_chat_1 = await service._migration_job_repository.get("job-running-chat-1")
    running_chat_2 = await service._migration_job_repository.get("job-running-chat-2")
    assert queued_record is not None
    assert queued_record.status is MigrationJobStatus.FAILED
    assert queued_record.last_error_code == "cancelled_by_operator"
    assert running_chat_1 is not None
    assert running_chat_1.last_error_code == "cancel_requested_by_operator"
    assert running_chat_2 is not None
    assert running_chat_2.last_error_code is None


@pytest.mark.asyncio
async def test_start_migrate_chat_returns_existing_private_chat_without_starting_duplicate_job():
    service, backfill_use_case = build_service(
        dialogs=[
            SourceDialog(
                dialog_id="1057191621",
                chat_type="private",
                title="Direct Chat",
                message_count=1,
                media_count=0,
                approximate_bytes=250,
            ),
        ],
        summaries=[
            MessageStatusSummary(
                migration_id="migration-bot-dynamic",
                source_chat_id="1057191621",
                import_status=MessageImportStatus.IMPORTED,
                count=1,
            ),
        ],
        checkpoints={
            (
                "migration-bot-dynamic",
                "1057191621",
                "backfill",
            ): MigrationCheckpoint(
                migration_id="migration-bot-dynamic",
                source_chat_id="1057191621",
                mode="backfill",
                cursor=HistoryCursor(last_source_message_id="1"),
                updated_at=datetime(2026, 3, 19, 10, 0, tzinfo=UTC),
            ),
        },
        telegram_session_service=StubTelegramSessionService(
            resolved_session_string="session-string-1",
        ),
    )
    await service._chat_mapping_repository.save(
        ChatMappingRecord(
            migration_id="migration-bot-dynamic",
            source_chat_id="1057191621",
            source_chat_type="private",
            source_chat_title="Direct Chat",
            target_chat_id="express-chat-1",
            target_chat_title="Imported Direct Chat",
            status="completed",
            created_at=datetime(2026, 3, 19, 10, 0, tzinfo=UTC),
            updated_at=datetime(2026, 3, 19, 10, 0, tzinfo=UTC),
        ),
    )
    await service._express_gateway.create_chat(
        "Imported Direct Chat",
        participant_huids=["operator-1"],
    )

    result = await service.start_migrate_chat(
        operator=BotOperatorContext(huid="operator-1", chat_id="operator-chat"),
        source_chat_id="1057191621",
        options=MigrationRunOptions(progress_policy="resume"),
    )

    assert result.status == "already_migrated"
    assert result.target_chat_id == "express-chat-1"
    assert result.target_chat_link == "https://express.example/chat/express-chat-1"
    assert not backfill_use_case.commands


@pytest.mark.asyncio
async def test_start_migrate_chat_uses_mapping_anchor_for_existing_chat_link():
    class HostAwareExpressGateway(FakeExpressGateway):
        async def create_chat_link(self, target_chat_id: str) -> str | None:
            return f"https://{get_current_express_cts_host() or 'primary'}/chat/{target_chat_id}"

    service, _ = build_service(
        dialogs=[
            SourceDialog(
                dialog_id="group-1",
                chat_type="supergroup",
                title="Engineering Group",
                message_count=10,
                media_count=0,
                approximate_bytes=100,
            ),
        ],
        express_gateway=HostAwareExpressGateway(),
        telegram_session_service=StubTelegramSessionService(
            resolved_session_string="session-string-1",
        ),
    )
    await service._chat_mapping_repository.save(
        ChatMappingRecord(
            migration_id="migration-bot-dynamic",
            source_chat_id="group-1",
            source_chat_type="supergroup",
            source_chat_title="Engineering Group",
            target_chat_id="express-group-1",
            target_chat_title="Imported Engineering Group",
            status="completed",
            created_at=datetime(2026, 3, 23, 10, 0, tzinfo=UTC),
            updated_at=datetime(2026, 3, 23, 10, 0, tzinfo=UTC),
            anchor_cts_host="cts2.example",
        ),
    )

    result = await service.start_migrate_chat(
        operator=BotOperatorContext(
            huid="operator-1",
            chat_id="operator-chat",
            current_cts_host="cts1.example",
            current_bot_id="bot-cts1",
        ),
        source_chat_id="group-1",
        options=MigrationRunOptions(progress_policy="resume"),
    )

    assert result.status == "already_migrated"
    assert result.target_chat_link == "https://cts2.example/chat/express-group-1"


@pytest.mark.asyncio
async def test_remigrate_chat_runs_for_completed_private_chat():
    service, backfill_use_case = build_service(
        dialogs=[
            SourceDialog(
                dialog_id="1057191621",
                chat_type="private",
                title="Direct Chat",
                message_count=1,
                media_count=0,
                approximate_bytes=250,
            ),
        ],
        telegram_session_service=StubTelegramSessionService(
            resolved_session_string="session-string-1",
        ),
    )
    await service._chat_mapping_repository.save(
        ChatMappingRecord(
            migration_id="migration-bot-dynamic",
            source_chat_id="1057191621",
            source_chat_type="private",
            source_chat_title="Direct Chat",
            target_chat_id="express-chat-1",
            target_chat_title="Imported Direct Chat",
            status="completed",
            created_at=datetime(2026, 3, 19, 10, 0, tzinfo=UTC),
            updated_at=datetime(2026, 3, 19, 10, 0, tzinfo=UTC),
        ),
    )

    result = await service.remigrate_chat(
        operator=BotOperatorContext(huid="operator-1", chat_id="operator-chat"),
        source_chat_id="1057191621",
        options=MigrationRunOptions(progress_policy="resume"),
    )

    assert result.status == "accepted"
    await drain_queued_jobs(service)
    await asyncio.wait_for(backfill_use_case.called.wait(), timeout=1)
    assert backfill_use_case.commands


@pytest.mark.asyncio
async def test_remigrate_chat_allows_repeat_run_after_previous_job_finished():
    service, backfill_use_case = build_service(
        dialogs=[
            SourceDialog(
                dialog_id="1057191621",
                chat_type="private",
                title="Direct Chat",
                message_count=1,
                media_count=0,
                approximate_bytes=250,
            ),
        ],
        telegram_session_service=StubTelegramSessionService(
            resolved_session_string="session-string-1",
        ),
    )
    await service._chat_mapping_repository.save(
        ChatMappingRecord(
            migration_id="migration-bot-dynamic",
            source_chat_id="1057191621",
            source_chat_type="private",
            source_chat_title="Direct Chat",
            target_chat_id="express-chat-1",
            target_chat_title="Imported Direct Chat",
            status="completed",
            created_at=datetime(2026, 3, 19, 10, 0, tzinfo=UTC),
            updated_at=datetime(2026, 3, 19, 10, 0, tzinfo=UTC),
        ),
    )
    await service._migration_job_repository.enqueue(
        MigrationJobRecord(
            job_key="legacy-remigrate-job",
            migration_id="migration-bot-dynamic",
            operation="remigrate_chat",
            operator_huid="operator-1",
            source_chat_ids=("1057191621",),
            batch_size=100,
            status=MigrationJobStatus.COMPLETED,
            requested_at=datetime(2026, 3, 19, 10, 0, tzinfo=UTC),
            started_at=datetime(2026, 3, 19, 10, 0, tzinfo=UTC),
            finished_at=datetime(2026, 3, 19, 10, 5, tzinfo=UTC),
        ),
    )

    result = await service.remigrate_chat(
        operator=BotOperatorContext(huid="operator-1", chat_id="operator-chat"),
        source_chat_id="1057191621",
        options=MigrationRunOptions(progress_policy="resume"),
    )

    assert result.status == "accepted"
    await drain_queued_jobs(service)
    await asyncio.wait_for(backfill_use_case.called.wait(), timeout=1)
    assert backfill_use_case.commands


@pytest.mark.asyncio
async def test_remigrate_chat_requires_existing_completed_chat():
    service, backfill_use_case = build_service(
        dialogs=[
            SourceDialog(
                dialog_id="1057191621",
                chat_type="private",
                title="Direct Chat",
                message_count=1,
                media_count=0,
                approximate_bytes=250,
            ),
        ],
        telegram_session_service=StubTelegramSessionService(
            resolved_session_string="session-string-1",
        ),
    )

    with pytest.raises(ConfigurationError, match="use /migrate first"):
        await service.remigrate_chat(
            operator=BotOperatorContext(huid="operator-1", chat_id="operator-chat"),
            source_chat_id="1057191621",
            options=MigrationRunOptions(progress_policy="resume"),
        )

    assert not backfill_use_case.commands


@pytest.mark.asyncio
async def test_start_migrate_chat_does_not_short_circuit_private_chat_when_status_is_not_completed():
    service, backfill_use_case = build_service(
        dialogs=[
            SourceDialog(
                dialog_id="1057191621",
                chat_type="private",
                title="Direct Chat",
                message_count=1,
                media_count=0,
                approximate_bytes=250,
            ),
        ],
        telegram_session_service=StubTelegramSessionService(
            resolved_session_string="session-string-1",
        ),
    )
    await service._chat_mapping_repository.save(
        ChatMappingRecord(
            migration_id="migration-bot-dynamic",
            source_chat_id="1057191621",
            source_chat_type="private",
            source_chat_title="Direct Chat",
            target_chat_id="express-chat-1",
            target_chat_title="Imported Direct Chat",
            status="active",
            created_at=datetime(2026, 3, 19, 10, 0, tzinfo=UTC),
            updated_at=datetime(2026, 3, 19, 10, 0, tzinfo=UTC),
        ),
    )

    result = await service.start_migrate_chat(
        operator=BotOperatorContext(huid="operator-1", chat_id="operator-chat"),
        source_chat_id="1057191621",
        options=MigrationRunOptions(progress_policy="resume"),
    )

    assert result.status == "accepted"
    await drain_queued_jobs(service)
    await asyncio.wait_for(backfill_use_case.called.wait(), timeout=1)


@pytest.mark.asyncio
async def test_prepare_private_chat_resolution_returns_unresolved_peer_for_private_chat():
    telegram_gateway = StubTelegramGateway(
        dialogs=[
            SourceDialog(
                dialog_id="1057191621",
                chat_type="private",
                title="Direct Chat",
                message_count=5,
                media_count=0,
                approximate_bytes=250,
            ),
        ],
    )
    telegram_gateway._participants["1057191621"] = [
        SourceParticipant(
            external_id="self-1",
            username="self.user",
            display_name="Self User",
            is_self=True,
        ),
        SourceParticipant(
            external_id="peer-1",
            username="peer.user",
            display_name="Peer User",
        ),
    ]
    service, _ = build_service(
        telegram_gateway=telegram_gateway,
        telegram_session_service=StubTelegramSessionService(
            resolved_session_string="session-string-1",
        ),
    )

    result = await service.prepare_private_chat_resolution(
        operator=BotOperatorContext(huid="operator-1", chat_id="operator-chat"),
        source_chat_id="1057191621",
        options=MigrationRunOptions(progress_policy="resume"),
    )

    assert result is not None
    assert result.peer_telegram_user_id == "peer-1"
    assert result.peer_telegram_username == "peer.user"
    assert result.peer_display_name == "Peer User"


@pytest.mark.asyncio
async def test_map_private_chat_peer_identity_supports_direct_huid() -> None:
    telegram_gateway = StubTelegramGateway(
        dialogs=[
            SourceDialog(
                dialog_id="1057191621",
                chat_type="private",
                title="Direct Chat",
                message_count=5,
                media_count=0,
                approximate_bytes=250,
            ),
        ],
    )
    telegram_gateway._participants["1057191621"] = [
        SourceParticipant(
            external_id="self-1",
            username="self.user",
            display_name="Self User",
            is_self=True,
        ),
        SourceParticipant(
            external_id="peer-1",
            username="peer.user",
            display_name="Peer User",
        ),
    ]
    service, _ = build_service(
        telegram_gateway=telegram_gateway,
        telegram_session_service=StubTelegramSessionService(
            resolved_session_string="session-string-1",
        ),
    )

    result = await service.map_private_chat_peer_identity(
        operator=BotOperatorContext(huid="operator-1", chat_id="operator-chat"),
        source_chat_id="1057191621",
        target_huid="huid-direct",
    )

    assert result.telegram_user_id == "peer-1"
    assert result.telegram_username == "peer.user"
    assert result.corporate_email is None
    assert result.target_huid == "huid-direct"


@pytest.mark.asyncio
async def test_map_identity_persists_mapping_and_resolves_target_huid():
    service, _ = build_service(
        user_huid_by_email={"peer.user@example.com": "huid-peer-user"},
        telegram_session_service=StubTelegramSessionService(
            resolved_session_string="session-string-1",
        ),
    )

    result = await service.map_identity(
        operator=BotOperatorContext(huid="operator-1", chat_id="operator-chat"),
        telegram_user_id="123456",
        telegram_username="@peer.user",
        telegram_display_name="Peer User",
        corporate_email="Peer.User@example.com",
    )

    saved = await service._identity_mapping_repository.get_by_username("peer.user")

    assert result.telegram_user_id == "123456"
    assert result.telegram_username == "peer.user"
    assert result.telegram_display_name == "Peer User"
    assert result.corporate_email == "peer.user@example.com"
    assert result.target_huid == "huid-peer-user"
    assert saved is not None
    assert saved.telegram_user_id == "123456"
    assert saved.telegram_username == "peer.user"
    assert saved.corporate_email == "peer.user@example.com"


@pytest.mark.asyncio
async def test_bot_operations_require_connected_operator_session() -> None:
    service, _ = build_service(
        telegram_session_service=StubTelegramSessionService(
            error=ConfigurationError(
                "telegram session is not connected for this user; use /connect first",
            ),
        ),
    )

    with pytest.raises(ConfigurationError, match="/connect"):
        await service.list_available_chats(
            operator=BotOperatorContext(huid="operator-1", chat_id="operator-chat"),
            limit=10,
        )


@pytest.mark.asyncio
async def test_bot_operations_allow_any_user_when_allowlist_is_empty() -> None:
    service, _ = build_service(
        operator_huids=(),
        telegram_session_service=StubTelegramSessionService(
            resolved_session_string="session-string-1",
        ),
    )

    chats = await service.list_available_chats(
        operator=BotOperatorContext(huid="any-user", chat_id="chat-1"),
        limit=10,
    )

    assert chats


@pytest.mark.asyncio
async def test_start_migrate_chat_returns_existing_completed_group_chat_without_duplicate_job():
    service, backfill_use_case = build_service(
        dialogs=[
            SourceDialog(
                dialog_id="group-1",
                chat_type="supergroup",
                title="Engineering Group",
                message_count=10,
                media_count=0,
                approximate_bytes=100,
            ),
        ],
        telegram_session_service=StubTelegramSessionService(
            resolved_session_string="session-string-1",
        ),
    )
    await service._chat_mapping_repository.save(
        ChatMappingRecord(
            migration_id="migration-bot-dynamic",
            source_chat_id="group-1",
            source_chat_type="supergroup",
            source_chat_title="Engineering Group",
            target_chat_id="express-group-1",
            target_chat_title="Imported Engineering Group",
            status="completed",
            created_at=datetime(2026, 3, 23, 10, 0, tzinfo=UTC),
            updated_at=datetime(2026, 3, 23, 10, 0, tzinfo=UTC),
        ),
    )
    await service._express_gateway.create_chat(
        "Imported Engineering Group",
        participant_huids=["operator-1"],
    )

    result = await service.start_migrate_chat(
        operator=BotOperatorContext(huid="operator-1", chat_id="operator-chat"),
        source_chat_id="group-1",
        options=MigrationRunOptions(progress_policy="resume"),
    )

    assert result.status == "already_migrated"
    assert result.target_chat_id == "express-group-1"
    assert not backfill_use_case.commands


@pytest.mark.asyncio
async def test_start_migrate_chat_rejects_chat_managed_by_another_operator():
    service, backfill_use_case = build_service(
        dialogs=[
            SourceDialog(
                dialog_id="group-1",
                chat_type="supergroup",
                title="Engineering Group",
                message_count=10,
                media_count=0,
                approximate_bytes=100,
            ),
        ],
        telegram_session_service=StubTelegramSessionService(
            resolved_session_string="session-string-1",
        ),
        operator_huids=("operator-1", "operator-2"),
    )
    await service._chat_migration_config_repository.save(
        ChatMigrationConfigRecord(
            migration_id="migration-bot-dynamic",
            source_chat_id="group-1",
            source_chat_type="supergroup",
            source_chat_title="Engineering Group",
            source_backend="telethon_user_session",
            target_strategy="create",
            target_title="Imported Engineering Group",
            target_chat_id=None,
            include_from=None,
            include_to=None,
            migrate_media=True,
            reply_mode="inline_quote",
            identity_policy="matrix_uploaded",
            access_strategy="direct_add",
            updated_by_huid="operator-1",
            created_at=datetime(2026, 3, 27, 12, 0, tzinfo=UTC),
            updated_at=datetime(2026, 3, 27, 12, 0, tzinfo=UTC),
        ),
    )

    with pytest.raises(ConfigurationError, match="managed by another operator"):
        await service.start_migrate_chat(
            operator=BotOperatorContext(huid="operator-2", chat_id="operator-chat"),
            source_chat_id="group-1",
            options=MigrationRunOptions(progress_policy="resume"),
        )

    saved = await service._chat_migration_config_repository.get(
        "migration-bot-dynamic",
        "group-1",
    )
    assert saved is not None
    assert saved.updated_by_huid == "operator-1"
    assert not backfill_use_case.commands


@pytest.mark.asyncio
async def test_prepare_group_chat_resolution_returns_identity_matrix_request_for_unresolved_users():
    telegram_gateway = StubTelegramGateway(
        dialogs=[
            SourceDialog(
                dialog_id="group-1",
                chat_type="supergroup",
                title="Engineering Group",
                message_count=10,
                media_count=0,
                approximate_bytes=100,
            ),
        ],
    )
    telegram_gateway._participants["group-1"] = [
        SourceParticipant(
            external_id="self-1",
            username="self.user",
            display_name="Self User",
            is_self=True,
        ),
        SourceParticipant(
            external_id="peer-1",
            username="peer.user",
            display_name="Peer User",
            is_self=False,
        ),
    ]
    service, _ = build_service(
        telegram_gateway=telegram_gateway,
        telegram_session_service=StubTelegramSessionService(
            resolved_session_string="session-string-1",
        ),
    )

    result = await service.prepare_group_chat_resolution(
        operator=BotOperatorContext(huid="operator-1", chat_id="operator-chat"),
        source_chat_id="group-1",
        options=MigrationRunOptions(progress_policy="resume"),
    )

    assert result is not None
    assert result.action == "upload_identity_matrix"
    assert result.unresolved_count == 1
    assert result.workbook_content
    assert result.workbook_filename == "identity_matrix_group-1.xlsx"


@pytest.mark.asyncio
async def test_prepare_group_chat_resolution_returns_topic_choice_for_forum_supergroup():
    telegram_gateway = StubTelegramGateway(
        dialogs=[
            SourceDialog(
                dialog_id="forum-1",
                chat_type="supergroup",
                title="Forum Group",
                message_count=10,
                media_count=0,
                approximate_bytes=100,
                has_topics=True,
            ),
        ],
    )
    telegram_gateway._topics["forum-1"] = [
        SourceTopic(topic_id="101", title="Topic 1"),
        SourceTopic(topic_id="102", title="Topic 2"),
    ]
    service, _ = build_service(
        telegram_gateway=telegram_gateway,
        telegram_session_service=StubTelegramSessionService(
            resolved_session_string="session-string-1",
        ),
    )

    result = await service.prepare_group_chat_resolution(
        operator=BotOperatorContext(huid="operator-1", chat_id="operator-chat"),
        source_chat_id="forum-1",
        options=MigrationRunOptions(progress_policy="resume"),
    )

    assert result is not None
    assert result.action == "choose_topic_strategy"
    assert result.topic_titles == ("Topic 1", "Topic 2")


@pytest.mark.asyncio
async def test_prepare_channel_chat_resolution_returns_identity_matrix_request_for_unresolved_users():
    telegram_gateway = StubTelegramGateway(
        dialogs=[
            SourceDialog(
                dialog_id="channel-1",
                chat_type="channel",
                title="Announcements",
                message_count=10,
                media_count=0,
                approximate_bytes=100,
            ),
        ],
    )
    telegram_gateway._participants["channel-1"] = [
        SourceParticipant(
            external_id="user-1",
            username="alice",
            display_name="Alice",
        ),
    ]
    telegram_gateway._channel_access_profiles["channel-1"] = SourceChannelAccessProfile(
        dialog_id="channel-1",
        is_admin=True,
        can_list_participants=True,
    )
    service, _ = build_service(
        telegram_gateway=telegram_gateway,
        telegram_session_service=StubTelegramSessionService(
            resolved_session_string="session-string-1",
        ),
    )

    result = await service.prepare_channel_chat_resolution(
        operator=BotOperatorContext(huid="operator-1", chat_id="operator-chat"),
        source_chat_id="channel-1",
        options=MigrationRunOptions(progress_policy="resume"),
    )

    assert result is not None
    assert result.action == "upload_identity_matrix"
    assert result.unresolved_count == 1
    assert result.workbook_content
    assert result.workbook_filename == "channel_identity_matrix_channel-1.xlsx"


@pytest.mark.asyncio
async def test_prepare_channel_chat_resolution_returns_access_strategy_choice_when_users_are_resolved():
    telegram_gateway = StubTelegramGateway(
        dialogs=[
            SourceDialog(
                dialog_id="channel-1",
                chat_type="channel",
                title="Announcements",
                message_count=10,
                media_count=0,
                approximate_bytes=100,
            ),
        ],
    )
    telegram_gateway._participants["channel-1"] = [
        SourceParticipant(
            external_id="user-1",
            username="alice",
            display_name="Alice",
        ),
    ]
    telegram_gateway._channel_access_profiles["channel-1"] = SourceChannelAccessProfile(
        dialog_id="channel-1",
        is_admin=True,
        can_list_participants=True,
    )
    service, _ = build_service(
        telegram_gateway=telegram_gateway,
        user_huid_by_email={"alice@example.com": "alice-huid"},
        telegram_session_service=StubTelegramSessionService(
            resolved_session_string="session-string-1",
        ),
    )
    await service._identity_mapping_repository.save(
        IdentityMappingRecord(
            telegram_user_id="user-1",
            telegram_username="alice",
            telegram_display_name="Alice",
            corporate_email="alice@example.com",
        ),
    )

    result = await service.prepare_channel_chat_resolution(
        operator=BotOperatorContext(huid="operator-1", chat_id="operator-chat"),
        source_chat_id="channel-1",
        options=MigrationRunOptions(progress_policy="resume"),
    )

    assert result is not None
    assert result.action == "choose_access_strategy"
    assert result.source_chat_id == "channel-1"


@pytest.mark.asyncio
async def test_prepare_channel_chat_resolution_blocks_non_admin_operator():
    telegram_gateway = StubTelegramGateway(
        dialogs=[
            SourceDialog(
                dialog_id="channel-1",
                chat_type="channel",
                title="Announcements",
                message_count=10,
                media_count=0,
                approximate_bytes=100,
            ),
        ],
    )
    telegram_gateway._channel_access_profiles["channel-1"] = SourceChannelAccessProfile(
        dialog_id="channel-1",
        is_admin=False,
        can_list_participants=False,
    )
    service, _ = build_service(
        telegram_gateway=telegram_gateway,
        telegram_session_service=StubTelegramSessionService(
            resolved_session_string="session-string-1",
        ),
    )

    result = await service.prepare_channel_chat_resolution(
        operator=BotOperatorContext(huid="operator-1", chat_id="operator-chat"),
        source_chat_id="channel-1",
        options=MigrationRunOptions(progress_policy="resume"),
    )

    assert result is not None
    assert result.action == "choose_access_strategy"
    assert result.can_list_participants is False


@pytest.mark.asyncio
async def test_apply_group_identity_matrix_imports_rows():
    service, _ = build_service(
        user_huid_by_email={"peer.user@example.com": "huid-peer-user"},
        telegram_session_service=StubTelegramSessionService(
            resolved_session_string="session-string-1",
        ),
    )
    workbook = service._identity_matrix_workbook_service.build_workbook(
        rows=[
            IdentityMatrixWorkbookRow(
                telegram_user_id="peer-1",
                telegram_username="peer.user",
                telegram_display_name="Peer User",
                corporate_email="peer.user@example.com",
                target_huid=None,
            ),
        ],
    )

    result = await service.apply_group_identity_matrix(
        operator=BotOperatorContext(huid="operator-1", chat_id="operator-chat"),
        source_chat_id="group-1",
        workbook_content=workbook,
    )

    saved = await service._identity_mapping_repository.get_by_username("peer.user")
    assert result.imported_count == 1
    assert result.skipped_count == 0
    assert saved is not None
    assert saved.telegram_user_id == "peer-1"
    assert saved.corporate_email == "peer.user@example.com"


@pytest.mark.asyncio
async def test_apply_group_identity_matrix_prefers_huid_over_email_ad_login_and_other_id():
    service, _ = build_service(
        user_huid_by_email={"wrong@example.com": "wrong-huid"},
        telegram_session_service=StubTelegramSessionService(
            resolved_session_string="session-string-1",
        ),
    )
    service._identity_directory._user_huid_by_ad_login["alice"] = "ad-huid"
    service._identity_directory._user_huid_by_other_id["hr-123"] = "other-huid"
    workbook = service._identity_matrix_workbook_service.build_workbook(
        rows=[
            IdentityMatrixWorkbookRow(
                telegram_user_id="peer-1",
                telegram_username="peer.user",
                telegram_display_name="Peer User",
                corporate_email="wrong@example.com",
                target_huid="preferred-huid",
                ad_login="alice",
                other_id="hr-123",
            ),
        ],
    )

    result = await service.apply_group_identity_matrix(
        operator=BotOperatorContext(huid="operator-1", chat_id="operator-chat"),
        source_chat_id="group-1",
        workbook_content=workbook,
    )

    saved = await service._identity_mapping_repository.get_by_username("peer.user")
    assert result.imported_count == 1
    assert result.skipped_count == 0
    assert saved is not None
    assert saved.target_huid == "preferred-huid"
    assert saved.corporate_email == "wrong@example.com"


@pytest.mark.asyncio
async def test_list_migrated_chats_for_user_addition_returns_operator_chats_with_mapping():
    service, _ = build_service(
        telegram_session_service=StubTelegramSessionService(
            resolved_session_string="session-string-1",
        ),
    )
    await service._chat_migration_config_repository.save(
        ChatMigrationConfigRecord(
            migration_id="migration-bot-dynamic",
            source_chat_id="group-1",
            source_chat_type="group",
            source_chat_title="Group One",
            source_backend="telethon_user_session",
            target_strategy="create",
            target_title="Imported Group One",
            target_chat_id=None,
            include_from=None,
            include_to=None,
            migrate_media=True,
            reply_mode="inline_quote",
            identity_policy="matrix_uploaded",
            access_strategy="direct_add",
            updated_by_huid="operator-1",
            created_at=datetime(2026, 3, 27, 12, 0, tzinfo=UTC),
            updated_at=datetime(2026, 3, 27, 12, 0, tzinfo=UTC),
        ),
    )
    await service._chat_mapping_repository.save(
        ChatMappingRecord(
            migration_id="migration-bot-dynamic",
            source_chat_id="group-1",
            source_chat_type="group",
            source_chat_title="Group One",
            target_chat_id="3b3fd173-f5fe-4a35-a50b-aac45f7d1abc",
            target_chat_title="Imported Group One",
            status="completed",
            created_at=datetime(2026, 3, 27, 12, 0, tzinfo=UTC),
            updated_at=datetime(2026, 3, 27, 12, 0, tzinfo=UTC),
        ),
    )
    await service._chat_migration_config_repository.save(
        ChatMigrationConfigRecord(
            migration_id="migration-bot-dynamic",
            source_chat_id="group-2",
            source_chat_type="group",
            source_chat_title="Group Two",
            source_backend="telethon_user_session",
            target_strategy="create",
            target_title="Imported Group Two",
            target_chat_id=None,
            include_from=None,
            include_to=None,
            migrate_media=True,
            reply_mode="inline_quote",
            identity_policy="matrix_uploaded",
            access_strategy="direct_add",
            updated_by_huid="operator-2",
            created_at=datetime(2026, 3, 27, 12, 0, tzinfo=UTC),
            updated_at=datetime(2026, 3, 27, 12, 0, tzinfo=UTC),
        ),
    )
    await service._chat_mapping_repository.save(
        ChatMappingRecord(
            migration_id="migration-bot-dynamic",
            source_chat_id="group-2",
            source_chat_type="group",
            source_chat_title="Group Two",
            target_chat_id="7d94fbe2-1f77-457c-90d1-b60a3ab9632f",
            target_chat_title="Imported Group Two",
            status="completed",
            created_at=datetime(2026, 3, 27, 12, 0, tzinfo=UTC),
            updated_at=datetime(2026, 3, 27, 12, 0, tzinfo=UTC),
        ),
    )

    chats = await service.list_migrated_chats_for_user_addition(
        operator=BotOperatorContext(huid="operator-1", chat_id="operator-chat"),
    )

    assert len(chats) == 1
    assert chats[0].source_chat_id == "group-1"
    assert chats[0].target_chat_id == "3b3fd173-f5fe-4a35-a50b-aac45f7d1abc"


@pytest.mark.asyncio
async def test_show_chat_rejects_chat_managed_by_another_operator():
    service, _ = build_service(
        telegram_session_service=StubTelegramSessionService(
            resolved_session_string="session-string-1",
        ),
        operator_huids=("operator-1", "operator-2"),
    )
    await service._chat_migration_config_repository.save(
        ChatMigrationConfigRecord(
            migration_id="migration-bot-dynamic",
            source_chat_id="group-1",
            source_chat_type="group",
            source_chat_title="Group One",
            source_backend="telethon_user_session",
            target_strategy="create",
            target_title="Imported Group One",
            target_chat_id=None,
            include_from=None,
            include_to=None,
            migrate_media=True,
            reply_mode="inline_quote",
            identity_policy="matrix_uploaded",
            access_strategy="direct_add",
            updated_by_huid="operator-1",
            created_at=datetime(2026, 3, 27, 12, 0, tzinfo=UTC),
            updated_at=datetime(2026, 3, 27, 12, 0, tzinfo=UTC),
        ),
    )

    with pytest.raises(ConfigurationError, match="managed by another operator"):
        await service.show_chat(
            operator=BotOperatorContext(huid="operator-2", chat_id="operator-chat"),
            source_chat_id="group-1",
        )


@pytest.mark.asyncio
async def test_start_migrate_all_uses_only_operator_owned_configs():
    service, _ = build_service(
        dialogs=[
            SourceDialog(
                dialog_id="group-1",
                chat_type="group",
                title="Group One",
                message_count=10,
                media_count=1,
                approximate_bytes=100,
            ),
            SourceDialog(
                dialog_id="group-2",
                chat_type="group",
                title="Group Two",
                message_count=20,
                media_count=2,
                approximate_bytes=200,
            ),
        ],
        telegram_session_service=StubTelegramSessionService(
            resolved_session_string="session-string-1",
        ),
        operator_huids=("operator-1", "operator-2"),
    )
    await service._chat_migration_config_repository.save(
        ChatMigrationConfigRecord(
            migration_id="migration-bot-dynamic",
            source_chat_id="group-1",
            source_chat_type="group",
            source_chat_title="Group One",
            source_backend="telethon_user_session",
            target_strategy="create",
            target_title="Imported Group One",
            target_chat_id=None,
            include_from=None,
            include_to=None,
            migrate_media=True,
            reply_mode="inline_quote",
            identity_policy="matrix_uploaded",
            access_strategy="direct_add",
            updated_by_huid="operator-1",
            created_at=datetime(2026, 3, 27, 12, 0, tzinfo=UTC),
            updated_at=datetime(2026, 3, 27, 12, 0, tzinfo=UTC),
        ),
    )
    await service._chat_migration_config_repository.save(
        ChatMigrationConfigRecord(
            migration_id="migration-bot-dynamic",
            source_chat_id="group-2",
            source_chat_type="group",
            source_chat_title="Group Two",
            source_backend="telethon_user_session",
            target_strategy="create",
            target_title="Imported Group Two",
            target_chat_id=None,
            include_from=None,
            include_to=None,
            migrate_media=True,
            reply_mode="inline_quote",
            identity_policy="matrix_uploaded",
            access_strategy="direct_add",
            updated_by_huid="operator-2",
            created_at=datetime(2026, 3, 27, 12, 0, tzinfo=UTC),
            updated_at=datetime(2026, 3, 27, 12, 0, tzinfo=UTC),
        ),
    )

    result = await service.start_migrate_all(
        operator=BotOperatorContext(huid="operator-1", chat_id="operator-chat"),
        options=MigrationRunOptions(progress_policy="resume"),
    )

    assert result.accepted_source_chat_ids == ("group-1",)
    inventory_command = service._inventory_use_case.commands[0]
    assert [dialog.source_chat_id for dialog in inventory_command.manifest.dialogs] == ["group-1"]


@pytest.mark.asyncio
async def test_stats_reports_foreign_managed_chats_separately():
    service, _ = build_service(
        telegram_session_service=StubTelegramSessionService(
            resolved_session_string="session-string-1",
        ),
        operator_huids=("operator-1", "operator-2"),
    )
    await service._chat_migration_config_repository.save(
        ChatMigrationConfigRecord(
            migration_id="migration-bot-dynamic",
            source_chat_id="group-1",
            source_chat_type="group",
            source_chat_title="Group One",
            source_backend="telethon_user_session",
            target_strategy="create",
            target_title="Imported Group One",
            target_chat_id=None,
            include_from=None,
            include_to=None,
            migrate_media=True,
            reply_mode="inline_quote",
            identity_policy="matrix_uploaded",
            access_strategy="direct_add",
            updated_by_huid="operator-1",
            created_at=datetime(2026, 3, 27, 12, 0, tzinfo=UTC),
            updated_at=datetime(2026, 3, 27, 12, 0, tzinfo=UTC),
        ),
    )

    result = await service.stats(
        operator=BotOperatorContext(huid="operator-2", chat_id="operator-chat"),
    )

    assert result.configured_chats == 0
    assert result.foreign_managed_chats == 1
    assert result.active_jobs == 0


@pytest.mark.asyncio
async def test_status_filters_foreign_chats_and_jobs():
    service, _ = build_service(
        telegram_session_service=StubTelegramSessionService(
            resolved_session_string="session-string-1",
        ),
        operator_huids=("operator-1", "operator-2"),
    )
    await service._chat_migration_config_repository.save(
        ChatMigrationConfigRecord(
            migration_id="migration-bot-dynamic",
            source_chat_id="group-1",
            source_chat_type="group",
            source_chat_title="Group One",
            source_backend="telethon_user_session",
            target_strategy="create",
            target_title="Imported Group One",
            target_chat_id=None,
            include_from=None,
            include_to=None,
            migrate_media=True,
            reply_mode="inline_quote",
            identity_policy="matrix_uploaded",
            access_strategy="direct_add",
            updated_by_huid="operator-1",
            created_at=datetime(2026, 3, 27, 12, 0, tzinfo=UTC),
            updated_at=datetime(2026, 3, 27, 12, 0, tzinfo=UTC),
        ),
    )
    await service._chat_migration_config_repository.save(
        ChatMigrationConfigRecord(
            migration_id="migration-bot-dynamic",
            source_chat_id="group-2",
            source_chat_type="group",
            source_chat_title="Group Two",
            source_backend="telethon_user_session",
            target_strategy="create",
            target_title="Imported Group Two",
            target_chat_id=None,
            include_from=None,
            include_to=None,
            migrate_media=True,
            reply_mode="inline_quote",
            identity_policy="matrix_uploaded",
            access_strategy="direct_add",
            updated_by_huid="operator-2",
            created_at=datetime(2026, 3, 27, 12, 0, tzinfo=UTC),
            updated_at=datetime(2026, 3, 27, 12, 0, tzinfo=UTC),
        ),
    )
    await service._migration_job_repository.enqueue(
        MigrationJobRecord(
            job_key="job-1",
            migration_id="migration-bot-dynamic",
            operation="migrate_chat",
            operator_huid="operator-1",
            source_chat_ids=("group-1",),
            batch_size=100,
            status=MigrationJobStatus.QUEUED,
            requested_at=datetime(2026, 3, 27, 12, 0, tzinfo=UTC),
        ),
    )
    await service._migration_job_repository.enqueue(
        MigrationJobRecord(
            job_key="job-2",
            migration_id="migration-bot-dynamic",
            operation="migrate_chat",
            operator_huid="operator-2",
            source_chat_ids=("group-2",),
            batch_size=100,
            status=MigrationJobStatus.QUEUED,
            requested_at=datetime(2026, 3, 27, 12, 1, tzinfo=UTC),
        ),
    )
    service._reconcile_migration_use_case.result = ReconcileMigrationResult(
        migration_id="migration-bot-dynamic",
        chats_total=0,
        attention_chats=0,
        inventory_missing_chats=0,
        source_messages_total_known=0,
        imported_count=0,
        failed_count=0,
        ambiguous_count=0,
        processing_count=0,
        mapped_total=0,
        gap_total_known=0,
        source_media_total_known=0,
        attachment_imported_count=0,
        attachment_failed_count=0,
        attachment_ambiguous_count=0,
        attachment_processing_count=0,
        attachment_skipped_count=0,
        attachment_mapped_total=0,
        chats=[],
    )

    result = await service.status(
        operator=BotOperatorContext(huid="operator-1", chat_id="operator-chat"),
    )

    assert [dialog.source_chat_id for dialog in service._reconcile_migration_use_case.commands[0].manifest.dialogs] == ["group-1"]
    assert [job.job_key for job in result.active_jobs] == ["job-1"]
    assert [checkpoint.source_chat_id for checkpoint in result.chat_checkpoints] == ["group-1"]


@pytest.mark.asyncio
async def test_prepare_chat_members_workbook_includes_manual_template_rows():
    telegram_gateway = StubTelegramGateway(
        dialogs=[
            SourceDialog(
                dialog_id="group-1",
                chat_type="group",
                title="Group One",
                message_count=10,
                media_count=0,
                approximate_bytes=100,
            ),
        ],
    )
    telegram_gateway._participants["group-1"] = [
        SourceParticipant(
            external_id="tg-1",
            username="alice",
            display_name="Alice",
        ),
    ]
    service, _ = build_service(
        telegram_gateway=telegram_gateway,
        telegram_session_service=StubTelegramSessionService(
            resolved_session_string="session-string-1",
        ),
    )
    await service._chat_migration_config_repository.save(
        ChatMigrationConfigRecord(
            migration_id="migration-bot-dynamic",
            source_chat_id="group-1",
            source_chat_type="group",
            source_chat_title="Group One",
            source_backend="telethon_user_session",
            target_strategy="create",
            target_title="Imported Group One",
            target_chat_id=None,
            include_from=None,
            include_to=None,
            migrate_media=True,
            reply_mode="inline_quote",
            identity_policy="matrix_uploaded",
            access_strategy="direct_add",
            updated_by_huid="operator-1",
            created_at=datetime(2026, 3, 27, 12, 0, tzinfo=UTC),
            updated_at=datetime(2026, 3, 27, 12, 0, tzinfo=UTC),
        ),
    )
    await service._chat_mapping_repository.save(
        ChatMappingRecord(
            migration_id="migration-bot-dynamic",
            source_chat_id="group-1",
            source_chat_type="group",
            source_chat_title="Group One",
            target_chat_id="3b3fd173-f5fe-4a35-a50b-aac45f7d1abc",
            target_chat_title="Imported Group One",
            status="completed",
            created_at=datetime(2026, 3, 27, 12, 0, tzinfo=UTC),
            updated_at=datetime(2026, 3, 27, 12, 0, tzinfo=UTC),
        ),
    )

    request = await service.prepare_chat_members_workbook(
        operator=BotOperatorContext(huid="operator-1", chat_id="operator-chat"),
        chat_selector="group-1",
    )
    rows = service._identity_matrix_workbook_service.parse_workbook(request.workbook_content)
    workbook = load_workbook(filename=BytesIO(request.workbook_content), data_only=True)
    sheet = workbook[service._identity_matrix_workbook_service.SHEET_NAME]

    assert request.target_chat_id == "3b3fd173-f5fe-4a35-a50b-aac45f7d1abc"
    assert request.mapped_rows == 0
    assert request.unresolved_rows == 1
    assert request.workbook_filename == "chat_members_group-1.xlsx"
    assert len(rows) == 1
    assert rows[0].telegram_username == "alice"
    assert sheet.max_row >= 12


@pytest.mark.asyncio
async def test_add_chat_members_from_workbook_supports_ad_login_and_manual_rows():
    service, _ = build_service(
        telegram_session_service=StubTelegramSessionService(
            resolved_session_string="session-string-1",
        ),
    )
    service._identity_directory._user_huid_by_ad_login["alice"] = "alice-huid"
    service._identity_directory._user_huid_by_other_id["hr-123"] = "bob-huid"
    service._target_chat_provisioning_service.member_add_result = ParticipantAccessApplyResult(
        added_huids=("alice-huid",),
        invited_huids=("bob-huid",),
    )
    await service._chat_migration_config_repository.save(
        ChatMigrationConfigRecord(
            migration_id="migration-bot-dynamic",
            source_chat_id="group-1",
            source_chat_type="group",
            source_chat_title="Group One",
            source_backend="telethon_user_session",
            target_strategy="create",
            target_title="Imported Group One",
            target_chat_id=None,
            include_from=None,
            include_to=None,
            migrate_media=True,
            reply_mode="inline_quote",
            identity_policy="matrix_uploaded",
            access_strategy="direct_add",
            updated_by_huid="operator-1",
            created_at=datetime(2026, 3, 27, 12, 0, tzinfo=UTC),
            updated_at=datetime(2026, 3, 27, 12, 0, tzinfo=UTC),
            anchor_cts_host="cts-main.example.test",
        ),
    )
    await service._chat_mapping_repository.save(
        ChatMappingRecord(
            migration_id="migration-bot-dynamic",
            source_chat_id="group-1",
            source_chat_type="group",
            source_chat_title="Group One",
            target_chat_id="3b3fd173-f5fe-4a35-a50b-aac45f7d1abc",
            target_chat_title="Imported Group One",
            status="completed",
            created_at=datetime(2026, 3, 27, 12, 0, tzinfo=UTC),
            updated_at=datetime(2026, 3, 27, 12, 0, tzinfo=UTC),
            anchor_cts_host="cts-main.example.test",
        ),
    )
    workbook = service._identity_matrix_workbook_service.build_workbook(
        rows=[
            IdentityMatrixWorkbookRow(
                telegram_user_id="tg-1",
                telegram_username="alice",
                telegram_display_name="Alice",
                corporate_email=None,
                target_huid=None,
                ad_login="alice",
            ),
            IdentityMatrixWorkbookRow(
                telegram_user_id=None,
                telegram_username=None,
                telegram_display_name="Bob Manual",
                corporate_email=None,
                target_huid=None,
                other_id="hr-123",
            ),
        ],
    )

    result = await service.add_chat_members_from_workbook(
        operator=BotOperatorContext(
            huid="operator-1",
            chat_id="operator-chat",
            ad_domain="corp.example",
        ),
        chat_selector="group-1",
        workbook_content=workbook,
    )

    saved = await service._identity_mapping_repository.get_by_username("alice")
    assert saved is not None
    assert saved.target_huid == "alice-huid"
    assert result.processed_rows == 2
    assert result.imported_identity_mappings == 1
    assert result.mapping_skipped_rows == 1


@pytest.mark.asyncio
async def test_add_chat_members_from_workbook_prefers_email_over_ad_login_and_other_id():
    service, _ = build_service(
        telegram_session_service=StubTelegramSessionService(
            resolved_session_string="session-string-1",
        ),
    )
    service._identity_directory._user_huid_by_email["preferred@example.com"] = "preferred-huid"
    service._identity_directory._user_huid_by_ad_login["alice"] = "ad-huid"
    service._identity_directory._user_huid_by_other_id["hr-123"] = "other-huid"
    service._target_chat_provisioning_service.member_add_result = ParticipantAccessApplyResult(
        added_huids=("preferred-huid",),
    )
    await service._chat_migration_config_repository.save(
        ChatMigrationConfigRecord(
            migration_id="migration-bot-dynamic",
            source_chat_id="group-1",
            source_chat_type="group",
            source_chat_title="Group One",
            source_backend="telethon_user_session",
            target_strategy="create",
            target_title="Imported Group One",
            target_chat_id=None,
            include_from=None,
            include_to=None,
            migrate_media=True,
            reply_mode="inline_quote",
            identity_policy="matrix_uploaded",
            access_strategy="direct_add",
            updated_by_huid="operator-1",
            created_at=datetime(2026, 3, 27, 12, 0, tzinfo=UTC),
            updated_at=datetime(2026, 3, 27, 12, 0, tzinfo=UTC),
            anchor_cts_host="cts-main.example.test",
        ),
    )
    await service._chat_mapping_repository.save(
        ChatMappingRecord(
            migration_id="migration-bot-dynamic",
            source_chat_id="group-1",
            source_chat_type="group",
            source_chat_title="Group One",
            target_chat_id="3b3fd173-f5fe-4a35-a50b-aac45f7d1abc",
            target_chat_title="Imported Group One",
            status="completed",
            created_at=datetime(2026, 3, 27, 12, 0, tzinfo=UTC),
            updated_at=datetime(2026, 3, 27, 12, 0, tzinfo=UTC),
            anchor_cts_host="cts-main.example.test",
        ),
    )
    workbook = service._identity_matrix_workbook_service.build_workbook(
        rows=[
            IdentityMatrixWorkbookRow(
                telegram_user_id=None,
                telegram_username=None,
                telegram_display_name="Manual User",
                corporate_email="preferred@example.com",
                target_huid=None,
                ad_login="alice",
                other_id="hr-123",
            ),
        ],
    )

    result = await service.add_chat_members_from_workbook(
        operator=BotOperatorContext(
            huid="operator-1",
            chat_id="operator-chat",
            ad_domain="corp.example",
        ),
        chat_selector="group-1",
        workbook_content=workbook,
    )

    member_call = service._target_chat_provisioning_service.member_add_calls[-1]
    participant_targets = member_call["participant_targets"]
    assert result.processed_rows == 1
    assert result.imported_identity_mappings == 0
    assert result.mapping_skipped_rows == 1
    assert len(participant_targets) == 1
    assert participant_targets[0].target_huid == "preferred-huid"
    assert result.resolved_targets == 1
    assert result.direct_added == 1
    assert result.invited == 0


@pytest.mark.asyncio
async def test_add_chat_members_from_workbook_returns_failed_target_labels_for_partial_direct_add():
    service, _ = build_service(
        telegram_session_service=StubTelegramSessionService(
            resolved_session_string="session-string-1",
        ),
    )
    service._identity_directory._user_huid_by_email["alice@example.com"] = "alice-huid"
    service._identity_directory._user_huid_by_ad_login["bob"] = "bob-huid"
    service._target_chat_provisioning_service.member_add_result = ParticipantAccessApplyResult(
        added_huids=("alice-huid",),
        failed_targets=(
            ParticipantAccessFailure(
                target_huid="bob-huid",
                cts_host="cts-main.example.test",
                reason="direct_add_failed",
                error="Sender is not chat admin",
            ),
        ),
    )
    await service._chat_migration_config_repository.save(
        ChatMigrationConfigRecord(
            migration_id="migration-bot-dynamic",
            source_chat_id="group-1",
            source_chat_type="group",
            source_chat_title="Group One",
            source_backend="telethon_user_session",
            target_strategy="create",
            target_title="Imported Group One",
            target_chat_id=None,
            include_from=None,
            include_to=None,
            migrate_media=True,
            reply_mode="inline_quote",
            identity_policy="matrix_uploaded",
            access_strategy="direct_add",
            updated_by_huid="operator-1",
            created_at=datetime(2026, 3, 27, 12, 0, tzinfo=UTC),
            updated_at=datetime(2026, 3, 27, 12, 0, tzinfo=UTC),
            anchor_cts_host="cts-main.example.test",
        ),
    )
    await service._chat_mapping_repository.save(
        ChatMappingRecord(
            migration_id="migration-bot-dynamic",
            source_chat_id="group-1",
            source_chat_type="group",
            source_chat_title="Group One",
            target_chat_id="3b3fd173-f5fe-4a35-a50b-aac45f7d1abc",
            target_chat_title="Imported Group One",
            status="completed",
            created_at=datetime(2026, 3, 27, 12, 0, tzinfo=UTC),
            updated_at=datetime(2026, 3, 27, 12, 0, tzinfo=UTC),
            anchor_cts_host="cts-main.example.test",
        ),
    )
    workbook = service._identity_matrix_workbook_service.build_workbook(
        rows=[
            IdentityMatrixWorkbookRow(
                telegram_user_id=None,
                telegram_username=None,
                telegram_display_name="Alice",
                corporate_email="alice@example.com",
                target_huid=None,
            ),
            IdentityMatrixWorkbookRow(
                telegram_user_id=None,
                telegram_username=None,
                telegram_display_name="Bob",
                corporate_email=None,
                target_huid=None,
                ad_login="bob",
            ),
        ],
    )

    result = await service.add_chat_members_from_workbook(
        operator=BotOperatorContext(
            huid="operator-1",
            chat_id="operator-chat",
            ad_domain="corp.example",
        ),
        chat_selector="group-1",
        workbook_content=workbook,
    )

    assert result.effective_result == "partial_direct_add"
    assert result.failed_targets == ("ad_login=bob",)


@pytest.mark.asyncio
async def test_start_migrate_chat_split_by_topic_materializes_logical_topic_dialogs():
    telegram_gateway = StubTelegramGateway(
        dialogs=[
            SourceDialog(
                dialog_id="forum-1",
                chat_type="supergroup",
                title="Forum Chat",
                message_count=10,
                media_count=1,
                approximate_bytes=1000,
                has_topics=True,
            ),
        ],
    )
    telegram_gateway._topics["forum-1"] = [
        SourceTopic(topic_id="101", title="Topic 1", top_message_id="10"),
        SourceTopic(topic_id="102", title="Topic 2", top_message_id="20"),
    ]
    service, backfill_use_case = build_service(
        telegram_gateway=telegram_gateway,
        telegram_session_service=StubTelegramSessionService(
            resolved_session_string="session-string-1",
        ),
    )

    result = await service.start_migrate_chat(
        operator=BotOperatorContext(huid="operator-1", chat_id="operator-chat"),
        source_chat_id="forum-1",
        options=MigrationRunOptions(topic_strategy="split_by_topic", progress_policy="resume"),
    )

    assert result.status == "accepted"
    assert result.source_chat_ids == ("forum-1#topic:101", "forum-1#topic:102")
    await drain_queued_jobs(service)
    await asyncio.wait_for(backfill_use_case.called.wait(), timeout=1)
    await asyncio.sleep(0)
    assert len(backfill_use_case.commands) == 2
    manifest = backfill_use_case.commands[0].manifest
    assert {dialog.source_chat_id for dialog in manifest.dialogs} == {
        "forum-1#topic:101",
        "forum-1#topic:102",
    }
    assert {dialog.telegram_chat_id for dialog in manifest.dialogs} == {"forum-1"}
    assert {dialog.source_thread_id for dialog in manifest.dialogs} == {"10", "20"}


@pytest.mark.asyncio
async def test_start_migrate_chat_returns_existing_topic_targets_for_completed_split_migration():
    telegram_gateway = StubTelegramGateway(
        dialogs=[
            SourceDialog(
                dialog_id="forum-1",
                chat_type="supergroup",
                title="Forum Chat",
                message_count=10,
                media_count=1,
                approximate_bytes=1000,
                has_topics=True,
            ),
        ],
    )
    telegram_gateway._topics["forum-1"] = [
        SourceTopic(topic_id="101", title="Topic 1", top_message_id="10"),
        SourceTopic(topic_id="102", title="Topic 2", top_message_id="20"),
    ]
    service, backfill_use_case = build_service(
        telegram_gateway=telegram_gateway,
        telegram_session_service=StubTelegramSessionService(
            resolved_session_string="session-string-1",
        ),
    )

    await service.start_migrate_chat(
        operator=BotOperatorContext(huid="operator-1", chat_id="operator-chat"),
        source_chat_id="forum-1",
        options=MigrationRunOptions(topic_strategy="split_by_topic", progress_policy="resume"),
    )
    await drain_queued_jobs(service)
    await asyncio.wait_for(backfill_use_case.called.wait(), timeout=1)
    await service._chat_mapping_repository.save(
        ChatMappingRecord(
            migration_id="migration-bot-dynamic",
            source_chat_id="forum-1#topic:101",
            source_chat_type="supergroup",
            source_chat_title="Forum Chat / Topic 1",
            target_chat_id="express-chat-1",
            target_chat_title="Forum Chat / Topic 1",
            status="completed",
            created_at=datetime(2026, 3, 23, 12, 0, tzinfo=UTC),
            updated_at=datetime(2026, 3, 23, 12, 0, tzinfo=UTC),
        ),
    )
    await service._chat_mapping_repository.save(
        ChatMappingRecord(
            migration_id="migration-bot-dynamic",
            source_chat_id="forum-1#topic:102",
            source_chat_type="supergroup",
            source_chat_title="Forum Chat / Topic 2",
            target_chat_id="express-chat-2",
            target_chat_title="Forum Chat / Topic 2",
            status="completed",
            created_at=datetime(2026, 3, 23, 12, 0, tzinfo=UTC),
            updated_at=datetime(2026, 3, 23, 12, 0, tzinfo=UTC),
        ),
    )
    await service._express_gateway.create_chat("Forum Chat / Topic 1")
    await service._express_gateway.create_chat("Forum Chat / Topic 2")

    result = await service.start_migrate_chat(
        operator=BotOperatorContext(huid="operator-1", chat_id="operator-chat"),
        source_chat_id="forum-1",
        options=MigrationRunOptions(progress_policy="resume"),
    )

    assert result.status == "already_migrated"
    assert len(result.topic_targets) == 2
    assert {target.target_chat_id for target in result.topic_targets} == {
        "express-chat-1",
        "express-chat-2",
    }


@pytest.mark.asyncio
async def test_prepare_channel_chat_resolution_requires_channel_admin():
    telegram_gateway = StubTelegramGateway(
        dialogs=[
            SourceDialog(
                dialog_id="channel-1",
                chat_type="channel",
                title="Release Notes",
                message_count=10,
                media_count=0,
                approximate_bytes=100,
            ),
        ],
    )
    telegram_gateway._channel_access_profiles["channel-1"] = SourceChannelAccessProfile(
        dialog_id="channel-1",
        is_admin=False,
        can_list_participants=False,
    )
    service, _ = build_service(
        telegram_gateway=telegram_gateway,
        telegram_session_service=StubTelegramSessionService(
            resolved_session_string="session-string-1",
        ),
    )

    result = await service.prepare_channel_chat_resolution(
        operator=BotOperatorContext(huid="operator-1", chat_id="operator-chat"),
        source_chat_id="channel-1",
        options=MigrationRunOptions(progress_policy="resume"),
    )

    assert result is not None
    assert result.action == "choose_access_strategy"
    assert result.can_list_participants is False


@pytest.mark.asyncio
async def test_prepare_channel_chat_resolution_returns_identity_matrix_request_for_unresolved_users():
    telegram_gateway = StubTelegramGateway(
        dialogs=[
            SourceDialog(
                dialog_id="channel-1",
                chat_type="channel",
                title="Release Notes",
                message_count=10,
                media_count=0,
                approximate_bytes=100,
            ),
        ],
    )
    telegram_gateway._channel_access_profiles["channel-1"] = SourceChannelAccessProfile(
        dialog_id="channel-1",
        is_admin=True,
        can_list_participants=True,
    )
    telegram_gateway._participants["channel-1"] = [
        SourceParticipant(
            external_id="self-1",
            username="self.user",
            display_name="Self User",
            is_self=True,
        ),
        SourceParticipant(
            external_id="peer-1",
            username="peer.user",
            display_name="Peer User",
            is_self=False,
        ),
    ]
    service, _ = build_service(
        telegram_gateway=telegram_gateway,
        telegram_session_service=StubTelegramSessionService(
            resolved_session_string="session-string-1",
        ),
    )

    result = await service.prepare_channel_chat_resolution(
        operator=BotOperatorContext(huid="operator-1", chat_id="operator-chat"),
        source_chat_id="channel-1",
        options=MigrationRunOptions(progress_policy="resume"),
    )

    assert result is not None
    assert result.action == "upload_identity_matrix"
    assert result.unresolved_count == 1
    assert result.workbook_content
    assert result.workbook_filename == "channel_identity_matrix_channel-1.xlsx"
    assert result.can_list_participants is True


@pytest.mark.asyncio
async def test_prepare_channel_chat_resolution_requests_access_strategy_after_resolved_users():
    telegram_gateway = StubTelegramGateway(
        dialogs=[
            SourceDialog(
                dialog_id="channel-1",
                chat_type="channel",
                title="Release Notes",
                message_count=10,
                media_count=0,
                approximate_bytes=100,
            ),
        ],
    )
    telegram_gateway._channel_access_profiles["channel-1"] = SourceChannelAccessProfile(
        dialog_id="channel-1",
        is_admin=True,
        can_list_participants=False,
    )
    telegram_gateway._participants["channel-1"] = [
        SourceParticipant(
            external_id="peer-1",
            username="peer.user",
            display_name="Peer User",
            is_self=False,
        ),
    ]
    service, _ = build_service(
        telegram_gateway=telegram_gateway,
        user_huid_by_email={"peer.user@example.com": "huid-peer-user"},
        telegram_session_service=StubTelegramSessionService(
            resolved_session_string="session-string-1",
        ),
    )
    await service.map_identity(
        operator=BotOperatorContext(huid="operator-1", chat_id="operator-chat"),
        telegram_user_id="peer-1",
        telegram_username="@peer.user",
        telegram_display_name="Peer User",
        corporate_email="peer.user@example.com",
    )

    result = await service.prepare_channel_chat_resolution(
        operator=BotOperatorContext(huid="operator-1", chat_id="operator-chat"),
        source_chat_id="channel-1",
        options=MigrationRunOptions(progress_policy="resume"),
    )

    assert result is not None
    assert result.action == "choose_access_strategy"
    assert result.workbook_content is None
    assert result.can_list_participants is False


@pytest.mark.asyncio
async def test_apply_channel_identity_matrix_imports_rows():
    service, _ = build_service(
        user_huid_by_email={"peer.user@example.com": "huid-peer-user"},
        telegram_session_service=StubTelegramSessionService(
            resolved_session_string="session-string-1",
        ),
    )
    workbook = service._identity_matrix_workbook_service.build_workbook(
        rows=[
            IdentityMatrixWorkbookRow(
                telegram_user_id="peer-1",
                telegram_username="peer.user",
                telegram_display_name="Peer User",
                corporate_email="peer.user@example.com",
                target_huid=None,
            ),
        ],
    )

    result = await service.apply_channel_identity_matrix(
        operator=BotOperatorContext(huid="operator-1", chat_id="operator-chat"),
        source_chat_id="channel-1",
        workbook_content=workbook,
    )

    saved = await service._identity_mapping_repository.get_by_username("peer.user")
    assert result.imported_count == 1
    assert result.skipped_count == 0
    assert saved is not None
    assert saved.telegram_user_id == "peer-1"
    assert saved.corporate_email == "peer.user@example.com"


@pytest.mark.asyncio
async def test_list_chat_users_returns_best_effort_note_for_channel_listing_failure():
    telegram_gateway = StubTelegramGateway(
        dialogs=[
            SourceDialog(
                dialog_id="channel-1",
                chat_type="channel",
                title="Release Notes",
                message_count=10,
                media_count=0,
                approximate_bytes=100,
            ),
        ],
    )
    telegram_gateway._participant_errors["channel-1"] = FatalItemError("participants hidden")
    service, _ = build_service(
        telegram_gateway=telegram_gateway,
        telegram_session_service=StubTelegramSessionService(
            resolved_session_string="session-string-1",
        ),
    )

    result = await service.list_chat_users(
        operator=BotOperatorContext(huid="operator-1", chat_id="operator-chat"),
        source_chat_id="channel-1",
    )

    assert result.source_chat_type == "channel"
    assert result.entries == ()
    assert result.note is not None
    assert "best-effort" in result.note


@pytest.mark.asyncio
async def test_start_migrate_chat_returns_existing_completed_channel_without_duplicate_job():
    service, backfill_use_case = build_service(
        dialogs=[
            SourceDialog(
                dialog_id="channel-1",
                chat_type="channel",
                title="Release Notes",
                message_count=10,
                media_count=0,
                approximate_bytes=100,
            ),
        ],
        telegram_session_service=StubTelegramSessionService(
            resolved_session_string="session-string-1",
        ),
    )
    target_chat_id = await service._express_gateway.create_chat("Imported Release Notes")
    await service._chat_mapping_repository.save(
        ChatMappingRecord(
            migration_id="migration-bot-dynamic",
            source_chat_id="channel-1",
            source_chat_type="channel",
            source_chat_title="Release Notes",
            target_chat_id=target_chat_id,
            target_chat_title="Imported Release Notes",
            status="completed",
            created_at=datetime(2026, 3, 23, 13, 0, tzinfo=UTC),
            updated_at=datetime(2026, 3, 23, 13, 0, tzinfo=UTC),
        ),
    )

    result = await service.start_migrate_chat(
        operator=BotOperatorContext(huid="operator-1", chat_id="operator-chat"),
        source_chat_id="channel-1",
        options=MigrationRunOptions(progress_policy="resume"),
    )

    assert result.status == "already_migrated"
    assert result.target_chat_id == target_chat_id
    assert result.target_chat_link == f"https://express.example/chat/{target_chat_id}"
    assert not backfill_use_case.commands
