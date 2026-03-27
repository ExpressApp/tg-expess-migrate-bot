from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO

from openpyxl import Workbook, load_workbook

from extg_shared.contracts.errors import ConfigurationError


@dataclass(frozen=True, slots=True)
class IdentityMatrixWorkbookRow:
    telegram_user_id: str | None
    telegram_username: str | None
    telegram_display_name: str
    corporate_email: str | None
    target_huid: str | None
    ad_login: str | None = None
    other_id: str | None = None
    is_self: bool = False
    resolution_source: str = "display_only"
    reason: str | None = None


class IdentityMatrixWorkbookService:
    SHEET_NAME = "identity_mapping"
    LEGACY_HEADER = (
        "telegram_user_id",
        "telegram_username",
        "telegram_display_name",
        "corporate_email",
        "target_huid",
        "is_self",
        "resolution_source",
        "reason",
    )
    HEADER = (
        "telegram_user_id",
        "telegram_username",
        "telegram_display_name",
        "corporate_email",
        "target_huid",
        "ad_login",
        "other_id",
        "is_self",
        "resolution_source",
        "reason",
    )

    def build_workbook(
        self,
        *,
        rows: list[IdentityMatrixWorkbookRow],
        manual_template_rows: int = 0,
    ) -> bytes:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = self.SHEET_NAME
        sheet.append(self.HEADER)
        for row in rows:
            sheet.append(
                [
                    row.telegram_user_id,
                    row.telegram_username,
                    row.telegram_display_name,
                    row.corporate_email,
                    row.target_huid,
                    row.ad_login,
                    row.other_id,
                    "yes" if row.is_self else "no",
                    row.resolution_source,
                    row.reason,
                ],
            )
        for _ in range(max(manual_template_rows, 0)):
            sheet.append(
                [
                    None,
                    None,
                    "",
                    None,
                    None,
                    None,
                    None,
                    "no",
                    "manual_add_template",
                    "You can fill several selectors; priority: target_huid, corporate_email, ad_login, other_id",
                ],
            )
        buffer = BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()

    def parse_workbook(
        self,
        payload: bytes,
    ) -> list[IdentityMatrixWorkbookRow]:
        try:
            workbook = load_workbook(filename=BytesIO(payload), data_only=True)
        except Exception as error:  # pragma: no cover - openpyxl specific
            raise ConfigurationError("uploaded file is not a valid Excel workbook") from error
        if self.SHEET_NAME not in workbook.sheetnames:
            raise ConfigurationError(
                f"uploaded workbook must contain sheet {self.SHEET_NAME!r}",
            )
        sheet = workbook[self.SHEET_NAME]
        header = tuple(
            str(value).strip() if value is not None else ""
            for value in next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
        )
        header_variant: tuple[str, ...]
        if header[: len(self.HEADER)] == self.HEADER:
            header_variant = self.HEADER
        elif header[: len(self.LEGACY_HEADER)] == self.LEGACY_HEADER:
            header_variant = self.LEGACY_HEADER
        else:
            raise ConfigurationError(
                "uploaded workbook has unexpected header; download a fresh matrix and fill it",
            )
        rows: list[IdentityMatrixWorkbookRow] = []
        for values in sheet.iter_rows(min_row=2, values_only=True):
            telegram_user_id = self._optional_str(values[0] if len(values) > 0 else None)
            telegram_username = self._optional_str(values[1] if len(values) > 1 else None)
            telegram_display_name = self._optional_str(values[2] if len(values) > 2 else None) or ""
            corporate_email = self._optional_str(values[3] if len(values) > 3 else None)
            target_huid = self._optional_str(values[4] if len(values) > 4 else None)
            ad_login = None
            other_id = None
            offset = 0
            if header_variant == self.HEADER:
                ad_login = self._optional_str(values[5] if len(values) > 5 else None)
                other_id = self._optional_str(values[6] if len(values) > 6 else None)
                offset = 2
            is_self = self._parse_yes_no(values[5 + offset] if len(values) > 5 + offset else None)
            resolution_source = (
                self._optional_str(values[6 + offset] if len(values) > 6 + offset else None)
                or "display_only"
            )
            reason = self._optional_str(values[7 + offset] if len(values) > 7 + offset else None)
            if (
                telegram_user_id is None
                and telegram_username is None
                and telegram_display_name == ""
                and corporate_email is None
                and target_huid is None
                and ad_login is None
                and other_id is None
            ):
                continue
            rows.append(
                IdentityMatrixWorkbookRow(
                    telegram_user_id=telegram_user_id,
                    telegram_username=telegram_username,
                    telegram_display_name=telegram_display_name,
                    corporate_email=corporate_email,
                    target_huid=target_huid,
                    ad_login=ad_login,
                    other_id=other_id,
                    is_self=is_self,
                    resolution_source=resolution_source,
                    reason=reason,
                ),
            )
        return rows

    def _optional_str(self, value: object) -> str | None:
        if value is None:
            return None
        rendered = str(value).strip()
        return rendered or None

    def _parse_yes_no(self, value: object) -> bool:
        if value is None:
            return False
        return str(value).strip().lower() in {"yes", "true", "1"}
