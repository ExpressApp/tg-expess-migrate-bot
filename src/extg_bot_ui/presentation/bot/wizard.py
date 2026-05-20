from __future__ import annotations

from datetime import UTC, datetime
from io import BytesIO
from dataclasses import dataclass
from enum import Enum, auto
from uuid import UUID

from openpyxl import Workbook
from pybotx import Bot, BubbleMarkup, IncomingMessage
from pybotx.models.attachments import OutgoingAttachment
from pybotx_fsm import FSMCollector

from extg_bot_ui.application.bot_control import (
    BotAvailableChat,
    BotChatConfigurationResult,
    BotChatMembersAddResult,
    BotChatMembersWorkbookRequest,
    BotChannelChatResolutionRequest,
    BotGroupChatResolutionRequest,
    BotGroupIdentityImportResult,
    BotMigratedChat,
    BotPrivateChatResolutionRequest,
    BotOperationAcceptedResult,
    BotOperatorContext,
    BotOperatorDefaultsResult,
    BotProgressHint,
    MigrationBotControlService,
    MigrationRunOptions,
)
from extg_bot_ui.presentation.bot.command_parser import parse_arguments
from extg_bot_ui.presentation.bot.menus import show_connected_menu, show_main_menu
from extg_bot_ui.presentation.bot.screen import render_screen
from extg_shared.contracts.errors import (
    AmbiguousDeliveryError,
    ConfigurationError,
    FatalItemError,
    RecoverableItemError,
)
from extg_shared.contracts.output_template import (
    OUTPUT_TEMPLATE_PLACEHOLDERS,
    validate_output_template,
)

WIZARD_STATE_REPO_KEY = "extg_fsm_repo"
WIZARD_TTL_SECONDS = 3600
USER_VISIBLE_WIZARD_ERRORS = (
    ConfigurationError,
    FatalItemError,
    RecoverableItemError,
    AmbiguousDeliveryError,
)

FIGMA_FLOW_STYLE = "figma"
FIGMA_FLOW_SINGLE_CHAT = "single_chat"
FIGMA_FLOW_MIGRATE_ALL = "migrate_all"
FIGMA_FLOW_OPERATOR_DEFAULTS = "operator_defaults"
DEFAULT_FIGMA_MEDIA_KINDS = ("photo", "file")
ALL_FIGMA_MEDIA_KINDS = ("photo", "video", "voice", "video_note", "file")
FIGMA_MEDIA_KIND_LABELS = {
    "photo": "Фотографии",
    "video": "Видеозаписи",
    "voice": "Голосовые сообщения",
    "video_note": "Видеосообщения",
    "file": "Файлы",
}
SUPPORTED_REPLY_MODES = ("inline_quote", "source_id", "none")


class MigrationWizardState(Enum):
    SELECT_CHAT = auto()
    INPUT_SINGLE_CHAT_CONFIGURATION = auto()
    INPUT_OPERATOR_DEFAULTS_CONFIGURATION = auto()
    INPUT_MIGRATE_ALL_EXCLUSIONS = auto()
    CONFIRM_MIGRATE_ALL_EXCLUSIONS = auto()
    INPUT_MIGRATE_ALL_CONFIGURATION = auto()
    SELECT_MEMBERS_CHAT = auto()
    INPUT_CHAT_MEMBERS_WORKBOOK = auto()
    INPUT_ARCHIVE_IMPORT = auto()
    INPUT_ARCHIVE_IDENTITY_MATRIX = auto()
    INPUT_ARCHIVE_PARTICIPANT_DECISION = auto()
    INPUT_PRIVATE_PEER = auto()
    INPUT_GROUP_TOPIC_STRATEGY = auto()
    INPUT_GROUP_IDENTITY_MATRIX = auto()
    INPUT_CHANNEL_IDENTITY_MATRIX = auto()
    INPUT_CHANNEL_ACCESS_STRATEGY = auto()
    INPUT_FROM = auto()
    INPUT_TO = auto()
    INPUT_MEDIA = auto()
    INPUT_SERVICE_MESSAGES = auto()
    INPUT_REPLY = auto()
    CONFIRM = auto()
    EDIT_CUSTOM_CONFIGURATION = auto()
    INPUT_PARTICIPANT_MIGRATION_DECISION = auto()


@dataclass(frozen=True, slots=True)
class PrivatePeerResolutionInput:
    corporate_email: str | None = None
    target_huid: str | None = None
    source: str = "unknown"


def build_wizard_collector(
    service: MigrationBotControlService,
    *,
    default_batch_size: int,
) -> FSMCollector:
    fsm = FSMCollector(MigrationWizardState)

    @fsm.on(MigrationWizardState.SELECT_CHAT)
    async def select_chat(message: IncomingMessage, bot: Bot) -> None:
        body = _message_text(message)
        choice = _normalized_choice(body)
        if _is_cancel(choice):
            await _cancel(message, bot)
            return
        if _is_back(choice):
            await message.state.fsm.drop_state()
            await show_connected_menu(bot, message=message)
            return
        if choice == "list":
            chats = await service.list_available_chats(
                operator=_operator_from_message(message),
                limit=1000,
                query=None,
            )
            await _show_chat_migration_selection(message=message, bot=bot, chats=chats)
            return
        if choice == "all":
            await _prompt_migrate_all_exclusions(message=message, bot=bot)
            return
        if not body:
            await bot.answer_message(
                "Нужен source_chat_id. Отправь id чата, нажми `Все чаты`, `Назад` или `Отмена`.",
                wait_callback=False,
            )
            return

        try:
            chats = await service.list_available_chats(
                operator=_operator_from_message(message),
                limit=1000,
                query=None,
            )
        except USER_VISIBLE_WIZARD_ERRORS as error:
            await bot.answer_message(f"Ошибка: {error}", wait_callback=False)
            return
        if not any(chat.source_chat_id == body for chat in chats):
            await bot.answer_message(
                "Такой source_chat_id не найден. Отправь корректный id, `list`, `Все чаты`, `Назад` или `Отмена`.",
                wait_callback=False,
            )
            return
        selected_chat = next(chat for chat in chats if chat.source_chat_id == body)
        await _begin_single_chat_configuration_choice(
            message=message,
            bot=bot,
            service=service,
            source_chat_id=body,
            source_chat_type=selected_chat.source_chat_type,
            source_chat_title=selected_chat.source_chat_title,
        )

    @fsm.on(MigrationWizardState.INPUT_SINGLE_CHAT_CONFIGURATION)
    async def input_single_chat_configuration(message: IncomingMessage, bot: Bot) -> None:
        choice = _normalized_choice(_message_text(message))
        if _is_cancel(choice):
            await _cancel(message, bot)
            return
        if _is_back(choice):
            await _return_to_chat_selection(message=message, bot=bot, service=service)
            return
        if choice == "custom":
            await _prompt_figma_media_selection(message=message, bot=bot)
            return
        if choice != "default":
            await bot.answer_message(
                "Выбери `Оставить по умолчанию`, `Настроить кастомную`, `Назад` или `Отмена`.",
                wait_callback=False,
            )
            return
        await _start_single_chat_figma_flow(
            message=message,
            bot=bot,
            service=service,
            default_batch_size=default_batch_size,
        )

    @fsm.on(MigrationWizardState.INPUT_OPERATOR_DEFAULTS_CONFIGURATION)
    async def input_operator_defaults_configuration(message: IncomingMessage, bot: Bot) -> None:
        choice = _normalized_choice(_message_text(message))
        if _is_cancel(choice) or _is_back(choice) or choice in {"menu", "main_menu"}:
            await message.state.fsm.drop_state()
            await show_main_menu(bot, message=message)
            return
        if choice != "custom":
            await bot.answer_message(
                "Выбери `Настроить кастомную`, `Главное меню` или `Отмена`.",
                wait_callback=False,
            )
            return
        await _prompt_figma_media_selection(
            message=message,
            bot=bot,
            flow_target=FIGMA_FLOW_OPERATOR_DEFAULTS,
        )

    @fsm.on(MigrationWizardState.INPUT_MIGRATE_ALL_EXCLUSIONS)
    async def input_migrate_all_exclusions(message: IncomingMessage, bot: Bot) -> None:
        choice = _normalized_choice(_message_text(message))
        if _is_cancel(choice):
            await _cancel(message, bot)
            return
        if _is_back(choice):
            await _return_to_chat_selection(message=message, bot=bot, service=service)
            return
        if _is_skip(choice):
            await _prompt_migrate_all_configuration(
                message=message,
                bot=bot,
                service=service,
                excluded_source_chat_ids=(),
            )
            return
        try:
            chats = await service.list_available_chats(
                operator=_operator_from_message(message),
                limit=1000,
                query=None,
            )
        except USER_VISIBLE_WIZARD_ERRORS as error:
            await bot.answer_message(f"Ошибка: {error}", wait_callback=False)
            return
        valid_chat_ids = {chat.source_chat_id for chat in chats}
        excluded_source_chat_ids = _parse_source_chat_id_list(_message_text(message))
        if not excluded_source_chat_ids:
            await bot.answer_message(
                "Отправь id чатов через запятую, нажми `Пропустить`, `Назад` или `Отмена`.",
                wait_callback=False,
            )
            return
        unknown_chat_ids = [
            source_chat_id
            for source_chat_id in excluded_source_chat_ids
            if source_chat_id not in valid_chat_ids
        ]
        if unknown_chat_ids:
            await bot.answer_message(
                "Не удалось найти чаты: "
                + ", ".join(unknown_chat_ids)
                + ". Отправь список заново, `Пропустить`, `Назад` или `Отмена`.",
                wait_callback=False,
            )
            return
        await message.state.fsm.change_state(
            MigrationWizardState.CONFIRM_MIGRATE_ALL_EXCLUSIONS,
            ttl_seconds=WIZARD_TTL_SECONDS,
            **_wizard_state_payload(
                message,
                excluded_source_chat_ids=excluded_source_chat_ids,
            ),
        )
        await render_screen(
            message=message,
            bot=bot,
            body=_build_migrate_all_exclusion_confirmation(excluded_source_chat_ids),
            bubbles=_migrate_all_exclusion_confirmation_keyboard(),
        )

    @fsm.on(MigrationWizardState.CONFIRM_MIGRATE_ALL_EXCLUSIONS)
    async def confirm_migrate_all_exclusions(message: IncomingMessage, bot: Bot) -> None:
        choice = _normalized_choice(_message_text(message))
        if _is_cancel(choice):
            await _cancel(message, bot)
            return
        if choice in {"edit", "change", "back"}:
            await _prompt_migrate_all_exclusions(message=message, bot=bot)
            return
        if choice != "confirm":
            await bot.answer_message(
                "Нажми `Все верно`, `Изменить список` или `Отмена`.",
                wait_callback=False,
            )
            return
        excluded_source_chat_ids = tuple(
            getattr(message.state.fsm_storage, "excluded_source_chat_ids", ()) or ()
        )
        await _prompt_migrate_all_configuration(
            message=message,
            bot=bot,
            service=service,
            excluded_source_chat_ids=excluded_source_chat_ids,
        )

    @fsm.on(MigrationWizardState.INPUT_MIGRATE_ALL_CONFIGURATION)
    async def input_migrate_all_configuration(message: IncomingMessage, bot: Bot) -> None:
        choice = _normalized_choice(_message_text(message))
        if _is_cancel(choice):
            await _cancel(message, bot)
            return
        if _is_back(choice):
            excluded_source_chat_ids = tuple(
                getattr(message.state.fsm_storage, "excluded_source_chat_ids", ()) or ()
            )
            if excluded_source_chat_ids:
                await message.state.fsm.change_state(
                    MigrationWizardState.CONFIRM_MIGRATE_ALL_EXCLUSIONS,
                    ttl_seconds=WIZARD_TTL_SECONDS,
                    **_wizard_state_payload(
                        message,
                        excluded_source_chat_ids=excluded_source_chat_ids,
                    ),
                )
                await render_screen(
                    message=message,
                    bot=bot,
                    body=_build_migrate_all_exclusion_confirmation(excluded_source_chat_ids),
                    bubbles=_migrate_all_exclusion_confirmation_keyboard(),
                )
                return
            await _prompt_migrate_all_exclusions(message=message, bot=bot)
            return
        if choice == "custom":
            await _prompt_figma_media_selection(
                message=message,
                bot=bot,
                flow_target=FIGMA_FLOW_MIGRATE_ALL,
            )
            return
        if choice != "default":
            await bot.answer_message(
                "Выбери `Оставить по умолчанию`, `Настроить кастомную`, `Назад` или `Отмена`.",
                wait_callback=False,
            )
            return
        try:
            result = await service.start_migrate_all(
                operator=_operator_from_message(message),
                options=MigrationRunOptions(
                    excluded_source_chat_ids=tuple(
                        getattr(message.state.fsm_storage, "excluded_source_chat_ids", ()) or ()
                    ),
                    progress_policy="resume",
                    batch_size=default_batch_size,
                ),
            )
        except USER_VISIBLE_WIZARD_ERRORS as error:
            await bot.answer_message(f"Ошибка: {error}", wait_callback=False)
            return
        await message.state.fsm.drop_state()
        if result.status == "accepted":
            await show_connected_menu(
                bot,
                message=message,
                notice=(
                    "Миграция истории сообщений чатов запущена.\n"
                    "Статус миграции можно посмотреть в главном меню по одноименной кнопке."
                ),
            )
            return
        await show_connected_menu(
            bot,
            message=message,
            notice=_format_background_operation_for_wizard(result),
        )

    @fsm.on(MigrationWizardState.INPUT_SERVICE_MESSAGES)
    async def input_service_messages(message: IncomingMessage, bot: Bot) -> None:
        choice = _normalized_choice(_message_text(message))
        if _is_cancel(choice):
            await _cancel(message, bot)
            return
        if _is_back(choice):
            if _figma_edit_field(message) == "service_messages":
                await _prompt_figma_edit_menu(message=message, bot=bot)
                return
            await _prompt_figma_media_selection(
                message=message,
                bot=bot,
                flow_target=_figma_flow_target(message),
            )
            return
        if choice not in {"yes", "no"}:
            await bot.answer_message(
                "Нажми `Да`, `Нет`, `Назад` или `Отмена`.",
                wait_callback=False,
            )
            return
        service_messages = choice == "yes"
        if _figma_edit_field(message) == "service_messages":
            await _clear_figma_edit_field_and_prompt_confirm(
                message=message,
                bot=bot,
                service_messages=service_messages,
            )
            return
        await _prompt_figma_from(
            message=message,
            bot=bot,
            service_messages=service_messages,
        )

    @fsm.on(MigrationWizardState.SELECT_MEMBERS_CHAT)
    async def select_members_chat(message: IncomingMessage, bot: Bot) -> None:
        choice = _normalized_choice(_message_text(message))
        if _is_cancel(choice):
            await _cancel(message, bot)
            return
        if _is_back(choice):
            await message.state.fsm.drop_state()
            await show_main_menu(bot, message=message)
            return
        if choice == "list":
            chats = await service.list_migrated_chats_for_user_addition(
                operator=_operator_from_message(message),
            )
            await bot.answer_message(
                _build_members_wizard_intro(chats),
                wait_callback=False,
            )
            return
        try:
            chat_selector = _parse_chat_selection_message(message)
            request = await service.prepare_chat_members_workbook(
                operator=_operator_from_message(message),
                chat_selector=chat_selector,
            )
        except USER_VISIBLE_WIZARD_ERRORS as error:
            await bot.answer_message(f"Ошибка: {error}", wait_callback=False)
            return
        await begin_chat_members_workbook(
            message=message,
            bot=bot,
            request=request,
        )

    @fsm.on(MigrationWizardState.INPUT_CHAT_MEMBERS_WORKBOOK)
    async def input_chat_members_workbook(message: IncomingMessage, bot: Bot) -> None:
        choice = _normalized_choice(_message_text(message))
        if _is_cancel(choice):
            await _cancel(message, bot)
            return
        if _is_back(choice):
            try:
                chats = await service.list_migrated_chats_for_user_addition(
                    operator=_operator_from_message(message),
                )
            except USER_VISIBLE_WIZARD_ERRORS as error:
                await bot.answer_message(f"Ошибка: {error}", wait_callback=False)
                return
            await begin_chat_members_selection(
                message=message,
                bot=bot,
                chats=chats,
            )
            return
        if message.file is None:
            await bot.answer_message(
                "Загрузи заполненный Excel-файл или отправь `/cancel`.",
                wait_callback=False,
            )
            return
        try:
            result = await service.add_chat_members_from_workbook(
                operator=_operator_from_message(message),
                chat_selector=message.state.fsm_storage.source_chat_id,
                workbook_content=message.file.content,
            )
        except USER_VISIBLE_WIZARD_ERRORS as error:
            await bot.answer_message(f"Ошибка: {error}", wait_callback=False)
            return
        await render_screen(
            message=message,
            bot=bot,
            body=_format_chat_members_add_result(result),
            bubbles=_chat_members_result_keyboard(),
        )
        await message.state.fsm.drop_state()

    @fsm.on(MigrationWizardState.INPUT_ARCHIVE_IMPORT)
    async def input_archive_import(message: IncomingMessage, bot: Bot) -> None:
        choice = _normalized_choice(_message_text(message))
        if _is_cancel(choice):
            await _cancel(message, bot)
            return
        if _is_back(choice):
            await message.state.fsm.drop_state()
            await show_main_menu(bot, message=message)
            return
        if message.file is None:
            await bot.answer_message(
                "Загрузите Telegram export `.json`, `.zip` или `.rar`, либо нажмите `Назад` или отправьте `/cancel`.",
                wait_callback=False,
            )
            return
        await render_screen(
            message=message,
            bot=bot,
            body="Подождите, файл обрабатывается...",
        )
        try:
            result = await service.prepare_archive_import(
                operator=_operator_from_message(message),
                archive_content=message.file.content,
                archive_filename=getattr(message.file, "filename", None),
            )
        except USER_VISIBLE_WIZARD_ERRORS as error:
            await bot.answer_message(f"Ошибка: {error}", wait_callback=False)
            return
        await message.state.fsm.change_state(
            MigrationWizardState.INPUT_ARCHIVE_IDENTITY_MATRIX,
            ttl_seconds=WIZARD_TTL_SECONDS,
            **_wizard_state_payload(
                message,
                source_chat_id=result.source_chat_id,
                source_chat_type=result.source_chat_type,
                source_chat_title=result.source_chat_title,
            ),
        )
        await render_screen(
            message=message,
            bot=bot,
            body=_build_archive_identity_matrix_prompt(),
            file=OutgoingAttachment(
                content=result.workbook_content,
                filename=result.workbook_filename,
            ),
            bubbles=_archive_identity_matrix_keyboard(),
        )

    @fsm.on(MigrationWizardState.INPUT_ARCHIVE_IDENTITY_MATRIX)
    async def input_archive_identity_matrix(message: IncomingMessage, bot: Bot) -> None:
        choice = _normalized_choice(_message_text(message))
        if _is_cancel(choice):
            await _cancel(message, bot)
            return
        if _is_back(choice):
            await begin_archive_import(message=message, bot=bot)
            return
        if _is_skip(choice):
            try:
                result = await service.start_prepared_archive_import(
                    operator=_operator_from_message(message),
                    source_chat_id=message.state.fsm_storage.source_chat_id,
                    migrate_members=False,
                    identity_matrix_uploaded=False,
                )
            except USER_VISIBLE_WIZARD_ERRORS as error:
                await bot.answer_message(f"Ошибка: {error}", wait_callback=False)
                return
            notice = _build_archive_import_start_notice(
                source_chat_type=getattr(message.state.fsm_storage, "source_chat_type", "group"),
                source_chat_title=getattr(message.state.fsm_storage, "source_chat_title", "Telegram Export"),
                result=result,
            )
            await message.state.fsm.drop_state()
            await show_main_menu(
                bot,
                message=message,
                notice=notice,
            )
            return
        if message.file is None:
            await bot.answer_message(
                "Загрузите заполненный `identity_matrix.xlsx`, либо нажмите `Пропустить`, `Назад` или `Отмена`.",
                wait_callback=False,
            )
            return
        await render_screen(
            message=message,
            bot=bot,
            body="Подождите, файл обрабатывается...",
        )
        try:
            await service.apply_archive_identity_matrix(
                operator=_operator_from_message(message),
                source_chat_id=message.state.fsm_storage.source_chat_id,
                workbook_content=message.file.content,
            )
        except USER_VISIBLE_WIZARD_ERRORS as error:
            await bot.answer_message(f"Ошибка: {error}", wait_callback=False)
            return
        await message.state.fsm.change_state(
            MigrationWizardState.INPUT_ARCHIVE_PARTICIPANT_DECISION,
            ttl_seconds=WIZARD_TTL_SECONDS,
            **_wizard_state_payload(
                message,
                source_chat_id=message.state.fsm_storage.source_chat_id,
                source_chat_type=getattr(message.state.fsm_storage, "source_chat_type", None),
                source_chat_title=getattr(message.state.fsm_storage, "source_chat_title", None),
                archive_identity_matrix_uploaded=True,
            ),
        )
        await render_screen(
            message=message,
            bot=bot,
            body=_build_participant_migration_prompt(),
            bubbles=_participant_migration_keyboard(),
        )

    @fsm.on(MigrationWizardState.INPUT_ARCHIVE_PARTICIPANT_DECISION)
    async def input_archive_participant_decision(message: IncomingMessage, bot: Bot) -> None:
        choice = _normalized_choice(_message_text(message))
        if _is_cancel(choice):
            await _cancel(message, bot)
            return
        if _is_back(choice):
            await message.state.fsm.change_state(
                MigrationWizardState.INPUT_ARCHIVE_IDENTITY_MATRIX,
                ttl_seconds=WIZARD_TTL_SECONDS,
                **_wizard_state_payload(
                    message,
                    source_chat_id=getattr(message.state.fsm_storage, "source_chat_id", None),
                    source_chat_type=getattr(message.state.fsm_storage, "source_chat_type", None),
                    source_chat_title=getattr(message.state.fsm_storage, "source_chat_title", None),
                ),
            )
            await render_screen(
                message=message,
                bot=bot,
                body=_build_archive_identity_matrix_prompt(),
                bubbles=_archive_identity_matrix_keyboard(),
            )
            return
        if choice not in {"yes", "no"}:
            await bot.answer_message(
                "Нажми `Да`, `Нет`, `Назад` или `Отмена`.",
                wait_callback=False,
            )
            return
        try:
            result = await service.start_prepared_archive_import(
                operator=_operator_from_message(message),
                source_chat_id=message.state.fsm_storage.source_chat_id,
                migrate_members=choice == "yes",
                identity_matrix_uploaded=bool(
                    getattr(message.state.fsm_storage, "archive_identity_matrix_uploaded", False),
                ),
            )
        except USER_VISIBLE_WIZARD_ERRORS as error:
            await bot.answer_message(f"Ошибка: {error}", wait_callback=False)
            return
        notice = _build_archive_import_start_notice(
            source_chat_type=getattr(message.state.fsm_storage, "source_chat_type", "group"),
            source_chat_title=getattr(message.state.fsm_storage, "source_chat_title", "Telegram Export"),
            result=result,
        )
        await message.state.fsm.drop_state()
        await show_main_menu(
            bot,
            message=message,
            notice=notice,
        )

    @fsm.on(MigrationWizardState.INPUT_PRIVATE_PEER)
    async def input_private_peer(message: IncomingMessage, bot: Bot) -> None:
        body = _message_text(message)
        choice = _normalized_choice(body)
        flow_mode = getattr(message.state.fsm_storage, "flow_mode", "wizard")
        flow_style = _figma_flow_style(message)
        if _is_skip(choice):
            if flow_mode == "command":
                await _continue_private_chat_command_without_peer(
                    message=message,
                    bot=bot,
                    service=service,
                    default_batch_size=default_batch_size,
                )
                return
            if flow_style == FIGMA_FLOW_STYLE and _figma_flow_target(message) == FIGMA_FLOW_SINGLE_CHAT:
                await _start_single_chat_with_current_options(
                    message=message,
                    bot=bot,
                    service=service,
                    default_batch_size=default_batch_size,
                    overrides={"access_strategy": "none"},
                )
                return
            await _continue_wizard_after_private_resolution(
                message=message,
                bot=bot,
                source_chat_id=message.state.fsm_storage.source_chat_id,
            )
            return
        if _is_cancel(choice):
            await _cancel(message, bot)
            return
        try:
            parsed = _parse_private_peer_input(message)
            result = await service.map_private_chat_peer_identity(
                operator=_operator_from_message(message),
                source_chat_id=message.state.fsm_storage.source_chat_id,
                corporate_email=parsed.corporate_email,
                target_huid=parsed.target_huid,
            )
        except USER_VISIBLE_WIZARD_ERRORS as error:
            await bot.answer_message(
                f"Ошибка: {error}",
                wait_callback=False,
            )
            return
        if flow_mode == "command":
            await _continue_private_chat_command_after_mapping(
                message=message,
                bot=bot,
                service=service,
                result=result,
                default_batch_size=default_batch_size,
            )
            return
        if flow_style == FIGMA_FLOW_STYLE and _figma_flow_target(message) == FIGMA_FLOW_SINGLE_CHAT:
            await _start_single_chat_with_current_options(
                message=message,
                bot=bot,
                service=service,
                default_batch_size=default_batch_size,
            )
            return
        await _continue_wizard_after_private_resolution(
            message=message,
            bot=bot,
            source_chat_id=message.state.fsm_storage.source_chat_id,
        )

    @fsm.on(MigrationWizardState.INPUT_GROUP_TOPIC_STRATEGY)
    async def input_group_topic_strategy(message: IncomingMessage, bot: Bot) -> None:
        choice = _normalized_choice(_message_text(message))
        flow_mode = getattr(message.state.fsm_storage, "flow_mode", "wizard")
        flow_style = _figma_flow_style(message)
        flow_target = _figma_flow_target(message)
        if _is_cancel(choice):
            await _cancel(message, bot)
            return
        if _is_back(choice):
            if _figma_edit_field(message) == "topic_strategy":
                await _prompt_figma_edit_menu(message=message, bot=bot)
                return
            if flow_style == FIGMA_FLOW_STYLE and flow_target == FIGMA_FLOW_OPERATOR_DEFAULTS:
                await begin_operator_defaults_configuration(
                    message=message,
                    bot=bot,
                    service=service,
                )
                return
            if flow_style == FIGMA_FLOW_STYLE and flow_target == FIGMA_FLOW_MIGRATE_ALL:
                excluded_source_chat_ids = tuple(
                    getattr(message.state.fsm_storage, "excluded_source_chat_ids", ()) or ()
                )
                await _prompt_migrate_all_configuration(
                    message=message,
                    bot=bot,
                    service=service,
                    excluded_source_chat_ids=excluded_source_chat_ids,
                )
                return
            if flow_style == FIGMA_FLOW_STYLE and _figma_flow_target(message) == FIGMA_FLOW_SINGLE_CHAT:
                await _return_to_single_chat_configuration_origin(
                    message=message,
                    bot=bot,
                    service=service,
                )
                return
            await bot.answer_message("Нажми `Отмена`, чтобы выйти из этого шага.", wait_callback=False)
            return
        if choice not in {
            "single",
            "single_chat",
            "one",
            "one_chat",
            "all_in_one",
            "split",
            "split_by_topic",
            "discussion_chat",
        }:
            await bot.answer_message(
                "Выбери вариант обсуждений, `Назад` или `Отмена`.",
                wait_callback=False,
            )
            return
        normalized_topic_strategy = (
            "split_by_topic"
            if choice in {"split", "split_by_topic", "discussion_chat"}
            else "single_chat"
        )
        if flow_style == FIGMA_FLOW_STYLE and _figma_edit_field(message) == "topic_strategy":
            await _clear_figma_edit_field_and_prompt_confirm(
                message=message,
                bot=bot,
                topic_strategy=normalized_topic_strategy,
            )
            return
        options = _options_from_flow_state(
            message,
            default_batch_size=default_batch_size,
            topic_strategy=normalized_topic_strategy,
        )
        group_request = await service.prepare_group_chat_resolution(
            operator=_operator_from_message(message),
            source_chat_id=message.state.fsm_storage.source_chat_id,
            options=options,
        )
        if group_request is not None:
            await begin_group_chat_resolution(
                message=message,
                bot=bot,
                request=group_request,
                options=options if flow_mode == "command" else None,
            )
            return
        if flow_style == FIGMA_FLOW_STYLE and _figma_flow_target(message) == FIGMA_FLOW_SINGLE_CHAT:
            await _start_single_chat_with_current_options(
                message=message,
                bot=bot,
                service=service,
                default_batch_size=default_batch_size,
                overrides={"topic_strategy": normalized_topic_strategy},
            )
            return
        if flow_mode == "command":
            await _start_command_migration_from_state(
                message=message,
                bot=bot,
                service=service,
                default_batch_size=default_batch_size,
                overrides={"topic_strategy": normalized_topic_strategy},
            )
            return
        await _continue_wizard_after_group_resolution(
            message=message,
            bot=bot,
            source_chat_id=message.state.fsm_storage.source_chat_id,
            topic_strategy=normalized_topic_strategy,
            legacy_origin="group_topic",
        )

    @fsm.on(MigrationWizardState.INPUT_GROUP_IDENTITY_MATRIX)
    async def input_group_identity_matrix(message: IncomingMessage, bot: Bot) -> None:
        choice = _normalized_choice(_message_text(message))
        flow_mode = getattr(message.state.fsm_storage, "flow_mode", "wizard")
        flow_style = _figma_flow_style(message)
        if _is_back(choice):
            if flow_style == FIGMA_FLOW_STYLE and _figma_flow_target(message) == FIGMA_FLOW_SINGLE_CHAT:
                await _return_to_single_chat_configuration_origin(
                    message=message,
                    bot=bot,
                    service=service,
                )
                return
            await bot.answer_message("Нажми `Отмена`, чтобы остановить wizard.", wait_callback=False)
            return
        if _is_skip(choice):
            if flow_style == FIGMA_FLOW_STYLE and _figma_flow_target(message) == FIGMA_FLOW_SINGLE_CHAT:
                await _start_single_chat_with_current_options(
                    message=message,
                    bot=bot,
                    service=service,
                    default_batch_size=default_batch_size,
                    overrides={
                        "identity_policy": "skip_unresolved",
                        "access_strategy": "none",
                    },
                )
                return
            if flow_mode == "command":
                await _start_command_migration_from_state(
                    message=message,
                    bot=bot,
                    service=service,
                    default_batch_size=default_batch_size,
                    overrides={"identity_policy": "skip_unresolved"},
                )
                return
            await _continue_wizard_after_group_resolution(
                message=message,
                bot=bot,
                source_chat_id=message.state.fsm_storage.source_chat_id,
                identity_policy="skip_unresolved",
                legacy_origin="group_identity",
            )
            return
        if _is_cancel(choice):
            await _cancel(message, bot)
            return
        if message.file is None:
            await bot.answer_message(
                "Загрузи Excel-файл матрицы, отправь `/skip` для partial migration или `/cancel`.",
                wait_callback=False,
            )
            return
        if flow_style == FIGMA_FLOW_STYLE and _figma_flow_target(message) == FIGMA_FLOW_SINGLE_CHAT:
            await render_screen(
                message=message,
                bot=bot,
                body="Подождите, файл обрабатывается...",
            )
        try:
            import_result = await service.apply_group_identity_matrix(
                operator=_operator_from_message(message),
                source_chat_id=message.state.fsm_storage.source_chat_id,
                workbook_content=message.file.content,
            )
        except USER_VISIBLE_WIZARD_ERRORS as error:
            await bot.answer_message(f"Ошибка: {error}", wait_callback=False)
            return
        if flow_style == FIGMA_FLOW_STYLE and _figma_flow_target(message) == FIGMA_FLOW_SINGLE_CHAT:
            await message.state.fsm.change_state(
                MigrationWizardState.INPUT_PARTICIPANT_MIGRATION_DECISION,
                ttl_seconds=WIZARD_TTL_SECONDS,
                **_wizard_state_payload(
                    message,
                    identity_policy="matrix_uploaded",
                    source_chat_type="group",
                ),
            )
            await render_screen(
                message=message,
                bot=bot,
                body=_build_participant_migration_prompt(),
                bubbles=_participant_migration_keyboard(),
            )
            return
        if flow_mode == "command":
            await _continue_group_command_after_matrix(
                message=message,
                bot=bot,
                service=service,
                import_result=import_result,
                default_batch_size=default_batch_size,
            )
            return
        await _continue_wizard_after_group_resolution(
            message=message,
            bot=bot,
            source_chat_id=message.state.fsm_storage.source_chat_id,
            identity_policy="matrix_uploaded",
            legacy_origin="group_identity",
        )

    @fsm.on(MigrationWizardState.INPUT_CHANNEL_IDENTITY_MATRIX)
    async def input_channel_identity_matrix(message: IncomingMessage, bot: Bot) -> None:
        choice = _normalized_choice(_message_text(message))
        flow_mode = getattr(message.state.fsm_storage, "flow_mode", "wizard")
        flow_style = _figma_flow_style(message)
        if _is_back(choice):
            if flow_style == FIGMA_FLOW_STYLE and _figma_flow_target(message) == FIGMA_FLOW_SINGLE_CHAT:
                await _return_to_single_chat_configuration_origin(
                    message=message,
                    bot=bot,
                    service=service,
                )
                return
            await bot.answer_message("Нажми `Отмена`, чтобы остановить wizard.", wait_callback=False)
            return
        if _is_skip(choice):
            if flow_style == FIGMA_FLOW_STYLE and _figma_flow_target(message) == FIGMA_FLOW_SINGLE_CHAT:
                await _start_single_chat_with_current_options(
                    message=message,
                    bot=bot,
                    service=service,
                    default_batch_size=default_batch_size,
                    overrides={
                        "identity_policy": "skip_unresolved",
                        "access_strategy": "none",
                    },
                )
                return
            if flow_mode == "command":
                access_strategy = getattr(message.state.fsm_storage, "access_strategy", None)
                if access_strategy:
                    await _start_command_migration_from_state(
                        message=message,
                        bot=bot,
                        service=service,
                        default_batch_size=default_batch_size,
                        overrides={"identity_policy": "skip_unresolved"},
                    )
                    return
                await _prompt_channel_access_strategy(
                    message=message,
                    bot=bot,
                    identity_policy="skip_unresolved",
                )
                return
            await _prompt_channel_access_strategy(
                message=message,
                bot=bot,
                identity_policy="skip_unresolved",
            )
            return
        if _is_cancel(choice):
            await _cancel(message, bot)
            return
        if message.file is None:
            await bot.answer_message(
                "Загрузи Excel-файл матрицы, отправь `/skip` для partial migration или `/cancel`.",
                wait_callback=False,
            )
            return
        if flow_style == FIGMA_FLOW_STYLE and _figma_flow_target(message) == FIGMA_FLOW_SINGLE_CHAT:
            await render_screen(
                message=message,
                bot=bot,
                body="Подождите, файл обрабатывается...",
            )
        try:
            import_result = await service.apply_channel_identity_matrix(
                operator=_operator_from_message(message),
                source_chat_id=message.state.fsm_storage.source_chat_id,
                workbook_content=message.file.content,
            )
        except USER_VISIBLE_WIZARD_ERRORS as error:
            await bot.answer_message(f"Ошибка: {error}", wait_callback=False)
            return
        if flow_style == FIGMA_FLOW_STYLE and _figma_flow_target(message) == FIGMA_FLOW_SINGLE_CHAT:
            await message.state.fsm.change_state(
                MigrationWizardState.INPUT_PARTICIPANT_MIGRATION_DECISION,
                ttl_seconds=WIZARD_TTL_SECONDS,
                **_wizard_state_payload(
                    message,
                    identity_policy="matrix_uploaded",
                    source_chat_type="channel",
                ),
            )
            await render_screen(
                message=message,
                bot=bot,
                body=_build_participant_migration_prompt(),
                bubbles=_participant_migration_keyboard(),
            )
            return
        if flow_mode == "command":
            access_strategy = getattr(message.state.fsm_storage, "access_strategy", None)
            if access_strategy:
                await _start_command_migration_from_state(
                    message=message,
                    bot=bot,
                    service=service,
                    default_batch_size=default_batch_size,
                    overrides={"identity_policy": "matrix_uploaded"},
                )
                return
            await bot.answer_message(
                (
                    "Матрица соответствий импортирована.\n"
                    f"imported={import_result.imported_count}\n"
                    f"skipped={import_result.skipped_count}"
                ),
                wait_callback=False,
            )
            await _prompt_channel_access_strategy(
                message=message,
                bot=bot,
                identity_policy="matrix_uploaded",
            )
            return
        await bot.answer_message(
            (
                "Матрица соответствий импортирована.\n"
                f"imported={import_result.imported_count}\n"
                f"skipped={import_result.skipped_count}"
            ),
            wait_callback=False,
        )
        await _prompt_channel_access_strategy(
            message=message,
            bot=bot,
            identity_policy="matrix_uploaded",
        )

    @fsm.on(MigrationWizardState.INPUT_CHANNEL_ACCESS_STRATEGY)
    async def input_channel_access_strategy(message: IncomingMessage, bot: Bot) -> None:
        choice = _normalized_choice(_message_text(message))
        flow_mode = getattr(message.state.fsm_storage, "flow_mode", "wizard")
        flow_style = _figma_flow_style(message)
        if _is_cancel(choice):
            await _cancel(message, bot)
            return
        if choice not in {"direct_add", "direct", "invite_link", "invite", "link"}:
            await bot.answer_message(
                "Ожидаю `direct_add` или `invite_link`. Либо `cancel`.",
                wait_callback=False,
            )
            return
        normalized_access_strategy = (
            "invite_link" if choice in {"invite_link", "invite", "link"} else "direct_add"
        )
        if flow_mode == "command":
            await _start_command_migration_from_state(
                message=message,
                bot=bot,
                service=service,
                default_batch_size=default_batch_size,
                overrides={"access_strategy": normalized_access_strategy},
            )
            return
        if flow_style == FIGMA_FLOW_STYLE and _figma_flow_target(message) == FIGMA_FLOW_SINGLE_CHAT:
            await _start_single_chat_with_current_options(
                message=message,
                bot=bot,
                service=service,
                default_batch_size=default_batch_size,
                overrides={"access_strategy": normalized_access_strategy},
            )
            return
        await _continue_wizard_after_channel_resolution(
            message=message,
            bot=bot,
            source_chat_id=message.state.fsm_storage.source_chat_id,
            access_strategy=normalized_access_strategy,
            legacy_origin="channel_access",
        )

    @fsm.on(MigrationWizardState.INPUT_FROM)
    async def input_from(message: IncomingMessage, bot: Bot) -> None:
        body = _message_text(message)
        choice = _normalized_choice(body)
        if _is_cancel(choice):
            await _cancel(message, bot)
            return
        if _figma_flow_style(message) == FIGMA_FLOW_STYLE:
            if _is_back(choice):
                if _figma_edit_field(message) == "include_from":
                    await _prompt_figma_edit_menu(message=message, bot=bot)
                    return
                await _prompt_service_messages_info(message=message, bot=bot)
                return
            try:
                include_from = _parse_figma_date_value(body, end_of_day=False)
            except ConfigurationError as error:
                await bot.answer_message(f"Ошибка: {error}", wait_callback=False)
                return
            if _figma_edit_field(message) == "include_from":
                await _clear_figma_edit_field_and_prompt_confirm(
                    message=message,
                    bot=bot,
                    include_from=include_from,
                )
                return
            await message.state.fsm.change_state(
                MigrationWizardState.INPUT_TO,
                ttl_seconds=WIZARD_TTL_SECONDS,
                **_wizard_state_payload(
                    message,
                    source_chat_id=message.state.fsm_storage.source_chat_id,
                    include_from=include_from,
                ),
            )
            await render_screen(
                message=message,
                bot=bot,
                body=_build_figma_to_prompt(),
                bubbles=_figma_to_keyboard(),
            )
            return
        if _is_back(choice):
            await _return_legacy_from_to_origin(message=message, bot=bot)
            return
        include_from = None if choice in {"", "skip"} else body
        await _prompt_legacy_to(
            message=message,
            bot=bot,
            include_from=include_from,
        )

    @fsm.on(MigrationWizardState.INPUT_TO)
    async def input_to(message: IncomingMessage, bot: Bot) -> None:
        body = _message_text(message)
        choice = _normalized_choice(body)
        if _is_cancel(choice):
            await _cancel(message, bot)
            return
        if _figma_flow_style(message) == FIGMA_FLOW_STYLE:
            if _is_back(choice):
                if _figma_edit_field(message) == "include_to":
                    await _prompt_figma_edit_menu(message=message, bot=bot)
                    return
                await _prompt_figma_from(message=message, bot=bot)
                return
            try:
                include_to = _parse_figma_date_value(body, end_of_day=True)
            except ConfigurationError as error:
                await bot.answer_message(f"Ошибка: {error}", wait_callback=False)
                return
            if _figma_edit_field(message) == "include_to":
                await _clear_figma_edit_field_and_prompt_confirm(
                    message=message,
                    bot=bot,
                    include_to=include_to,
                )
                return
            await _prompt_figma_reply(message=message, bot=bot, include_to=include_to)
            return
        if _is_back(choice):
            await _prompt_legacy_from(
                message=message,
                bot=bot,
                source_chat_id=message.state.fsm_storage.source_chat_id,
                legacy_origin=getattr(message.state.fsm_storage, "legacy_origin", None),
            )
            return
        include_to = None if choice in {"", "skip"} else body
        await _prompt_legacy_media(
            message=message,
            bot=bot,
            include_to=include_to,
        )

    @fsm.on(MigrationWizardState.INPUT_MEDIA)
    async def input_media(message: IncomingMessage, bot: Bot) -> None:
        body = _message_text(message)
        choice = _normalized_choice(body)
        if _is_cancel(choice):
            await _cancel(message, bot)
            return
        if _figma_flow_style(message) == FIGMA_FLOW_STYLE:
            if _is_back(choice):
                if _figma_edit_field(message) == "migrate_media":
                    await _prompt_figma_edit_menu(message=message, bot=bot)
                    return
                if _figma_flow_target(message) == FIGMA_FLOW_MIGRATE_ALL:
                    excluded_source_chat_ids = tuple(
                        getattr(message.state.fsm_storage, "excluded_source_chat_ids", ()) or ()
                    )
                    await _prompt_migrate_all_configuration(
                        message=message,
                        bot=bot,
                        service=service,
                        excluded_source_chat_ids=excluded_source_chat_ids,
                    )
                    return
                if _figma_flow_target(message) == FIGMA_FLOW_OPERATOR_DEFAULTS:
                    await begin_operator_defaults_configuration(
                        message=message,
                        bot=bot,
                        service=service,
                    )
                    return
                await _begin_single_chat_configuration_choice(
                    message=message,
                    bot=bot,
                    service=service,
                    source_chat_id=message.state.fsm_storage.source_chat_id,
                    source_chat_type=getattr(message.state.fsm_storage, "source_chat_type", ""),
                    source_chat_title=getattr(message.state.fsm_storage, "source_chat_title", ""),
                )
                return
            if choice == "save":
                selected_media_kinds = tuple(
                    getattr(message.state.fsm_storage, "selected_media_kinds", ()) or ()
                )
                if _figma_edit_field(message) == "migrate_media":
                    await _clear_figma_edit_field_and_prompt_confirm(
                        message=message,
                        bot=bot,
                        selected_media_kinds=selected_media_kinds,
                        migrate_media=bool(selected_media_kinds),
                    )
                    return
                await _prompt_service_messages_info(
                    message=message,
                    bot=bot,
                    selected_media_kinds=selected_media_kinds,
                    migrate_media=bool(selected_media_kinds),
                )
                return
            toggle_media_kind = _parse_figma_media_toggle(choice)
            if toggle_media_kind is None:
                await bot.answer_message(
                    "Выбери типы вложений, затем нажми `Сохранить`, `Назад` или `Отмена`.",
                    wait_callback=False,
                )
                return
            selected_media_kinds = _toggle_figma_media_kind(
                getattr(message.state.fsm_storage, "selected_media_kinds", ()) or (),
                toggle_media_kind,
            )
            await message.state.fsm.change_state(
                MigrationWizardState.INPUT_MEDIA,
                ttl_seconds=WIZARD_TTL_SECONDS,
                **_wizard_state_payload(
                    message,
                    selected_media_kinds=selected_media_kinds,
                    migrate_media=bool(selected_media_kinds),
                ),
            )
            await render_screen(
                message=message,
                bot=bot,
                body=_build_figma_media_prompt(selected_media_kinds),
                bubbles=_figma_media_keyboard(selected_media_kinds),
            )
            return
        if _is_back(choice):
            await _prompt_legacy_to(
                message=message,
                bot=bot,
                include_from=getattr(message.state.fsm_storage, "include_from", None),
            )
            return
        migrate_media = _parse_on_off(choice)
        if migrate_media is None:
            await bot.answer_message(
                "Ожидаю `on` или `off`. Либо `cancel`.",
                wait_callback=False,
            )
            return
        await _prompt_legacy_reply(
            message=message,
            bot=bot,
            migrate_media=migrate_media,
        )

    @fsm.on(MigrationWizardState.INPUT_REPLY)
    async def input_reply(message: IncomingMessage, bot: Bot) -> None:
        body = _message_text(message)
        choice = _normalized_choice(body)
        if _is_cancel(choice):
            await _cancel(message, bot)
            return
        if _figma_flow_style(message) == FIGMA_FLOW_STYLE:
            if _is_back(choice):
                if _figma_edit_field(message) == "reply_mode":
                    await _prompt_figma_edit_menu(message=message, bot=bot)
                    return
                await _prompt_figma_to(message=message, bot=bot)
                return
            if choice == "default":
                flow_target = _figma_flow_target(message)
                reply_mode = getattr(message.state.fsm_storage, "default_reply_mode", "inline_quote")
                output_template = (
                    ""
                    if flow_target == FIGMA_FLOW_OPERATOR_DEFAULTS
                    else getattr(message.state.fsm_storage, "default_output_template", None)
                )
            elif choice in SUPPORTED_REPLY_MODES:
                reply_mode = choice
                output_template = None
            else:
                try:
                    output_template = validate_output_template(body)
                except ConfigurationError as error:
                    await bot.answer_message(
                        f"Ошибка: {error}",
                        wait_callback=False,
                    )
                    return
                reply_mode = getattr(
                    message.state.fsm_storage,
                    "reply_mode",
                    getattr(message.state.fsm_storage, "default_reply_mode", "inline_quote"),
                )
            if _figma_edit_field(message) == "reply_mode":
                await _clear_figma_edit_field_and_prompt_confirm(
                    message=message,
                    bot=bot,
                    reply_mode=reply_mode,
                    output_template=output_template,
                )
                return
            await _prompt_figma_confirm(
                message=message,
                bot=bot,
                reply_mode=reply_mode,
                output_template=output_template,
            )
            return
        if _is_back(choice):
            await _prompt_legacy_media(
                message=message,
                bot=bot,
                include_to=getattr(message.state.fsm_storage, "include_to", None),
            )
            return
        if choice not in {"inline_quote", "source_id", "none"}:
            await bot.answer_message(
                "Ожидаю `inline_quote`, `source_id` или `none`. Либо `cancel`.",
                wait_callback=False,
            )
            return
        await _prompt_legacy_confirm(
            message=message,
            bot=bot,
            reply_mode=choice,
        )

    @fsm.on(MigrationWizardState.CONFIRM)
    async def confirm(message: IncomingMessage, bot: Bot) -> None:
        body = _message_text(message)
        choice = _normalized_choice(body)
        if _is_cancel(choice):
            await _cancel(message, bot)
            return
        if _figma_flow_style(message) == FIGMA_FLOW_STYLE:
            if _is_back(choice):
                await _prompt_figma_reply(message=message, bot=bot)
                return
            if choice == "edit":
                await _prompt_figma_edit_menu(message=message, bot=bot)
                return
            if choice != "confirm":
                await bot.answer_message(
                    "Нажми `Все верно`, `Изменить`, `Назад` или `Отмена`.",
                    wait_callback=False,
                )
                return
            if _figma_flow_target(message) == FIGMA_FLOW_MIGRATE_ALL:
                try:
                    result = await service.start_migrate_all(
                        operator=_operator_from_message(message),
                        options=_options_from_flow_state(
                            message,
                            default_batch_size=default_batch_size,
                            progress_policy="resume",
                        ),
                    )
                except USER_VISIBLE_WIZARD_ERRORS as error:
                    await bot.answer_message(f"Ошибка: {error}", wait_callback=False)
                    return
                await message.state.fsm.drop_state()
                await show_connected_menu(
                    bot,
                    message=message,
                    notice=(
                        "Миграция истории сообщений чатов запущена.\n"
                        "Статус миграции можно посмотреть в главном меню по одноименной кнопке."
                        if result.status == "accepted"
                        else _format_background_operation_for_wizard(result)
                    ),
                )
                return
            if _figma_flow_target(message) == FIGMA_FLOW_OPERATOR_DEFAULTS:
                try:
                    await service.configure_operator_defaults(
                        operator=_operator_from_message(message),
                        options=_options_from_flow_state(
                            message,
                            default_batch_size=default_batch_size,
                        ),
                    )
                except USER_VISIBLE_WIZARD_ERRORS as error:
                    await bot.answer_message(f"Ошибка: {error}", wait_callback=False)
                    return
                await message.state.fsm.drop_state()
                await show_main_menu(
                    bot,
                    message=message,
                    notice="Конфигурация по умолчанию обновлена.",
                )
                return
            await _start_single_chat_figma_flow(
                message=message,
                bot=bot,
                service=service,
                default_batch_size=default_batch_size,
            )
            return
        if _is_back(choice):
            await _prompt_legacy_reply(
                message=message,
                bot=bot,
                migrate_media=getattr(message.state.fsm_storage, "migrate_media", True),
            )
            return
        if choice not in {"start", "resume"}:
            await bot.answer_message(
                "Отправь `start`, чтобы сохранить конфиг и запустить перенос, "
                "`resume`, чтобы явно продолжить найденный прогресс, или `cancel`.",
                wait_callback=False,
            )
            return

        run_options = MigrationRunOptions(
            include_from=getattr(message.state.fsm_storage, "include_from", None),
            include_to=getattr(message.state.fsm_storage, "include_to", None),
            migrate_media=getattr(message.state.fsm_storage, "migrate_media", True),
            reply_mode=getattr(message.state.fsm_storage, "reply_mode", "inline_quote"),
            source_backend=getattr(message.state.fsm_storage, "source_backend", None),
            topic_strategy=getattr(message.state.fsm_storage, "topic_strategy", None),
            identity_policy=getattr(message.state.fsm_storage, "identity_policy", None),
            access_strategy=getattr(message.state.fsm_storage, "access_strategy", None),
            progress_policy="resume" if choice == "resume" else "ask",
            batch_size=default_batch_size,
        )
        operator = _operator_from_message(message)
        source_chat_id = message.state.fsm_storage.source_chat_id

        try:
            await service.configure_chat(
                operator=operator,
                source_chat_id=source_chat_id,
                options=run_options,
            )
            result = await service.start_migrate_chat(
                operator=operator,
                source_chat_id=source_chat_id,
                options=run_options,
            )
        except USER_VISIBLE_WIZARD_ERRORS as error:
            await bot.answer_message(f"Ошибка: {error}", wait_callback=False)
            return

        if result.status == "blocked":
            await render_screen(
                message=message,
                bot=bot,
                body=_format_blocked_result(result),
                bubbles=_legacy_confirm_keyboard(),
            )
            return

        await message.state.fsm.drop_state()
        await bot.answer_message(
            _format_start_result(result),
            wait_callback=False,
        )

    @fsm.on(MigrationWizardState.EDIT_CUSTOM_CONFIGURATION)
    async def edit_custom_configuration(message: IncomingMessage, bot: Bot) -> None:
        choice = _normalized_choice(_message_text(message))
        if _is_cancel(choice):
            await _cancel(message, bot)
            return
        if _is_back(choice):
            await _prompt_figma_confirm(message=message, bot=bot)
            return
        if choice == "media":
            await _prompt_figma_media_selection(message=message, bot=bot, edit_field="migrate_media")
            return
        if choice == "service":
            await _prompt_service_messages_info(
                message=message,
                bot=bot,
                edit_field="service_messages",
            )
            return
        if choice == "topic" and _should_offer_topic_strategy_edit(message):
            await _prompt_figma_topic_strategy(
                message=message,
                bot=bot,
                edit_field="topic_strategy",
            )
            return
        if choice == "from":
            await _prompt_figma_from(message=message, bot=bot, edit_field="include_from")
            return
        if choice == "to":
            await _prompt_figma_to(message=message, bot=bot, edit_field="include_to")
            return
        if choice == "reply":
            await _prompt_figma_reply(message=message, bot=bot, edit_field="reply_mode")
            return
        await bot.answer_message(
            "Выбери параметр для изменения, `Назад` или `Отмена`.",
            wait_callback=False,
        )

    @fsm.on(MigrationWizardState.INPUT_PARTICIPANT_MIGRATION_DECISION)
    async def input_participant_migration_decision(message: IncomingMessage, bot: Bot) -> None:
        choice = _normalized_choice(_message_text(message))
        if _is_cancel(choice):
            await _cancel(message, bot)
            return
        if _is_back(choice):
            await _return_to_identity_upload_state(message=message, bot=bot)
            return
        if choice not in {"yes", "no"}:
            await bot.answer_message(
                "Нажми `Да`, `Нет`, `Назад` или `Отмена`.",
                wait_callback=False,
            )
            return
        access_strategy = getattr(message.state.fsm_storage, "access_strategy", None)
        if choice == "no":
            access_strategy = "none"
        elif access_strategy in {None, "", "none"}:
            access_strategy = "direct_add"
        await _start_single_chat_with_current_options(
            message=message,
            bot=bot,
            service=service,
            default_batch_size=default_batch_size,
            overrides={
                "identity_policy": getattr(message.state.fsm_storage, "identity_policy", "matrix_uploaded"),
                "access_strategy": access_strategy,
            },
        )

    return fsm


def _build_wizard_intro(chats: tuple[BotAvailableChat, ...]) -> str:
    if not chats:
        return "Доступные Telegram-чаты не найдены."
    lines = [
        "Wizard переноса чата.",
        "Шаг 1/5. Выбери source_chat_id и отправь его следующим сообщением.",
        "Можно отправить `list`, чтобы показать список еще раз, или `cancel`.",
        "",
        "Доступные чаты:",
    ]
    for chat in chats[:10]:
        lines.append(
            f"- {chat.source_chat_id} | {chat.source_chat_title} "
            f"| configured={'yes' if chat.configured else 'no'} "
            f"| progress={'yes' if chat.has_progress else 'no'}",
        )
    if len(chats) > 10:
        lines.append(f"... +{len(chats) - 10} chats")
    return "\n".join(lines)


def _build_members_wizard_intro(chats: tuple[BotMigratedChat, ...]) -> str:
    if not chats:
        return "У тебя пока нет мигрированных чатов, доступных для добавления участников."
    lines = [
        "Добавление участников.",
        "",
        "Шаг 1/2. Выберите чат и отправьте `source_chat_id` следующим сообщением.",
        "Можно также отправить `list`, чтобы показать список еще раз.",
        "",
        "Доступные чаты:",
    ]
    for chat in chats[:10]:
        lines.append(
            f"- {chat.source_chat_id} | {chat.source_chat_title} -> {chat.target_chat_title}",
        )
    if len(chats) > 10:
        lines.append(f"... +{len(chats) - 10} chats")
    return "\n".join(lines)


async def begin_chat_migration_selection(
    *,
    message: IncomingMessage,
    bot: Bot,
    chats: tuple[BotAvailableChat, ...],
) -> None:
    if not chats:
        await render_screen(
            message=message,
            bot=bot,
            body="Доступные Telegram-чаты не найдены.",
        )
        return
    await message.state.fsm.change_state(
        MigrationWizardState.SELECT_CHAT,
        ttl_seconds=WIZARD_TTL_SECONDS,
        **_wizard_state_payload(message),
    )
    await _show_chat_migration_selection(message=message, bot=bot, chats=chats)


async def begin_chat_members_selection(
    *,
    message: IncomingMessage,
    bot: Bot,
    chats: tuple[BotMigratedChat, ...],
) -> None:
    if not chats:
        await render_screen(
            message=message,
            bot=bot,
            body=_build_members_wizard_intro(chats),
        )
        return
    await message.state.fsm.change_state(
        MigrationWizardState.SELECT_MEMBERS_CHAT,
        ttl_seconds=WIZARD_TTL_SECONDS,
        **_wizard_state_payload(message),
    )
    await render_screen(
        message=message,
        bot=bot,
        body=_build_members_wizard_intro(chats),
        bubbles=_chat_members_selection_keyboard(),
    )


async def begin_archive_import(
    *,
    message: IncomingMessage,
    bot: Bot,
) -> None:
    await message.state.fsm.change_state(
        MigrationWizardState.INPUT_ARCHIVE_IMPORT,
        ttl_seconds=WIZARD_TTL_SECONDS,
        **_wizard_state_payload(message),
    )
    await render_screen(
        message=message,
        bot=bot,
        body=_build_archive_import_intro(),
        bubbles=_archive_import_keyboard(),
    )


async def _show_chat_migration_selection(
    *,
    message: IncomingMessage,
    bot: Bot,
    chats: tuple[BotAvailableChat, ...],
) -> None:
    attachment = None
    if len(chats) > 100:
        attachment = OutgoingAttachment(
            content=_build_available_chats_workbook(chats),
            filename="available_chats.xlsx",
        )
    await render_screen(
        message=message,
        bot=bot,
        body=_build_chat_selection_intro(chats),
        bubbles=_chat_selection_keyboard(),
        file=attachment,
    )


async def _return_to_chat_selection(
    *,
    message: IncomingMessage,
    bot: Bot,
    service: MigrationBotControlService,
) -> None:
    try:
        chats = await service.list_available_chats(
            operator=_operator_from_message(message),
            limit=1000,
            query=None,
        )
    except USER_VISIBLE_WIZARD_ERRORS as error:
        await bot.answer_message(f"Ошибка: {error}", wait_callback=False)
        return
    await begin_chat_migration_selection(
        message=message,
        bot=bot,
        chats=chats,
    )


async def _prompt_migrate_all_exclusions(
    *,
    message: IncomingMessage,
    bot: Bot,
) -> None:
    await message.state.fsm.change_state(
        MigrationWizardState.INPUT_MIGRATE_ALL_EXCLUSIONS,
        ttl_seconds=WIZARD_TTL_SECONDS,
        **_wizard_state_payload(message),
    )
    await render_screen(
        message=message,
        bot=bot,
        body=_build_migrate_all_exclusion_prompt(),
        bubbles=_migrate_all_exclusions_keyboard(),
    )


async def begin_operator_defaults_configuration(
    *,
    message: IncomingMessage,
    bot: Bot,
    service: MigrationBotControlService,
) -> None:
    try:
        operator_defaults = await service.show_operator_defaults(
            operator=_operator_from_message(message),
        )
    except USER_VISIBLE_WIZARD_ERRORS as error:
        await bot.answer_message(f"Ошибка: {error}", wait_callback=False)
        return
    await message.state.fsm.change_state(
        MigrationWizardState.INPUT_OPERATOR_DEFAULTS_CONFIGURATION,
        ttl_seconds=WIZARD_TTL_SECONDS,
        **_wizard_state_payload(
            message,
            flow_style=FIGMA_FLOW_STYLE,
            flow_target=FIGMA_FLOW_OPERATOR_DEFAULTS,
            include_from=operator_defaults.include_from,
            include_to=operator_defaults.include_to,
            migrate_media=operator_defaults.migrate_media,
            selected_media_kinds=_selected_media_kinds(
                migrate_media=operator_defaults.migrate_media,
                media_kinds=operator_defaults.media_kinds,
            ),
            service_messages=operator_defaults.service_messages,
            reply_mode=operator_defaults.reply_mode,
            output_template=operator_defaults.output_template,
            default_reply_mode=operator_defaults.base_reply_mode,
            default_output_template=operator_defaults.base_output_template,
            access_strategy=operator_defaults.access_strategy,
            topic_strategy=operator_defaults.topic_strategy,
            figma_custom_configured=operator_defaults.configured_via == "db",
        ),
    )
    await render_screen(
        message=message,
        bot=bot,
        body=_build_operator_defaults_prompt(operator_defaults),
        bubbles=_operator_defaults_configuration_keyboard(),
    )


async def _prompt_migrate_all_configuration(
    *,
    message: IncomingMessage,
    bot: Bot,
    service: MigrationBotControlService,
    excluded_source_chat_ids: tuple[str, ...],
) -> None:
    try:
        operator_defaults = await service.show_operator_defaults(
            operator=_operator_from_message(message),
        )
    except USER_VISIBLE_WIZARD_ERRORS as error:
        await bot.answer_message(f"Ошибка: {error}", wait_callback=False)
        return
    await message.state.fsm.change_state(
        MigrationWizardState.INPUT_MIGRATE_ALL_CONFIGURATION,
        ttl_seconds=WIZARD_TTL_SECONDS,
        **_wizard_state_payload(
            message,
            excluded_source_chat_ids=excluded_source_chat_ids,
            flow_style=FIGMA_FLOW_STYLE,
            flow_target=FIGMA_FLOW_MIGRATE_ALL,
            include_from=operator_defaults.include_from,
            include_to=operator_defaults.include_to,
            migrate_media=operator_defaults.migrate_media,
            reply_mode=operator_defaults.reply_mode,
            output_template=operator_defaults.output_template,
            default_reply_mode=operator_defaults.reply_mode,
            default_output_template=operator_defaults.output_template,
            service_messages=operator_defaults.service_messages,
            selected_media_kinds=_selected_media_kinds(
                migrate_media=operator_defaults.migrate_media,
                media_kinds=operator_defaults.media_kinds,
            ),
            access_strategy=operator_defaults.access_strategy,
            topic_strategy=operator_defaults.topic_strategy,
        ),
    )
    await render_screen(
        message=message,
        bot=bot,
        body=_build_migrate_all_configuration_prompt(
            excluded_source_chat_ids,
            operator_defaults=operator_defaults,
        ),
        bubbles=_migrate_all_configuration_keyboard(),
    )


def _build_chat_selection_intro(chats: tuple[BotAvailableChat, ...]) -> str:
    lines = [
        "Миграция истории сообщений осуществляется для одного чата/канала или для всех сразу.",
        "При выборе `Все чаты` можно исключить отдельные чаты перед запуском.",
        "",
        "Для выбора одного чата отправь его `source_chat_id` следующим сообщением.",
        "Чтобы показать список еще раз, отправь `list`.",
        "",
    ]
    if len(chats) > 100:
        lines.extend(
            [
                "Чатов больше 100, поэтому список приложен Excel-файлом.",
                "В первом столбце указан `source_chat_id`, во втором — название чата.",
            ],
        )
        return "\n".join(lines)
    lines.append("Список чатов:")
    for chat in chats:
        lines.append(f"- {chat.source_chat_id} | {chat.source_chat_title}")
    return "\n".join(lines)


def _build_available_chats_workbook(chats: tuple[BotAvailableChat, ...]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "available_chats"
    sheet.append(("source_chat_id", "source_chat_title", "source_chat_type"))
    for chat in chats:
        sheet.append((chat.source_chat_id, chat.source_chat_title, chat.source_chat_type))
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _build_migrate_all_exclusion_prompt() -> str:
    return (
        "Нужно исключить какие-либо чаты или каналы из текущей миграции?\n"
        "Отправь их id через запятую или нажми `Пропустить`."
    )


def _build_migrate_all_exclusion_confirmation(
    excluded_source_chat_ids: tuple[str, ...],
) -> str:
    lines = [
        "Вы хотите исключить из текущей миграции следующие чаты:",
        "",
    ]
    lines.extend(excluded_source_chat_ids)
    return "\n".join(lines)


def _build_migrate_all_configuration_prompt(
    excluded_source_chat_ids: tuple[str, ...],
    *,
    operator_defaults: BotOperatorDefaultsResult,
) -> str:
    lines = [
        _build_figma_configuration_summary(
            source_chat_type=None,
            selected_media_kinds=_selected_media_kinds(
                migrate_media=operator_defaults.migrate_media,
                media_kinds=operator_defaults.media_kinds,
            ),
            service_messages=operator_defaults.service_messages,
            include_from=operator_defaults.include_from,
            include_to=operator_defaults.include_to,
            reply_mode=operator_defaults.reply_mode,
            output_template=operator_defaults.output_template,
            topic_strategy=operator_defaults.topic_strategy,
            show_topic_strategy=True,
        ),
        "",
    ]
    if excluded_source_chat_ids:
        lines.append("Исключены чаты: " + ", ".join(excluded_source_chat_ids))
        lines.append("")
    lines.append(
        "Оставить конфигурацию по умолчанию для этой миграции или перейти к кастомной настройке?"
    )
    return "\n".join(lines)


def _build_operator_defaults_prompt(operator_defaults: BotOperatorDefaultsResult) -> str:
    return "\n".join(
        [
            _build_figma_configuration_summary(
                source_chat_type=None,
                selected_media_kinds=_selected_media_kinds(
                    migrate_media=operator_defaults.migrate_media,
                    media_kinds=operator_defaults.media_kinds,
                ),
                service_messages=operator_defaults.service_messages,
                include_from=operator_defaults.include_from,
                include_to=operator_defaults.include_to,
                reply_mode=operator_defaults.reply_mode,
                output_template=operator_defaults.output_template,
                topic_strategy=operator_defaults.topic_strategy,
                show_topic_strategy=True,
            ),
            "",
            "Желаете изменить конфигурацию по умолчанию?",
        ]
    )


async def _begin_single_chat_configuration_choice(
    *,
    message: IncomingMessage,
    bot: Bot,
    service: MigrationBotControlService,
    source_chat_id: str,
    source_chat_type: str,
    source_chat_title: str,
) -> None:
    try:
        operator_defaults = await service.show_operator_defaults(
            operator=_operator_from_message(message),
        )
    except USER_VISIBLE_WIZARD_ERRORS as error:
        await bot.answer_message(f"Ошибка: {error}", wait_callback=False)
        return
    await message.state.fsm.change_state(
        MigrationWizardState.INPUT_SINGLE_CHAT_CONFIGURATION,
        ttl_seconds=WIZARD_TTL_SECONDS,
        **_wizard_state_payload(
            message,
            flow_style=FIGMA_FLOW_STYLE,
            flow_target=FIGMA_FLOW_SINGLE_CHAT,
            source_chat_id=source_chat_id,
            source_chat_type=source_chat_type,
            source_chat_title=source_chat_title,
            include_from=operator_defaults.include_from,
            include_to=operator_defaults.include_to,
            migrate_media=operator_defaults.migrate_media,
            reply_mode=operator_defaults.reply_mode,
            output_template=operator_defaults.output_template,
            default_reply_mode=operator_defaults.reply_mode,
            default_output_template=operator_defaults.output_template,
            service_messages=operator_defaults.service_messages,
            access_strategy=operator_defaults.access_strategy,
            topic_strategy=operator_defaults.topic_strategy,
            selected_media_kinds=_selected_media_kinds(
                migrate_media=operator_defaults.migrate_media,
                media_kinds=operator_defaults.media_kinds,
            ),
            figma_custom_configured=False,
        ),
    )
    await render_screen(
        message=message,
        bot=bot,
        body=_build_single_chat_configuration_prompt(
            operator_defaults,
            source_chat_type=source_chat_type,
        ),
        bubbles=_single_chat_configuration_keyboard(),
    )


async def _start_single_chat_figma_flow(
    *,
    message: IncomingMessage,
    bot: Bot,
    service: MigrationBotControlService,
    default_batch_size: int,
) -> None:
    await _start_single_chat_with_current_options(
        message=message,
        bot=bot,
        service=service,
        default_batch_size=default_batch_size,
        overrides=None,
    )


async def _start_single_chat_with_current_options(
    *,
    message: IncomingMessage,
    bot: Bot,
    service: MigrationBotControlService,
    default_batch_size: int,
    overrides: dict[str, object] | None,
) -> None:
    source_chat_id = message.state.fsm_storage.source_chat_id
    options = _options_from_flow_state(
        message,
        default_batch_size=default_batch_size,
        progress_policy="resume",
        **(overrides or {}),
    )
    try:
        resolution_request = await service.prepare_private_chat_resolution(
            operator=_operator_from_message(message),
            source_chat_id=source_chat_id,
            options=options,
        )
        if resolution_request is not None:
            await begin_private_chat_resolution(
                message=message,
                bot=bot,
                request=resolution_request,
                options=None,
            )
            return
        group_resolution_request = await service.prepare_group_chat_resolution(
            operator=_operator_from_message(message),
            source_chat_id=source_chat_id,
            options=options,
        )
        if group_resolution_request is not None:
            await begin_group_chat_resolution(
                message=message,
                bot=bot,
                request=group_resolution_request,
                options=None,
            )
            return
        channel_resolution_request = await service.prepare_channel_chat_resolution(
            operator=_operator_from_message(message),
            source_chat_id=source_chat_id,
            options=options,
        )
        if channel_resolution_request is not None:
            await begin_channel_chat_resolution(
                message=message,
                bot=bot,
                request=channel_resolution_request,
                options=None,
            )
            return
        result = await service.start_migrate_chat(
            operator=_operator_from_message(message),
            source_chat_id=source_chat_id,
            options=options,
        )
    except USER_VISIBLE_WIZARD_ERRORS as error:
        await bot.answer_message(f"Ошибка: {error}", wait_callback=False)
        return
    await message.state.fsm.drop_state()
    await show_connected_menu(
        bot,
        message=message,
        notice=(
            _build_single_chat_started_notice(message)
            if result.status == "accepted"
            else _format_blocked_result(result)
            if result.status == "blocked"
            else _format_start_result(result)
        ),
    )


def _build_single_chat_configuration_prompt(
    operator_defaults: BotOperatorDefaultsResult,
    *,
    source_chat_type: str,
) -> str:
    lines = [
        _build_figma_configuration_summary(
            source_chat_type=source_chat_type,
            selected_media_kinds=_selected_media_kinds(
                migrate_media=operator_defaults.migrate_media,
                media_kinds=operator_defaults.media_kinds,
            ),
            service_messages=operator_defaults.service_messages,
            include_from=operator_defaults.include_from,
            include_to=operator_defaults.include_to,
            reply_mode=operator_defaults.reply_mode,
            output_template=operator_defaults.output_template,
            topic_strategy=operator_defaults.topic_strategy,
        ),
        "",
        "Оставить конфигурацию по умолчанию для этой миграции?",
    ]
    return "\n".join(lines)


def _build_single_chat_started_notice(message: IncomingMessage) -> str:
    source_chat_title = getattr(message.state.fsm_storage, "source_chat_title", None)
    source_chat_type = getattr(message.state.fsm_storage, "source_chat_type", "") or "chat"
    title = source_chat_title or "чата"
    if source_chat_type == "channel":
        subject = f'канала "{title}"'
    elif source_chat_type == "supergroup":
        subject = f'супергруппы "{title}"'
    elif source_chat_type == "group":
        subject = f'чата "{title}"'
    else:
        subject = f'чата "{title}"'
    return (
        f"Миграция истории сообщений {subject} запущена.\n"
        "Статус миграции можно посмотреть в главном меню по одноименной кнопке."
    )


def _build_figma_configuration_summary(
    *,
    source_chat_type: str | None,
    selected_media_kinds: tuple[str, ...],
    service_messages: bool,
    include_from: str | None,
    include_to: str | None,
    reply_mode: str | None,
    output_template: str | None,
    topic_strategy: str | None = None,
    show_topic_strategy: bool = False,
) -> str:
    lines = ["Настройки конфигурации по умолчанию:", ""]
    lines.append("Вложения: " + _format_figma_media_summary(selected_media_kinds))
    lines.append("Сервисные сообщения: " + ("да" if service_messages else "нет"))
    if show_topic_strategy or source_chat_type == "supergroup":
        lines.append(
            "Обсуждения в супергруппе: "
            + ("чат под обсуждение" if topic_strategy == "split_by_topic" else "в один чат")
        )
    lines.append("Дата начала: " + _format_figma_date_summary(include_from, default_label="первое сообщение"))
    lines.append("Дата окончания: " + _format_figma_date_summary(include_to, default_label="текущая"))
    if output_template:
        lines.append("Формат вывода:")
        lines.append(output_template)
    else:
        lines.append("Формат вывода: " + _format_reply_mode_label(reply_mode))
    return "\n".join(lines)


def _format_figma_media_summary(selected_media_kinds: tuple[str, ...]) -> str:
    if not selected_media_kinds:
        return "нет"
    return ", ".join(FIGMA_MEDIA_KIND_LABELS[kind] for kind in selected_media_kinds)


def _selected_media_kinds(
    *,
    migrate_media: bool,
    media_kinds: tuple[str, ...] | None,
) -> tuple[str, ...]:
    if not migrate_media:
        return ()
    if media_kinds is not None:
        return tuple(media_kinds)
    return DEFAULT_FIGMA_MEDIA_KINDS


def _format_figma_date_summary(value: str | None, *, default_label: str) -> str:
    if not value:
        return default_label
    try:
        parsed = datetime.fromisoformat(value.replace("Z", "+00:00"))
    except ValueError:
        return value
    return parsed.astimezone(UTC).strftime("%d.%m.%Y")


def _format_reply_mode_label(value: str | None) -> str:
    mapping = {
        "inline_quote": "по умолчанию",
        "source_id": "source_id",
        "none": "без reply",
    }
    return mapping.get((value or "inline_quote").strip().lower(), value or "по умолчанию")


def _single_chat_configuration_keyboard() -> BubbleMarkup:
    bubbles = BubbleMarkup()
    bubbles.add_button("Оставить по умолчанию", command="default")
    bubbles.add_button("Настроить кастомную", command="custom")
    bubbles.add_button("Назад", command="back")
    bubbles.add_button("Отмена", command="/cancel", new_row=False)
    return bubbles


def _build_figma_media_prompt(selected_media_kinds: tuple[str, ...]) -> str:
    lines = [
        "Выберите типы вложений, которые необходимо мигрировать.",
        "Если нужна миграция без вложений, нажмите сразу `Сохранить`.",
        "",
        "Сейчас выбрано: " + _format_figma_media_summary(selected_media_kinds),
    ]
    return "\n".join(lines)


def _figma_media_keyboard(selected_media_kinds: tuple[str, ...]) -> BubbleMarkup:
    selected = set(selected_media_kinds)
    keyboard = BubbleMarkup()
    for index, kind in enumerate(ALL_FIGMA_MEDIA_KINDS):
        prefix = "[x]" if kind in selected else "[ ]"
        keyboard.add_button(
            f"{prefix} {FIGMA_MEDIA_KIND_LABELS[kind]}",
            command=f"media:{kind}",
            new_row=index % 2 == 0,
        )
    keyboard.add_button("Сохранить", command="save")
    keyboard.add_button("Назад", command="back")
    keyboard.add_button("Отмена", command="/cancel", new_row=False)
    return keyboard


async def _prompt_figma_media_selection(
    *,
    message: IncomingMessage,
    bot: Bot,
    flow_target: str | None = None,
    edit_field: str | None = None,
) -> None:
    selected_media_kinds = tuple(
        getattr(message.state.fsm_storage, "selected_media_kinds", ()) or DEFAULT_FIGMA_MEDIA_KINDS
    )
    await message.state.fsm.change_state(
        MigrationWizardState.INPUT_MEDIA,
        ttl_seconds=WIZARD_TTL_SECONDS,
        **_wizard_state_payload(
            message,
            flow_style=FIGMA_FLOW_STYLE,
            flow_target=flow_target or _figma_flow_target(message),
            selected_media_kinds=selected_media_kinds,
            migrate_media=bool(selected_media_kinds),
            figma_custom_configured=(
                True
                if (flow_target or _figma_flow_target(message)) == FIGMA_FLOW_SINGLE_CHAT
                and edit_field is None
                else getattr(message.state.fsm_storage, "figma_custom_configured", False)
            ),
            figma_edit_field=edit_field,
        ),
    )
    await render_screen(
        message=message,
        bot=bot,
        body=_build_figma_media_prompt(selected_media_kinds),
        bubbles=_figma_media_keyboard(selected_media_kinds),
    )


async def _prompt_service_messages_info(
    *,
    message: IncomingMessage,
    bot: Bot,
    selected_media_kinds: tuple[str, ...] | None = None,
    migrate_media: bool | None = None,
    edit_field: str | None = None,
) -> None:
    await message.state.fsm.change_state(
        MigrationWizardState.INPUT_SERVICE_MESSAGES,
        ttl_seconds=WIZARD_TTL_SECONDS,
        **_wizard_state_payload(
            message,
            selected_media_kinds=selected_media_kinds
            if selected_media_kinds is not None
            else getattr(message.state.fsm_storage, "selected_media_kinds", ()),
            migrate_media=(
                migrate_media
                if migrate_media is not None
                else getattr(message.state.fsm_storage, "migrate_media", None)
            ),
            service_messages=getattr(message.state.fsm_storage, "service_messages", True),
            figma_edit_field=edit_field,
        ),
    )
    await render_screen(
        message=message,
        bot=bot,
        body="Необходимо ли переносить сервисные сообщения?",
        bubbles=_figma_service_messages_keyboard(),
    )


def _figma_service_messages_keyboard() -> BubbleMarkup:
    bubbles = BubbleMarkup()
    bubbles.add_button("Да", command="yes")
    bubbles.add_button("Нет", command="no")
    bubbles.add_button("Назад", command="back")
    bubbles.add_button("Отмена", command="/cancel", new_row=False)
    return bubbles


async def _prompt_figma_from(
    *,
    message: IncomingMessage,
    bot: Bot,
    edit_field: str | None = None,
    service_messages: bool | None = None,
) -> None:
    await message.state.fsm.change_state(
        MigrationWizardState.INPUT_FROM,
        ttl_seconds=WIZARD_TTL_SECONDS,
        **_wizard_state_payload(
            message,
            figma_edit_field=edit_field,
            service_messages=(
                service_messages
                if service_messages is not None
                else getattr(message.state.fsm_storage, "service_messages", True)
            ),
        ),
    )
    await render_screen(
        message=message,
        bot=bot,
        body=(
            "С какого периода нужно переносить сообщения?\n"
            "Нажми `С первого сообщения` или отправь дату в формате `дд.мм.гггг`."
        ),
        bubbles=_figma_from_keyboard(),
    )


def _figma_from_keyboard() -> BubbleMarkup:
    bubbles = BubbleMarkup()
    bubbles.add_button("С первого сообщения", command="first")
    bubbles.add_button("Назад", command="back")
    bubbles.add_button("Отмена", command="/cancel", new_row=False)
    return bubbles


def _build_figma_to_prompt() -> str:
    return (
        "До какой даты нужно переносить сообщения?\n"
        "Нажми `По текущий день` или отправь дату в формате `дд.мм.гггг`."
    )


def _build_figma_reply_prompt(message: IncomingMessage) -> str:
    current_output_template = getattr(message.state.fsm_storage, "output_template", None)
    current_reply_mode = getattr(message.state.fsm_storage, "reply_mode", "inline_quote")
    current_format = (
        current_output_template
        if current_output_template
        else _format_reply_mode_label(current_reply_mode)
    )
    placeholders = ", ".join(
        f"`{{{{{placeholder}}}}}`"
        for placeholder in sorted(OUTPUT_TEMPLATE_PLACEHOLDERS)
    )
    return (
        "Текущий формат вывода сообщения при миграции:\n"
        f"{current_format}\n\n"
        "Если нужно изменить формат, отправьте новую структуру следующим сообщением.\n"
        f"Доступные переменные: {placeholders}.\n"
        "`Оставить по умолчанию` вернет формат из текущего уровня defaults."
    )


async def _prompt_figma_to(
    *,
    message: IncomingMessage,
    bot: Bot,
    edit_field: str | None = None,
) -> None:
    await message.state.fsm.change_state(
        MigrationWizardState.INPUT_TO,
        ttl_seconds=WIZARD_TTL_SECONDS,
        **_wizard_state_payload(
            message,
            figma_edit_field=edit_field,
        ),
    )
    await render_screen(
        message=message,
        bot=bot,
        body=_build_figma_to_prompt(),
        bubbles=_figma_to_keyboard(),
    )


def _figma_to_keyboard() -> BubbleMarkup:
    bubbles = BubbleMarkup()
    bubbles.add_button("По текущий день", command="today")
    bubbles.add_button("Назад", command="back")
    bubbles.add_button("Отмена", command="/cancel", new_row=False)
    return bubbles


async def _prompt_figma_reply(
    *,
    message: IncomingMessage,
    bot: Bot,
    include_to: str | None = None,
    edit_field: str | None = None,
) -> None:
    await message.state.fsm.change_state(
        MigrationWizardState.INPUT_REPLY,
        ttl_seconds=WIZARD_TTL_SECONDS,
        **_wizard_state_payload(
            message,
            include_to=include_to
            if include_to is not None
            else getattr(message.state.fsm_storage, "include_to", None),
            figma_edit_field=edit_field,
        ),
    )
    await render_screen(
        message=message,
        bot=bot,
        body=_build_figma_reply_prompt(message),
        bubbles=_figma_reply_keyboard(),
    )


async def _prompt_figma_topic_strategy(
    *,
    message: IncomingMessage,
    bot: Bot,
    edit_field: str | None = None,
) -> None:
    await message.state.fsm.change_state(
        MigrationWizardState.INPUT_GROUP_TOPIC_STRATEGY,
        ttl_seconds=WIZARD_TTL_SECONDS,
        **_wizard_state_payload(message, figma_edit_field=edit_field),
    )
    await render_screen(
        message=message,
        bot=bot,
        body="Обсуждения из супергруппы переносить в один чат или под каждое обсуждение нужно создать новый чат?",
        bubbles=_group_topic_keyboard(figma=True),
    )


def _figma_reply_keyboard() -> BubbleMarkup:
    bubbles = BubbleMarkup()
    bubbles.add_button("Оставить по умолчанию", command="default")
    bubbles.add_button("Назад", command="back")
    bubbles.add_button("Отмена", command="/cancel", new_row=False)
    return bubbles


def _figma_confirm_keyboard() -> BubbleMarkup:
    bubbles = BubbleMarkup()
    bubbles.add_button("Все верно", command="confirm")
    bubbles.add_button("Изменить", command="edit")
    bubbles.add_button("Назад", command="back")
    bubbles.add_button("Отмена", command="/cancel", new_row=False)
    return bubbles


def _figma_edit_keyboard(message: IncomingMessage) -> BubbleMarkup:
    bubbles = BubbleMarkup()
    bubbles.add_button("Вложения", command="media")
    bubbles.add_button("Сервисные сообщения", command="service")
    if _should_offer_topic_strategy_edit(message):
        bubbles.add_button("Обсуждения в супергруппе", command="topic")
    bubbles.add_button("Дата начала", command="from")
    bubbles.add_button("Дата окончания", command="to")
    bubbles.add_button("Формат вывода", command="reply")
    bubbles.add_button("Назад", command="back")
    bubbles.add_button("Отмена", command="/cancel", new_row=False)
    return bubbles


def _participant_migration_keyboard() -> BubbleMarkup:
    bubbles = BubbleMarkup()
    bubbles.add_button("Да", command="yes")
    bubbles.add_button("Нет", command="no")
    bubbles.add_button("Назад", command="back")
    bubbles.add_button("Отмена", command="/cancel", new_row=False)
    return bubbles


def _archive_import_keyboard() -> BubbleMarkup:
    bubbles = BubbleMarkup()
    bubbles.add_button("Назад", command="back")
    bubbles.add_button("Отмена", command="/cancel")
    return bubbles


def _archive_identity_matrix_keyboard() -> BubbleMarkup:
    bubbles = BubbleMarkup()
    bubbles.add_button("Пропустить", command="skip")
    bubbles.add_button("Назад", command="back")
    bubbles.add_button("Отмена", command="/cancel", new_row=False)
    return bubbles

async def begin_chat_members_workbook(
    *,
    message: IncomingMessage,
    bot: Bot,
    request: BotChatMembersWorkbookRequest,
) -> None:
    await message.state.fsm.change_state(
        MigrationWizardState.INPUT_CHAT_MEMBERS_WORKBOOK,
        ttl_seconds=WIZARD_TTL_SECONDS,
        **_wizard_state_payload(
            message,
            source_chat_id=request.source_chat_id,
        ),
    )
    lines = [
        "Вам направлен файл с составом участников чата/канала.",
        "После заполнения нужных колонок отправьте файл в чат.",
        "",
        f"Источник: {request.source_chat_title} ({request.source_chat_id})",
        f"Целевой чат: {request.target_chat_title} ({request.target_chat_id})",
        f"Режим добавления: {request.access_strategy}",
        f"Уже сопоставлено: {request.mapped_rows}",
        f"Требуют ручного заполнения: {request.unresolved_rows}",
        "",
        "Можно заполнить колонки: `corporate_email`, `target_huid`, `ad_login`, `other_id`.",
        "Приоритет разрешения: `target_huid` -> `corporate_email` -> `ad_login` -> `other_id`.",
        "Можно добавить новых пользователей: оставьте Telegram-поля пустыми и заполните только корпоративный идентификатор.",
    ]
    if request.note:
        lines.extend(["", request.note])
    await render_screen(
        message=message,
        bot=bot,
        body="\n".join(lines),
        bubbles=_chat_members_workbook_keyboard(),
        file=OutgoingAttachment(
            content=request.workbook_content,
            filename=request.workbook_filename,
        ),
    )


def _format_chat_members_add_result(result: BotChatMembersAddResult) -> str:
    lines = [
        "Добавление участников завершено.",
        f"Источник: {result.source_chat_id}",
        f"Целевой чат: {result.target_chat_title} ({result.target_chat_id})",
        f"Запрошенный режим: {result.access_strategy}",
        f"Фактический результат: {result.effective_result}",
    ]
    if result.invite_fallback_used:
        lines.append(
            "Часть пользователей не удалось добавить напрямую, был использован fallback через invite link.",
        )
    lines.extend(
        [
            "",
            f"- обработано строк: {result.processed_rows}",
            f"- импортировано identity mappings: {result.imported_identity_mappings}",
            f"- строк без Telegram identity: {result.mapping_skipped_rows}",
            f"- resolved targets: {result.resolved_targets}",
            f"- добавлено напрямую: {result.direct_added}",
            f"- приглашено: {result.invited}",
            f"- пропущено: {result.skipped_rows}",
        ],
    )
    if result.failed_targets:
        lines.append(f"- не удалось добавить: {len(result.failed_targets)}")
        lines.extend(
            f"  - {label}"
            for label in result.failed_targets
        )
    return "\n".join(lines)


def _format_background_operation_for_wizard(result: BotOperationAcceptedResult) -> str:
    lines = [
        "Задача принята." if result.status == "accepted" else "Новая задача не запущена.",
        f"operation={result.operation}",
        f"migration_id={result.migration_id}",
    ]
    job_keys = result.job_keys or ((result.job_key,) if result.job_key else ())
    if len(job_keys) == 1:
        lines.append(f"job_key={job_keys[0]}")
    elif job_keys:
        lines.append(f"accepted_jobs={len(job_keys)}")
        lines.append("job_keys:")
        lines.extend(f"- {job_key}" for job_key in job_keys[:10])
        if len(job_keys) > 10:
            lines.append(f"... +{len(job_keys) - 10} jobs")
    if result.source_chat_ids:
        lines.append(f"source_chat_ids={','.join(result.source_chat_ids)}")
    if result.reason:
        lines.extend(["", f"reason={result.reason}"])
    lines.extend(["", "Проверка прогресса: /status"])
    return "\n".join(lines)


async def begin_private_chat_resolution(
    *,
    message: IncomingMessage,
    bot: Bot,
    request: BotPrivateChatResolutionRequest,
    options: MigrationRunOptions | None,
) -> None:
    payload = {
        "source_chat_id": request.source_chat_id,
        "flow_mode": "command" if options is not None else "wizard",
        "legacy_origin": "private_peer",
        "private_peer_display_name": request.peer_display_name,
        "private_peer_telegram_username": request.peer_telegram_username,
        "private_peer_telegram_user_id": request.peer_telegram_user_id,
    }
    if options is not None:
        payload.update(
            {
                "include_from": options.include_from,
                "include_to": options.include_to,
                "migrate_media": options.migrate_media,
                "reply_mode": options.reply_mode,
                "output_template": options.output_template,
                "source_backend": options.source_backend,
                "target_strategy": options.target_strategy,
                "target_title": options.target_title,
                "target_chat_id": options.target_chat_id,
                "identity_policy": options.identity_policy,
                "progress_policy": options.progress_policy,
                "batch_size": options.batch_size,
            },
        )
    await message.state.fsm.change_state(
        MigrationWizardState.INPUT_PRIVATE_PEER,
        ttl_seconds=WIZARD_TTL_SECONDS,
        **payload,
    )
    username = (
        f"@{request.peer_telegram_username}"
        if request.peer_telegram_username
        else "-"
    )
    await render_screen(
        message=message,
        bot=bot,
        body=(
            "Для персонального чата второй участник пока не сопоставлен.\n"
            f"source_chat_id={request.source_chat_id}\n"
            f"peer_name={request.peer_display_name}\n"
            f"peer_username={username}\n"
            f"peer_tg_id={request.peer_telegram_user_id or '-'}\n\n"
            "Отправь mention пользователя в eXpress, corporate email или huid. "
            "Отправь `/skip`, чтобы продолжить перенос без второго участника, "
            "или `/cancel`, чтобы остановить wizard."
        ),
        bubbles=_private_peer_keyboard(),
    )


async def begin_group_chat_resolution(
    *,
    message: IncomingMessage,
    bot: Bot,
    request: BotGroupChatResolutionRequest,
    options: MigrationRunOptions | None,
) -> None:
    payload = {
        "source_chat_id": request.source_chat_id,
        "flow_mode": "command" if options is not None else "wizard",
        "source_chat_type": "supergroup" if request.action == "choose_topic_strategy" else "group",
        "source_chat_title": request.source_chat_title,
    }
    if options is not None:
        payload.update(
            {
                "include_from": options.include_from,
                "include_to": options.include_to,
                "migrate_media": options.migrate_media,
                "reply_mode": options.reply_mode,
                "output_template": options.output_template,
                "source_backend": options.source_backend,
                "target_strategy": options.target_strategy,
                "target_title": options.target_title,
                "target_chat_id": options.target_chat_id,
                "identity_policy": options.identity_policy,
                "topic_strategy": options.topic_strategy,
                "progress_policy": options.progress_policy,
                "batch_size": options.batch_size,
            },
        )
    if request.action == "choose_topic_strategy":
        await message.state.fsm.change_state(
            MigrationWizardState.INPUT_GROUP_TOPIC_STRATEGY,
            ttl_seconds=WIZARD_TTL_SECONDS,
            **payload,
        )
        figma_flow = _figma_flow_style(message) == FIGMA_FLOW_STYLE
        lines = (
            [
                "Обсуждения из супергруппы переносить в один чат или под каждое обсуждение нужно создать новый чат?",
            ]
            if figma_flow
            else [
                "Для этой супергруппы обнаружены Telegram topics.",
                f"source_chat_id={request.source_chat_id}",
                f"source_chat_title={request.source_chat_title}",
                "",
                "Выбери стратегию: `single_chat` или `split_by_topic`.",
            ]
        )
        if request.topic_titles:
            lines.extend(["", "Темы:"])
            for title in request.topic_titles:
                lines.append(f"- {title}")
        await render_screen(
            message=message,
            bot=bot,
            body="\n".join(lines),
            bubbles=_group_topic_keyboard(figma=figma_flow),
        )
        return

    await message.state.fsm.change_state(
        MigrationWizardState.INPUT_GROUP_IDENTITY_MATRIX,
        ttl_seconds=WIZARD_TTL_SECONDS,
        **payload,
    )
    file = None
    if request.workbook_content is not None:
        file = OutgoingAttachment(
            content=request.workbook_content,
            filename=request.workbook_filename or "identity_matrix.xlsx",
        )
    figma_flow = _figma_flow_style(message) == FIGMA_FLOW_STYLE
    if figma_flow:
        lines = [
            "Вам направлен файл с составом участников чата/канала.",
            "Вы можете заполнить столбцы `express_email`, `express_ad_login`, `express_other_id`, `express_huid` для сопоставления пользователя в Telegram и пользователя в eXpress.",
            "После заполнения нужных столбцов направьте файл в чат.",
            "Можно нажать `Пропустить`, чтобы не загружать матрицу и сразу перейти к миграции сообщений.",
        ]
    else:
        lines = [
            "Есть участники Telegram-чата, которых пока не удалось сопоставить с eXpress.",
            f"source_chat_id={request.source_chat_id}",
            f"resolved={request.resolved_count}",
            f"unresolved={request.unresolved_count}",
            "",
            "Заполни Excel и загрузи его обратно, или отправь `/skip`, чтобы продолжить partial migration.",
            "`/cancel` остановит flow.",
        ]
    await render_screen(
        message=message,
        bot=bot,
        body="\n".join(lines),
        file=file,
        bubbles=_group_identity_keyboard(figma=figma_flow),
    )


async def begin_channel_chat_resolution(
    *,
    message: IncomingMessage,
    bot: Bot,
    request: BotChannelChatResolutionRequest,
    options: MigrationRunOptions | None,
) -> None:
    payload = {
        "source_chat_id": request.source_chat_id,
        "flow_mode": "command" if options is not None else "wizard",
        "source_chat_type": "channel",
        "source_chat_title": request.source_chat_title,
    }
    if options is not None:
        payload.update(
            {
                "include_from": options.include_from,
                "include_to": options.include_to,
                "migrate_media": options.migrate_media,
                "reply_mode": options.reply_mode,
                "output_template": options.output_template,
                "source_backend": options.source_backend,
                "target_strategy": options.target_strategy,
                "target_title": options.target_title,
                "target_chat_id": options.target_chat_id,
                "identity_policy": options.identity_policy,
                "access_strategy": options.access_strategy,
                "topic_strategy": options.topic_strategy,
                "progress_policy": options.progress_policy,
                "batch_size": options.batch_size,
            },
        )
    if request.action == "choose_access_strategy":
        await message.state.fsm.change_state(
            MigrationWizardState.INPUT_CHANNEL_ACCESS_STRATEGY,
            ttl_seconds=WIZARD_TTL_SECONDS,
            **payload,
        )
        lines = [
            "Для Telegram-канала выбери стратегию выдачи доступа.",
            f"source_chat_id={request.source_chat_id}",
            f"source_chat_title={request.source_chat_title}",
            "",
            "`direct_add` — попытаться добавить сопоставленных пользователей напрямую.",
            "`invite_link` — создать канал и подготовить инвайт-ссылку как fallback.",
        ]
        if not request.can_list_participants:
            lines.extend(
                [
                    "",
                    "Telegram не гарантирует полный список подписчиков. "
                    "Доступ будет выдан в best-effort режиме.",
                ],
            )
        await render_screen(
            message=message,
            bot=bot,
            body="\n".join(lines),
            bubbles=_channel_access_keyboard(),
        )
        return

    await message.state.fsm.change_state(
        MigrationWizardState.INPUT_CHANNEL_IDENTITY_MATRIX,
        ttl_seconds=WIZARD_TTL_SECONDS,
        **payload,
    )
    file = None
    if request.workbook_content is not None:
        file = OutgoingAttachment(
            content=request.workbook_content,
            filename=request.workbook_filename or "channel_identity_matrix.xlsx",
        )
    figma_flow = _figma_flow_style(message) == FIGMA_FLOW_STYLE
    if figma_flow:
        lines = [
            "Вам направлен файл с составом участников чата/канала.",
            "Вы можете заполнить столбцы `express_email`, `express_ad_login`, `express_other_id`, `express_huid` для сопоставления пользователя в Telegram и пользователя в eXpress.",
            "После заполнения нужных столбцов направьте файл в чат.",
            "Можно нажать `Пропустить`, чтобы не загружать матрицу и сразу перейти к миграции сообщений.",
        ]
    else:
        lines = [
            "Есть пользователи Telegram-канала, которых пока не удалось сопоставить с eXpress.",
            f"source_chat_id={request.source_chat_id}",
            f"resolved={request.resolved_count}",
            f"unresolved={request.unresolved_count}",
            "",
            "Заполни Excel и загрузи его обратно, или отправь `/skip`, чтобы продолжить partial migration.",
            "`/cancel` остановит flow.",
        ]
    if not request.can_list_participants and not figma_flow:
        lines.extend(
            [
                "",
                "Список подписчиков канала может быть неполным: Telegram отдает его только в best-effort режиме.",
            ],
        )
    await render_screen(
        message=message,
        bot=bot,
        body="\n".join(lines),
        file=file,
        bubbles=_group_identity_keyboard(figma=figma_flow),
    )


def _build_confirmation_prompt(
    *,
    source_chat_id: str,
    include_from: str | None,
    include_to: str | None,
    migrate_media: bool,
    reply_mode: str,
    access_strategy: str | None = None,
) -> str:
    lines = [
        "Проверь параметры:",
        f"source_chat_id={source_chat_id}",
        f"include_from={include_from or '-'}",
        f"include_to={include_to or '-'}",
        f"migrate_media={'on' if migrate_media else 'off'}",
        f"reply_mode={reply_mode}",
    ]
    if access_strategy:
        lines.append(f"access_strategy={access_strategy}")
    lines.extend(
        [
            "",
            "Отправь `start`, чтобы запустить перенос с проверкой existing progress, "
            "`resume`, чтобы явно продолжить уже начатый перенос, или `cancel`.",
        ],
    )
    return "\n".join(lines)


def _legacy_from_keyboard() -> BubbleMarkup:
    bubbles = BubbleMarkup()
    bubbles.add_button("Пропустить", command="skip")
    bubbles.add_button("Назад", command="back", new_row=False)
    bubbles.add_button("Отмена", command="/cancel")
    return bubbles


def _legacy_to_keyboard() -> BubbleMarkup:
    bubbles = BubbleMarkup()
    bubbles.add_button("Пропустить", command="skip")
    bubbles.add_button("Назад", command="back", new_row=False)
    bubbles.add_button("Отмена", command="/cancel")
    return bubbles


def _legacy_media_keyboard() -> BubbleMarkup:
    bubbles = BubbleMarkup()
    bubbles.add_button("Да", command="on")
    bubbles.add_button("Нет", command="off", new_row=False)
    bubbles.add_button("Назад", command="back")
    bubbles.add_button("Отмена", command="/cancel", new_row=False)
    return bubbles


def _legacy_reply_keyboard() -> BubbleMarkup:
    bubbles = BubbleMarkup()
    bubbles.add_button("inline_quote", command="inline_quote")
    bubbles.add_button("source_id", command="source_id", new_row=False)
    bubbles.add_button("none", command="none")
    bubbles.add_button("Назад", command="back")
    bubbles.add_button("Отмена", command="/cancel", new_row=False)
    return bubbles


def _legacy_confirm_keyboard() -> BubbleMarkup:
    bubbles = BubbleMarkup()
    bubbles.add_button("Запустить", command="start")
    bubbles.add_button("Продолжить", command="resume", new_row=False)
    bubbles.add_button("Назад", command="back")
    bubbles.add_button("Отмена", command="/cancel", new_row=False)
    return bubbles


async def _prompt_legacy_from(
    *,
    message: IncomingMessage,
    bot: Bot,
    source_chat_id: str,
    legacy_origin: str | None = None,
) -> None:
    await message.state.fsm.change_state(
        MigrationWizardState.INPUT_FROM,
        ttl_seconds=WIZARD_TTL_SECONDS,
        **_wizard_state_payload(
            message,
            source_chat_id=source_chat_id,
            legacy_origin=legacy_origin,
        ),
    )
    await render_screen(
        message=message,
        bot=bot,
        body="Шаг 2/5. Введи `from` в ISO-формате или `skip`.",
        bubbles=_legacy_from_keyboard(),
    )


async def _prompt_legacy_to(
    *,
    message: IncomingMessage,
    bot: Bot,
    include_from: str | None,
) -> None:
    await message.state.fsm.change_state(
        MigrationWizardState.INPUT_TO,
        ttl_seconds=WIZARD_TTL_SECONDS,
        **_wizard_state_payload(
            message,
            source_chat_id=message.state.fsm_storage.source_chat_id,
            include_from=include_from,
        ),
    )
    await render_screen(
        message=message,
        bot=bot,
        body="Шаг 3/5. Введи `to` в ISO-формате или `skip`.",
        bubbles=_legacy_to_keyboard(),
    )


async def _prompt_legacy_media(
    *,
    message: IncomingMessage,
    bot: Bot,
    include_to: str | None,
) -> None:
    await message.state.fsm.change_state(
        MigrationWizardState.INPUT_MEDIA,
        ttl_seconds=WIZARD_TTL_SECONDS,
        **_wizard_state_payload(
            message,
            source_chat_id=message.state.fsm_storage.source_chat_id,
            include_from=getattr(message.state.fsm_storage, "include_from", None),
            include_to=include_to,
        ),
    )
    await render_screen(
        message=message,
        bot=bot,
        body="Шаг 4/5. Переносить вложения?",
        bubbles=_legacy_media_keyboard(),
    )


async def _prompt_legacy_reply(
    *,
    message: IncomingMessage,
    bot: Bot,
    migrate_media: bool,
) -> None:
    await message.state.fsm.change_state(
        MigrationWizardState.INPUT_REPLY,
        ttl_seconds=WIZARD_TTL_SECONDS,
        **_wizard_state_payload(
            message,
            source_chat_id=message.state.fsm_storage.source_chat_id,
            include_from=getattr(message.state.fsm_storage, "include_from", None),
            include_to=getattr(message.state.fsm_storage, "include_to", None),
            migrate_media=migrate_media,
        ),
    )
    await render_screen(
        message=message,
        bot=bot,
        body="Шаг 5/5. Режим reply: `inline_quote`, `source_id` или `none`.",
        bubbles=_legacy_reply_keyboard(),
    )


async def _prompt_legacy_confirm(
    *,
    message: IncomingMessage,
    bot: Bot,
    reply_mode: str,
) -> None:
    await message.state.fsm.change_state(
        MigrationWizardState.CONFIRM,
        ttl_seconds=WIZARD_TTL_SECONDS,
        **_wizard_state_payload(
            message,
            source_chat_id=message.state.fsm_storage.source_chat_id,
            include_from=getattr(message.state.fsm_storage, "include_from", None),
            include_to=getattr(message.state.fsm_storage, "include_to", None),
            migrate_media=getattr(message.state.fsm_storage, "migrate_media", True),
            reply_mode=reply_mode,
        ),
    )
    await render_screen(
        message=message,
        bot=bot,
        body=_build_confirmation_prompt(
            source_chat_id=message.state.fsm_storage.source_chat_id,
            include_from=getattr(message.state.fsm_storage, "include_from", None),
            include_to=getattr(message.state.fsm_storage, "include_to", None),
            migrate_media=getattr(message.state.fsm_storage, "migrate_media", True),
            reply_mode=reply_mode,
            access_strategy=getattr(message.state.fsm_storage, "access_strategy", None),
        ),
        bubbles=_legacy_confirm_keyboard(),
    )


async def _return_legacy_from_to_origin(*, message: IncomingMessage, bot: Bot) -> None:
    legacy_origin = getattr(message.state.fsm_storage, "legacy_origin", None)
    if legacy_origin == "private_peer":
        await message.state.fsm.change_state(
            MigrationWizardState.INPUT_PRIVATE_PEER,
            ttl_seconds=WIZARD_TTL_SECONDS,
            **_wizard_state_payload(message),
        )
        username = getattr(message.state.fsm_storage, "private_peer_telegram_username", None)
        rendered_username = f"@{username}" if username else "-"
        await render_screen(
            message=message,
            bot=bot,
            body=(
                "Для персонального чата второй участник пока не сопоставлен.\n"
                f"source_chat_id={message.state.fsm_storage.source_chat_id}\n"
                f"peer_name={getattr(message.state.fsm_storage, 'private_peer_display_name', '-')}\n"
                f"peer_username={rendered_username}\n"
                f"peer_tg_id={getattr(message.state.fsm_storage, 'private_peer_telegram_user_id', '-') or '-'}\n\n"
                "Отправь mention пользователя в eXpress, corporate email или huid. "
                "Отправь `/skip`, чтобы продолжить перенос без второго участника, "
                "или `/cancel`, чтобы остановить wizard."
            ),
            bubbles=_private_peer_keyboard(),
        )
        return
    if legacy_origin == "channel_access":
        await _prompt_channel_access_strategy(message=message, bot=bot)
        return
    if legacy_origin == "group_identity":
        await _return_to_identity_upload_state(message=message, bot=bot)
        return
    await bot.answer_message("Назад недоступен на этом шаге.", wait_callback=False)


def _format_blocked_result(result: BotOperationAcceptedResult) -> str:
    lines = [
        "Обнаружен существующий прогресс. Wizard не продолжил перенос автоматически.",
        "",
        "Найденный прогресс:",
    ]
    for hint in result.progress_hints:
        lines.append(_format_progress_hint(hint))
    lines.append("")
    lines.append("Отправь `resume`, чтобы продолжить, или `cancel`.")
    return "\n".join(lines)


def _format_start_result(result: BotOperationAcceptedResult) -> str:
    if result.status == "already_migrated":
        lines = [
            "Чат уже был перенесен. Wizard не запускал новую задачу.",
            f"migration_id={result.migration_id}",
            f"source_chat_id={','.join(result.source_chat_ids)}",
        ]
        if result.topic_targets:
            lines.append("target_chats_by_topic:")
            for topic_target in result.topic_targets:
                lines.append(
                    f"- topic={topic_target.topic_title} "
                    f"| target_chat_id={topic_target.target_chat_id} "
                    f"| target_chat_link={topic_target.target_chat_link or '-'}",
                )
        else:
            lines.append(f"target_chat_id={result.target_chat_id or '-'}")
            if result.target_chat_link:
                lines.append(f"target_chat_link={result.target_chat_link}")
        return "\n".join(lines)
    if result.status == "accepted":
        return (
            "Wizard запустил перенос.\n"
            f"migration_id={result.migration_id}\n"
            f"job_key={result.job_key}\n"
            f"source_chat_id={','.join(result.source_chat_ids)}"
        )
    if result.status == "already_running":
        return "Новая задача не запущена: уже есть активный bot-managed job."
    return result.reason or "Операция заблокирована."


def _build_figma_confirm_prompt(message: IncomingMessage, payload: dict[str, object]) -> str:
    flow_target = str(
        payload.get("flow_target")
        or getattr(message.state.fsm_storage, "flow_target", "")
        or ""
    )
    return "\n".join(
        [
            _build_figma_configuration_summary(
                source_chat_type=str(
                    payload.get("source_chat_type")
                    or getattr(message.state.fsm_storage, "source_chat_type", "")
                ),
                selected_media_kinds=tuple(payload.get("selected_media_kinds", ()) or ()),
                include_from=payload.get("include_from")
                if isinstance(payload.get("include_from"), (str, type(None)))
                else None,
                include_to=payload.get("include_to")
                if isinstance(payload.get("include_to"), (str, type(None)))
                else None,
                service_messages=bool(payload.get("service_messages", True)),
                reply_mode=payload.get("reply_mode")
                if isinstance(payload.get("reply_mode"), (str, type(None)))
                else None,
                output_template=payload.get("output_template")
                if isinstance(payload.get("output_template"), (str, type(None)))
                else None,
                topic_strategy=payload.get("topic_strategy")
                if isinstance(payload.get("topic_strategy"), (str, type(None)))
                else None,
                show_topic_strategy=flow_target in {
                    FIGMA_FLOW_MIGRATE_ALL,
                    FIGMA_FLOW_OPERATOR_DEFAULTS,
                },
            ),
            "",
            (
                "Проверь конфигурацию по умолчанию."
                if flow_target == FIGMA_FLOW_OPERATOR_DEFAULTS
                else "Проверь настройки миграции."
            ),
        ]
    )


async def _prompt_figma_confirm(
    *,
    message: IncomingMessage,
    bot: Bot,
    reply_mode: str | None = None,
    output_template: str | None = None,
) -> None:
    payload = _wizard_state_payload(
        message,
        reply_mode=reply_mode
        if reply_mode is not None
        else getattr(message.state.fsm_storage, "reply_mode", None),
        output_template=output_template
        if output_template is not None
        else getattr(message.state.fsm_storage, "output_template", None),
        figma_edit_field=None,
    )
    await message.state.fsm.change_state(
        MigrationWizardState.CONFIRM,
        ttl_seconds=WIZARD_TTL_SECONDS,
        **payload,
    )
    await render_screen(
        message=message,
        bot=bot,
        body=_build_figma_confirm_prompt(message, payload),
        bubbles=_figma_confirm_keyboard(),
    )


async def _prompt_figma_edit_menu(*, message: IncomingMessage, bot: Bot) -> None:
    await message.state.fsm.change_state(
        MigrationWizardState.EDIT_CUSTOM_CONFIGURATION,
        ttl_seconds=WIZARD_TTL_SECONDS,
        **_wizard_state_payload(message, figma_edit_field=None),
    )
    await render_screen(
        message=message,
        bot=bot,
        body="Какой параметр нужно изменить?",
        bubbles=_figma_edit_keyboard(message),
    )


async def _clear_figma_edit_field_and_prompt_confirm(
    *,
    message: IncomingMessage,
    bot: Bot,
    **updates: object,
) -> None:
    await message.state.fsm.change_state(
        MigrationWizardState.CONFIRM,
        ttl_seconds=WIZARD_TTL_SECONDS,
        **_wizard_state_payload(message, figma_edit_field=None, **updates),
    )
    await _prompt_figma_confirm(message=message, bot=bot)


def _build_participant_migration_prompt() -> str:
    return "Необходимо ли мигрировать участников в созданный чат/канал?"


def _build_archive_import_intro() -> str:
    return (
        "Для миграции истории сообщений из Telegram export выполните следующие шаги:\n\n"
        "1. Откройте нужный групповой чат или канал.\n"
        "2. Нажмите на 3 точки в углу и выберите “Экспорт истории чата”.\n"
        "3. В модальном окне в поле “Формат” выберите “Машиночитаемый JSON”.\n"
        "4. Если планируете отправить только `.json`, снимите галочки со всех видов вложений.\n"
        "5. Нажмите “Экспортировать”.\n"
        "6. Найдите экспортированный `.json`, `.zip` или `.rar` у себя на устройстве.\n"
        "7. Отправьте файл в ответ на это сообщение.\n\n"
        "`.zip` и `.rar` позволяют перенести вложения, если они были включены в экспорт. "
        "Обычный `.json` импортируется без вложений."
    )


def _build_archive_identity_matrix_prompt() -> str:
    return (
        "Вам направлен пустой шаблон для заполнения таблицы соответствия пользователей в telegram "
        "и пользователей в express.\n"
        "Данная таблица необходима для добавления пользователей в создаваемый чат/канал в express "
        "и для корректных упоминаний.\n\n"
        "Заполните хотя бы один идентификатор пользователя из telegram "
        "(`telegram_user_id`, `telegram_username`, `telegram_display_name`) и хотя бы один "
        "идентификатор из express (`corporate_email`, `target_huid`, `ad_login`, `other_id`).\n\n"
        "Вы можете пропустить заполнение таблицы и миграция сообщений начнется сразу."
    )


def _archive_chat_label(source_chat_type: str) -> str:
    if source_chat_type == "channel":
        return "канала"
    if source_chat_type == "supergroup":
        return "супергруппы"
    return "чата"


def _build_archive_import_start_notice(
    *,
    source_chat_type: str,
    source_chat_title: str,
    result: BotOperationAcceptedResult,
) -> str:
    if result.status == "accepted":
        return (
            f"Миграция истории сообщений {_archive_chat_label(source_chat_type)} "
            f"“{source_chat_title}” запущена.\n"
            "Статус миграции можно посмотреть в главном меню по одноименной кнопке."
        )
    return _format_background_operation_for_wizard(result)


async def _return_to_identity_upload_state(*, message: IncomingMessage, bot: Bot) -> None:
    source_chat_type = getattr(message.state.fsm_storage, "source_chat_type", "")
    state = (
        MigrationWizardState.INPUT_CHANNEL_IDENTITY_MATRIX
        if source_chat_type == "channel"
        else MigrationWizardState.INPUT_GROUP_IDENTITY_MATRIX
    )
    await message.state.fsm.change_state(
        state,
        ttl_seconds=WIZARD_TTL_SECONDS,
        **_wizard_state_payload(message),
    )
    await render_screen(
        message=message,
        bot=bot,
        body="Загрузи файл матрицы заново, нажми `Пропустить`, `Назад` или `Отмена`.",
        bubbles=_group_identity_keyboard(),
    )


def _format_progress_hint(hint: BotProgressHint) -> str:
    return (
        f"- {hint.source_chat_id}: mapped={hint.mapped_total} "
        f"imported={hint.imported_count} "
        f"last_source_message_id={hint.last_source_message_id or '-'}"
    )


def _operator_from_message(message: IncomingMessage) -> BotOperatorContext:
    sender = message.sender
    username = getattr(sender, "username", None)
    ad_login = getattr(sender, "ad_login", None)
    ad_domain = getattr(sender, "ad_domain", None)
    display_name = ad_login or username or str(sender.huid)
    bot = getattr(message, "bot", None)
    bot_id = getattr(bot, "id", None)
    return BotOperatorContext(
        huid=str(sender.huid),
        chat_id=str(message.chat.id),
        username=username,
        ad_login=ad_login,
        ad_domain=ad_domain,
        display_name=display_name,
        current_cts_host=getattr(bot, "host", None),
        current_bot_id=str(bot_id) if bot_id is not None else None,
    )


async def _cancel(message: IncomingMessage, bot: Bot) -> None:
    await message.state.fsm.drop_state()
    await render_screen(message=message, bot=bot, body="Wizard остановлен.")


def _message_text(message: IncomingMessage) -> str:
    return (message.body or "").strip()


def _normalized_choice(value: str) -> str:
    return value.strip().lower()


def _parse_on_off(value: str) -> bool | None:
    if value in {"on", "true", "1", "yes"}:
        return True
    if value in {"off", "false", "0", "no"}:
        return False
    return None


def _parse_figma_date_value(raw_value: str, *, end_of_day: bool) -> str | None:
    choice = _normalized_choice(raw_value)
    if choice in {"", "skip", "first", "today"}:
        return None
    candidate = raw_value.strip()
    for parser in ("%d.%m.%Y", "%Y-%m-%d"):
        try:
            parsed = datetime.strptime(candidate, parser).replace(tzinfo=UTC)
            if end_of_day:
                parsed = parsed.replace(hour=23, minute=59, second=59)
            return parsed.isoformat().replace("+00:00", "Z")
        except ValueError:
            continue
    normalized = candidate.replace("Z", "+00:00")
    try:
        parsed = datetime.fromisoformat(normalized)
    except ValueError as error:
        raise ConfigurationError(
            "ожидаю дату в формате дд.мм.гггг или ISO-дату/датавремя",
        ) from error
    if parsed.tzinfo is None:
        parsed = parsed.replace(tzinfo=UTC)
    if end_of_day and len(candidate) <= 10:
        parsed = parsed.replace(hour=23, minute=59, second=59)
    return parsed.astimezone(UTC).isoformat().replace("+00:00", "Z")


def _is_cancel(value: str) -> bool:
    return value in {"cancel", "/cancel"}


def _is_skip(value: str) -> bool:
    return value in {"skip", "/skip"}


def _is_back(value: str) -> bool:
    return value in {"back", "/back"}


def _parse_figma_media_toggle(choice: str) -> str | None:
    if not choice.startswith("media:"):
        return None
    _, _, kind = choice.partition(":")
    normalized = kind.strip().lower()
    return normalized if normalized in FIGMA_MEDIA_KIND_LABELS else None


def _toggle_figma_media_kind(current: tuple[str, ...], kind: str) -> tuple[str, ...]:
    selected = list(current)
    if kind in selected:
        selected.remove(kind)
    else:
        selected.append(kind)
    return tuple(item for item in ALL_FIGMA_MEDIA_KINDS if item in selected)


def _parse_figma_reply_choice(choice: str, *, default_reply_mode: str) -> str | None:
    if choice == "default":
        normalized = (default_reply_mode or "inline_quote").strip().lower()
        return normalized if normalized in SUPPORTED_REPLY_MODES else "inline_quote"
    if choice in SUPPORTED_REPLY_MODES:
        return choice
    return None


def _figma_flow_style(message: IncomingMessage) -> str | None:
    return getattr(message.state.fsm_storage, "flow_style", None)


def _figma_flow_target(message: IncomingMessage) -> str | None:
    return getattr(message.state.fsm_storage, "flow_target", None)


def _figma_edit_field(message: IncomingMessage) -> str | None:
    return getattr(message.state.fsm_storage, "figma_edit_field", None)


def _should_offer_topic_strategy_edit(message: IncomingMessage) -> bool:
    flow_target = _figma_flow_target(message)
    if flow_target in {FIGMA_FLOW_MIGRATE_ALL, FIGMA_FLOW_OPERATOR_DEFAULTS}:
        return True
    return getattr(message.state.fsm_storage, "source_chat_type", None) == "supergroup"


async def _return_to_single_chat_configuration_origin(
    *,
    message: IncomingMessage,
    bot: Bot,
    service: MigrationBotControlService,
) -> None:
    if getattr(message.state.fsm_storage, "figma_custom_configured", False):
        await _prompt_figma_confirm(message=message, bot=bot)
        return
    await _begin_single_chat_configuration_choice(
        message=message,
        bot=bot,
        service=service,
        source_chat_id=message.state.fsm_storage.source_chat_id,
        source_chat_type=getattr(message.state.fsm_storage, "source_chat_type", ""),
        source_chat_title=getattr(message.state.fsm_storage, "source_chat_title", ""),
    )


async def _continue_wizard_after_private_resolution(
    *,
    message: IncomingMessage,
    bot: Bot,
    source_chat_id: str,
) -> None:
    await _prompt_legacy_from(
        message=message,
        bot=bot,
        source_chat_id=source_chat_id,
        legacy_origin="private_peer",
    )


async def _continue_wizard_after_group_resolution(
    *,
    message: IncomingMessage,
    bot: Bot,
    source_chat_id: str,
    identity_policy: str | None = None,
    topic_strategy: str | None = None,
    legacy_origin: str | None = None,
) -> None:
    await message.state.fsm.change_state(
        MigrationWizardState.INPUT_FROM,
        ttl_seconds=WIZARD_TTL_SECONDS,
        **_wizard_state_payload(
            message,
            source_chat_id=source_chat_id,
            identity_policy=identity_policy or getattr(message.state.fsm_storage, "identity_policy", None),
            topic_strategy=topic_strategy or getattr(message.state.fsm_storage, "topic_strategy", None),
            legacy_origin=legacy_origin or getattr(message.state.fsm_storage, "legacy_origin", None),
        ),
    )
    await render_screen(
        message=message,
        bot=bot,
        body="Шаг 2/5. Введи `from` в ISO-формате или `skip`.",
        bubbles=_legacy_from_keyboard(),
    )


async def _prompt_channel_access_strategy(
    *,
    message: IncomingMessage,
    bot: Bot,
    identity_policy: str | None = None,
) -> None:
    await message.state.fsm.change_state(
        MigrationWizardState.INPUT_CHANNEL_ACCESS_STRATEGY,
        ttl_seconds=WIZARD_TTL_SECONDS,
        **_wizard_state_payload(
            message,
            identity_policy=identity_policy or getattr(message.state.fsm_storage, "identity_policy", None),
        ),
    )
    await render_screen(
        message=message,
        bot=bot,
        body=(
            "Выбери стратегию доступа для канала.\n"
            "`direct_add` — попытаться добавить сопоставленных пользователей напрямую.\n"
            "`invite_link` — подготовить инвайт-ссылку и продолжить best-effort migration."
        ),
        bubbles=_channel_access_keyboard(),
    )


async def _continue_wizard_after_channel_resolution(
    *,
    message: IncomingMessage,
    bot: Bot,
    source_chat_id: str,
    access_strategy: str,
    legacy_origin: str | None = None,
) -> None:
    await message.state.fsm.change_state(
        MigrationWizardState.INPUT_FROM,
        ttl_seconds=WIZARD_TTL_SECONDS,
        **_wizard_state_payload(
            message,
            source_chat_id=source_chat_id,
            access_strategy=access_strategy,
            legacy_origin=legacy_origin or getattr(message.state.fsm_storage, "legacy_origin", None),
        ),
    )
    await render_screen(
        message=message,
        bot=bot,
        body="Шаг 2/5. Введи `from` в ISO-формате или `skip`.",
        bubbles=_legacy_from_keyboard(),
    )


async def _continue_private_chat_command_without_peer(
    *,
    message: IncomingMessage,
    bot: Bot,
    service: MigrationBotControlService,
    default_batch_size: int,
) -> None:
    options = _options_from_flow_state(
        message,
        default_batch_size=default_batch_size,
    )
    try:
        result = await service.start_migrate_chat(
            operator=_operator_from_message(message),
            source_chat_id=message.state.fsm_storage.source_chat_id,
            options=options,
        )
    except USER_VISIBLE_WIZARD_ERRORS as error:
        await bot.answer_message(f"Ошибка: {error}", wait_callback=False)
        return
    await message.state.fsm.drop_state()
    await bot.answer_message(
        _format_blocked_result(result)
        if result.status == "blocked"
        else _format_start_result(result),
        wait_callback=False,
    )


async def _continue_private_chat_command_after_mapping(
    *,
    message: IncomingMessage,
    bot: Bot,
    service: MigrationBotControlService,
    result,
    default_batch_size: int,
) -> None:
    options = _options_from_flow_state(
        message,
        default_batch_size=default_batch_size,
    )
    try:
        start_result = await service.start_migrate_chat(
            operator=_operator_from_message(message),
            source_chat_id=message.state.fsm_storage.source_chat_id,
            options=options,
        )
    except USER_VISIBLE_WIZARD_ERRORS as error:
        await bot.answer_message(f"Ошибка: {error}", wait_callback=False)
        return
    await message.state.fsm.drop_state()
    lines = [
        "Маппинг второго участника сохранен.",
        f"target_huid={result.target_huid or '-'}",
        f"corporate_email={result.corporate_email or '-'}",
        "",
        _format_blocked_result(start_result)
        if start_result.status == "blocked"
        else _format_start_result(start_result),
    ]
    await bot.answer_message("\n".join(lines), wait_callback=False)


async def _continue_group_command_after_matrix(
    *,
    message: IncomingMessage,
    bot: Bot,
    service: MigrationBotControlService,
    import_result: BotGroupIdentityImportResult,
    default_batch_size: int,
) -> None:
    try:
        start_result = await service.start_migrate_chat(
            operator=_operator_from_message(message),
            source_chat_id=message.state.fsm_storage.source_chat_id,
            options=_options_from_flow_state(
                message,
                default_batch_size=default_batch_size,
                identity_policy="matrix_uploaded",
            ),
        )
    except USER_VISIBLE_WIZARD_ERRORS as error:
        await bot.answer_message(f"Ошибка: {error}", wait_callback=False)
        return
    await message.state.fsm.drop_state()
    lines = [
        "Матрица соответствий импортирована.",
        f"imported={import_result.imported_count}",
        f"skipped={import_result.skipped_count}",
        "",
        _format_blocked_result(start_result)
        if start_result.status == "blocked"
        else _format_start_result(start_result),
    ]
    await bot.answer_message("\n".join(lines), wait_callback=False)


async def _start_command_migration_from_state(
    *,
    message: IncomingMessage,
    bot: Bot,
    service: MigrationBotControlService,
    default_batch_size: int,
    overrides: dict[str, object],
) -> None:
    try:
        result = await service.start_migrate_chat(
            operator=_operator_from_message(message),
            source_chat_id=message.state.fsm_storage.source_chat_id,
            options=_options_from_flow_state(
                message,
                default_batch_size=default_batch_size,
                **overrides,
            ),
        )
    except USER_VISIBLE_WIZARD_ERRORS as error:
        await bot.answer_message(f"Ошибка: {error}", wait_callback=False)
        return
    await message.state.fsm.drop_state()
    await bot.answer_message(
        _format_blocked_result(result) if result.status == "blocked" else _format_start_result(result),
        wait_callback=False,
    )


def _options_from_flow_state(
    message: IncomingMessage,
    *,
    default_batch_size: int,
    **overrides: object,
) -> MigrationRunOptions:
    return MigrationRunOptions(
        include_from=overrides.get(
            "include_from",
            getattr(message.state.fsm_storage, "include_from", None),
        ),
        include_to=overrides.get(
            "include_to",
            getattr(message.state.fsm_storage, "include_to", None),
        ),
        migrate_media=overrides.get(
            "migrate_media",
            getattr(message.state.fsm_storage, "migrate_media", None),
        ),
        media_kinds=overrides.get(
            "media_kinds",
            tuple(getattr(message.state.fsm_storage, "selected_media_kinds", ()) or ()),
        ),
        service_messages=overrides.get(
            "service_messages",
            getattr(message.state.fsm_storage, "service_messages", None),
        ),
        reply_mode=overrides.get(
            "reply_mode",
            getattr(message.state.fsm_storage, "reply_mode", None),
        ),
        output_template=overrides.get(
            "output_template",
            getattr(message.state.fsm_storage, "output_template", None),
        ),
        source_backend=overrides.get(
            "source_backend",
            getattr(message.state.fsm_storage, "source_backend", None),
        ),
        target_strategy=overrides.get(
            "target_strategy",
            getattr(message.state.fsm_storage, "target_strategy", None),
        ),
        target_title=overrides.get(
            "target_title",
            getattr(message.state.fsm_storage, "target_title", None),
        ),
        target_chat_id=overrides.get(
            "target_chat_id",
            getattr(message.state.fsm_storage, "target_chat_id", None),
        ),
        identity_policy=overrides.get(
            "identity_policy",
            getattr(message.state.fsm_storage, "identity_policy", None),
        ),
        access_strategy=overrides.get(
            "access_strategy",
            getattr(message.state.fsm_storage, "access_strategy", None),
        ),
        topic_strategy=overrides.get(
            "topic_strategy",
            getattr(message.state.fsm_storage, "topic_strategy", None),
        ),
        excluded_source_chat_ids=tuple(
            overrides.get(
                "excluded_source_chat_ids",
                getattr(message.state.fsm_storage, "excluded_source_chat_ids", ()),
            )
            or ()
        ),
        progress_policy=str(
            overrides.get(
                "progress_policy",
                getattr(message.state.fsm_storage, "progress_policy", "resume"),
            ),
        ),
        batch_size=int(
            overrides.get(
                "batch_size",
                getattr(message.state.fsm_storage, "batch_size", None) or default_batch_size,
            ),
        ),
    )


def _parse_private_peer_input(
    message: IncomingMessage,
) -> PrivatePeerResolutionInput:
    mention_users = list(getattr(getattr(message, "mentions", None), "users", []))
    if mention_users:
        return PrivatePeerResolutionInput(
            target_huid=str(mention_users[0].entity_id),
            source="mention",
        )
    body = _message_text(message)
    if not body:
        raise ConfigurationError("нужен mention, corporate email, huid, /skip или /cancel")
    parsed = parse_arguments(body)
    email = parsed.options.get("email")
    huid = parsed.options.get("huid") or parsed.options.get("user_huid")
    if email is None and huid is None and parsed.positionals:
        token = parsed.positionals[0]
        if _looks_like_email(token):
            email = token
        else:
            huid = token
    if email and huid:
        raise ConfigurationError("укажи либо email, либо huid/mention")
    if email:
        return PrivatePeerResolutionInput(
            corporate_email=email.strip(),
            source="email",
        )
    if huid:
        try:
            normalized_huid = str(UUID(huid.strip()))
        except ValueError as error:
            raise ConfigurationError("huid должен быть корректным UUID") from error
        return PrivatePeerResolutionInput(
            target_huid=normalized_huid,
            source="huid",
        )
    raise ConfigurationError("ожидаю mention пользователя в eXpress, corporate email или huid")


def _parse_chat_selection_message(message: IncomingMessage) -> str:
    mentions = getattr(message, "mentions", None)
    mention_chats = list(getattr(mentions, "chats", []))
    if mention_chats:
        return str(mention_chats[0].entity_id)
    mention_channels = list(getattr(mentions, "channels", []))
    if mention_channels:
        return str(mention_channels[0].entity_id)
    body = _message_text(message)
    if not body:
        raise ConfigurationError(
            "нужен source_chat_id или mention целевого eXpress чата; можно также отправить `list` или `/cancel`",
        )
    return body


def _looks_like_email(value: str) -> bool:
    candidate = value.strip()
    return "@" in candidate and " " not in candidate


def _parse_source_chat_id_list(value: str) -> tuple[str, ...]:
    normalized = value.replace("\n", ",").replace(";", ",")
    items: list[str] = []
    seen: set[str] = set()
    for raw_item in normalized.split(","):
        candidate = raw_item.strip()
        if not candidate or candidate in seen:
            continue
        items.append(candidate)
        seen.add(candidate)
    return tuple(items)


def _chat_selection_keyboard() -> BubbleMarkup:
    bubbles = BubbleMarkup()
    bubbles.add_button("Все чаты", command="all")
    bubbles.add_button("Назад", command="back", new_row=False)
    bubbles.add_button("Отмена", command="/cancel")
    return bubbles


def _migrate_all_exclusions_keyboard() -> BubbleMarkup:
    bubbles = BubbleMarkup()
    bubbles.add_button("Пропустить", command="/skip")
    bubbles.add_button("Назад", command="back", new_row=False)
    bubbles.add_button("Отмена", command="/cancel")
    return bubbles


def _migrate_all_exclusion_confirmation_keyboard() -> BubbleMarkup:
    bubbles = BubbleMarkup()
    bubbles.add_button("Все верно", command="confirm")
    bubbles.add_button("Изменить список", command="edit", new_row=False)
    bubbles.add_button("Отмена", command="/cancel")
    return bubbles


def _migrate_all_configuration_keyboard() -> BubbleMarkup:
    bubbles = BubbleMarkup()
    bubbles.add_button("Оставить по умолчанию", command="default")
    bubbles.add_button("Настроить кастомную", command="custom")
    bubbles.add_button("Назад", command="back")
    bubbles.add_button("Отмена", command="/cancel", new_row=False)
    return bubbles


def _chat_members_selection_keyboard() -> BubbleMarkup:
    bubbles = BubbleMarkup()
    bubbles.add_button("Показать список", command="list")
    bubbles.add_button("Назад", command="back", new_row=False)
    bubbles.add_button("Отмена", command="/cancel")
    return bubbles


def _chat_members_workbook_keyboard() -> BubbleMarkup:
    bubbles = BubbleMarkup()
    bubbles.add_button("Назад", command="back")
    bubbles.add_button("Отмена", command="/cancel", new_row=False)
    return bubbles


def _chat_members_result_keyboard() -> BubbleMarkup:
    bubbles = BubbleMarkup()
    bubbles.add_button("Добавление участников", command="/main_add_users")
    bubbles.add_button("Главное меню", command="/menu", new_row=False)
    return bubbles


def _operator_defaults_configuration_keyboard() -> BubbleMarkup:
    bubbles = BubbleMarkup()
    bubbles.add_button("Настроить кастомную", command="custom")
    bubbles.add_button("Главное меню", command="back")
    bubbles.add_button("Отмена", command="/cancel", new_row=False)
    return bubbles


def _private_peer_keyboard() -> BubbleMarkup:
    bubbles = BubbleMarkup()
    bubbles.add_button("Пропустить второго участника", command="/skip")
    bubbles.add_button("Отменить wizard", command="/cancel", new_row=False)
    return bubbles


def _group_topic_keyboard(*, figma: bool = False) -> BubbleMarkup:
    bubbles = BubbleMarkup()
    if figma:
        bubbles.add_button("Все в один чат", command="all_in_one")
        bubbles.add_button("Чат под обсуждение", command="discussion_chat")
        bubbles.add_button("Назад", command="back")
        bubbles.add_button("Отмена", command="/cancel", new_row=False)
        return bubbles
    bubbles.add_button("Один чат", command="single_chat")
    bubbles.add_button("По топикам", command="split_by_topic", new_row=False)
    bubbles.add_button("Отменить", command="/cancel")
    return bubbles


def _group_identity_keyboard(*, figma: bool = False) -> BubbleMarkup:
    bubbles = BubbleMarkup()
    bubbles.add_button("Пропустить", command="/skip")
    if figma:
        bubbles.add_button("Назад", command="back")
        bubbles.add_button("Отмена", command="/cancel", new_row=False)
        return bubbles
    bubbles.add_button("Отменить", command="/cancel", new_row=False)
    return bubbles


def _channel_access_keyboard() -> BubbleMarkup:
    bubbles = BubbleMarkup()
    bubbles.add_button("Добавить напрямую", command="direct_add")
    bubbles.add_button("Через ссылку", command="invite_link", new_row=False)
    bubbles.add_button("Отменить", command="/cancel")
    return bubbles


def _wizard_state_payload(message: IncomingMessage, **updates: object) -> dict[str, object]:
    fsm_storage = getattr(message.state, "fsm_storage", None)
    payload = {
        "flow_mode": getattr(fsm_storage, "flow_mode", "wizard"),
        "flow_style": getattr(fsm_storage, "flow_style", None),
        "flow_target": getattr(fsm_storage, "flow_target", None),
        "source_backend": getattr(fsm_storage, "source_backend", None),
        "source_chat_id": getattr(fsm_storage, "source_chat_id", None),
        "source_chat_type": getattr(fsm_storage, "source_chat_type", None),
        "source_chat_title": getattr(fsm_storage, "source_chat_title", None),
        "target_strategy": getattr(fsm_storage, "target_strategy", None),
        "target_title": getattr(fsm_storage, "target_title", None),
        "target_chat_id": getattr(fsm_storage, "target_chat_id", None),
        "identity_policy": getattr(fsm_storage, "identity_policy", None),
        "access_strategy": getattr(fsm_storage, "access_strategy", None),
        "topic_strategy": getattr(fsm_storage, "topic_strategy", None),
        "service_messages": getattr(fsm_storage, "service_messages", True),
        "selected_media_kinds": getattr(fsm_storage, "selected_media_kinds", ()),
        "default_reply_mode": getattr(fsm_storage, "default_reply_mode", "inline_quote"),
        "output_template": getattr(fsm_storage, "output_template", None),
        "default_output_template": getattr(fsm_storage, "default_output_template", None),
        "figma_edit_field": getattr(fsm_storage, "figma_edit_field", None),
        "figma_custom_configured": getattr(fsm_storage, "figma_custom_configured", False),
        "excluded_source_chat_ids": getattr(fsm_storage, "excluded_source_chat_ids", ()),
        "progress_policy": getattr(fsm_storage, "progress_policy", "ask"),
        "batch_size": getattr(fsm_storage, "batch_size", None),
        "legacy_origin": getattr(fsm_storage, "legacy_origin", None),
        "private_peer_display_name": getattr(fsm_storage, "private_peer_display_name", None),
        "private_peer_telegram_username": getattr(fsm_storage, "private_peer_telegram_username", None),
        "private_peer_telegram_user_id": getattr(fsm_storage, "private_peer_telegram_user_id", None),
        "screen_sync_id": getattr(fsm_storage, "screen_sync_id", None),
    }
    payload.update(updates)
    return payload
